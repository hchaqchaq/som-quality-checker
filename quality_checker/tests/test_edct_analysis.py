from __future__ import annotations

import re
import sqlite3
import tempfile
import unittest
from contextlib import closing
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo
from quality_checker.checkers.edct import runner as edct_analysis
from quality_checker.checkers.edct.config import (
    EDCT_COFOR_COLUMNS,
    EDCT_DATE_COLUMNS,
    EDCT_DATED_COMMENT_COLUMNS,
    EDCT_EDI_MODE_VALUES,
    EDCT_EMAIL_COLUMNS,
    EDCT_FORMULA_COLUMNS,
    EDCT_PN_REQUIRED_COLUMNS,
    EDCT_PORTAL_COLUMNS,
    EDCT_PORTAL_VALUES,
    EDCT_REQUIRED_COLUMNS,
    EDCT_RULE_CATALOGUE_ROWS,
    EDCT_TRIPLE_STATUS_VALUES,
    EDCT_UNCHECKED_COLUMNS,
)
from quality_checker.checkers.edct.runner import (
    EdctLoadError,
    export_edct_result,
    run_edct_analysis,
)
from quality_checker.db.repository import list_runs


def build_edct_workbook(
    directory: Path,
    *,
    include_open_task: bool = True,
    row_overrides: dict[int, dict[str, object]] | None = None,
    include_cofor_template: bool = True,
    optional_columns: tuple[str, ...] = (),
    pn_rows: tuple[tuple[object, object], ...] = (),
) -> Path:
    workbook = Workbook()
    all_columns = (*EDCT_REQUIRED_COLUMNS, *optional_columns)
    supplier = workbook.active
    supplier.title = "Supplier Level"
    supplier.append(["metadata"] * len(all_columns))
    supplier.append(list(all_columns))

    base: dict[str, object] = {column: "" for column in all_columns}
    for row_number, index in ((3, "Metz_01"), (4, "Metz_02")):
        values: dict[str, object] = dict(base)
        values.update(
            {
                "Index": index,
                "Supplier Punch code": 1000 + row_number,
                "Supplier name": f"Supplier {row_number}",
                "Effective kick-off date": "",
                "Cofor created date": "",
                "Overseas": "",
                "OPEN TASK": "YES" if row_number == 3 else "",
            }
        )
        for column in EDCT_FORMULA_COLUMNS:
            values[column] = f'=IF(A{row_number}="","",A{row_number})'
        values.update((row_overrides or {}).get(row_number, {}))
        supplier.append([values[column] for column in all_columns])

    supplier.append([""] * len(all_columns))
    supplier["A3"].fill = PatternFill("solid", fgColor="FF0000")
    table = Table(
        displayName="Tabella2",
        ref=f"A2:{get_column_letter(len(all_columns))}5",
    )
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    supplier.add_table(table)

    if include_open_task:
        open_task = workbook.create_sheet("Open Task")
        open_task.append(["metadata"])
        open_task.append(["Punch Code"])
        open_task.append([1003])
    if include_cofor_template:
        cofor_template = workbook.create_sheet("Template-Cofor-Creation")
        cofor_template["D2"] = "Punch Code"
        cofor_template["D3"] = "9999"
    workbook.create_sheet("Other Sheet")["A1"] = "keep me"
    pn = workbook.create_sheet("PN Level")
    pn.append(["Punch seller", "Triplet COFOR"])
    for seller, triplet in pn_rows:
        pn.append([seller, triplet])

    path = directory / "input.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def set_formula_caches(
    path: Path, caches: dict[str, dict[str, str]], *, data_type: str = "str"
) -> None:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    replacement = path.with_suffix(".zip")
    with ZipFile(path) as source, ZipFile(replacement, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename in caches:
                root = ElementTree.fromstring(content)
                for cell in root.iter(f"{{{namespace}}}c"):
                    if cell.attrib["r"] not in caches[item.filename]:
                        continue
                    cell.set("t", data_type)
                    value = cell.find(f"{{{namespace}}}v")
                    if value is None:
                        value = ElementTree.SubElement(cell, f"{{{namespace}}}v")
                    value.text = caches[item.filename][cell.attrib["r"]]
                content = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, content)
    replacement.replace(path)


def remap_archive_parts(path: Path) -> None:
    replacements = {
        "xl/tables/table1.xml": "xl/tables/table27.xml",
        "xl/worksheets/sheet5.xml": "xl/worksheets/sheet42.xml",
        "xl/worksheets/_rels/sheet5.xml.rels": "xl/worksheets/_rels/sheet42.xml.rels",
    }
    replacement = path.with_suffix(".zip")
    with ZipFile(path) as source, ZipFile(replacement, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename.endswith((".xml", ".rels")):
                content = content.replace(b"tables/table1.xml", b"tables/table27.xml")
                content = content.replace(b"worksheets/sheet5.xml", b"worksheets/sheet42.xml")
            if item.filename == "xl/tables/table1.xml":
                content = content.replace(b'id="1"', b'id="27"', 1)
            item.filename = replacements.get(item.filename, item.filename)
            target.writestr(item, content)
    replacement.replace(path)


def add_worksheet_extension(
    path: Path, marker: str, worksheet_part: str = "xl/worksheets/sheet1.xml"
) -> None:
    replacement = path.with_suffix(".zip")
    extension = (
        f'<extLst><ext uri="{marker}"><test:payload xmlns:test="urn:test">'
        "keep me</test:payload></ext></extLst>"
    ).encode()
    root_namespaces = (
        b' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
        b' mc:Ignorable="x14ac xr"'
        b' xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac"'
        b' xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision"'
        b' xr:uid="{97570D38-0785-4F6B-AB88-F8FBC149D0F1}"'
    )
    with ZipFile(path) as source, ZipFile(replacement, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == worksheet_part:
                root_end = content.index(b">", content.index(b"<worksheet"))
                content = content[:root_end] + root_namespaces + content[root_end:]
                content = content.replace(
                    b"<sheetFormatPr ",
                    b'<sheetFormatPr x14ac:dyDescent="0.25" ',
                    1,
                )
                content = content.replace(b"</worksheet>", extension + b"</worksheet>")
            target.writestr(item, content)
    replacement.replace(path)


def corrupt_first_cell_style(path: Path) -> None:
    replacement = path.with_suffix(".zip")
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with ZipFile(path) as source, ZipFile(replacement, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(content)
                cell = root.find(f".//{{{namespace}}}c")
                if cell is None:
                    raise AssertionError("Generated worksheet has no cells to corrupt")
                cell.set("s", "999999")
                content = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, content)
    replacement.replace(path)


def remap_table_relationship(path: Path, relationship_id: str) -> None:
    replacement = path.with_suffix(".zip")
    relationship_attribute = (
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    )
    with ZipFile(path) as source, ZipFile(replacement, "w", ZIP_DEFLATED) as target:
        relationship_root = ElementTree.fromstring(
            source.read("xl/worksheets/_rels/sheet1.xml.rels")
        )
        table_relationship = next(
            relationship
            for relationship in relationship_root
            if relationship.attrib["Type"].endswith("/table")
        )
        original_id = table_relationship.attrib["Id"]
        table_relationship.set("Id", relationship_id)

        worksheet_root = ElementTree.fromstring(source.read("xl/worksheets/sheet1.xml"))
        table_part = next(
            element
            for element in worksheet_root.iter()
            if element.tag.rsplit("}", 1)[-1] == "tablePart"
            and element.attrib[relationship_attribute] == original_id
        )
        table_part.set(relationship_attribute, relationship_id)

        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                content = ElementTree.tostring(
                    relationship_root,
                    encoding="utf-8",
                    xml_declaration=True,
                )
            elif item.filename == "xl/worksheets/sheet1.xml":
                content = ElementTree.tostring(
                    worksheet_root,
                    encoding="utf-8",
                    xml_declaration=True,
                )
            target.writestr(item, content)
    replacement.replace(path)


class EdctWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls._settings_patcher = patch.object(
            edct_analysis,
            "load_edct_settings",
            return_value=edct_analysis.EdctHeaderSettings.default(),
        )
        cls._settings_patcher.start()

    @classmethod
    def tearDownClass(cls) -> None:
        cls._settings_patcher.stop()

    def test_pn_required_structure_stops_analysis_and_records_failure(self) -> None:
        for missing in ("PN Level", "Punch seller", "Triplet COFOR"):
            with self.subTest(missing=missing), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(Path(temp))
                workbook = load_workbook(path)
                if missing == "PN Level":
                    del workbook["PN Level"]
                else:
                    column = 1 if missing == "Punch seller" else 2
                    workbook["PN Level"].cell(1, column).value = "Unexpected header"
                workbook.save(path)
                workbook.close()
                with closing(sqlite3.connect(":memory:")) as connection:
                    connection.row_factory = sqlite3.Row
                    with self.assertRaises(EdctLoadError) as raised:
                        run_edct_analysis(path, connection=connection)
                    self.assertIn(missing, str(raised.exception))
                    history = list_runs(connection, "eDCT")
                    self.assertEqual(len(history), 1)
                    self.assertEqual(history[0]["status"], "failed")
                    self.assertIsNone(history[0]["exported_file"])

    def test_required_formula_caches_fail_actionably_and_allow_recovery(self) -> None:
        scenarios = (
            ("Supplier Level", "Supplier Punch code", 3, "1003"),
            ("Supplier Level", "Triplet COFOR", 3, "A"),
            ("PN Level", "Punch seller", 2, "1003"),
            ("PN Level", "Triplet COFOR", 2, "A"),
        )
        for sheet_name, column_name, row, cached_value in scenarios:
            with self.subTest(sheet=sheet_name, column=column_name):
                with tempfile.TemporaryDirectory() as temp:
                    path = build_edct_workbook(
                        Path(temp),
                        row_overrides={
                            3: {"Triplet COFOR": "A"},
                            4: {"Triplet COFOR": "A"},
                        },
                        pn_rows=((1003, "A"), ("unknown", "=A3"), ("", "=A4")),
                    )
                    workbook = load_workbook(path)
                    sheet = workbook[sheet_name]
                    header_row = 1 if sheet_name == "PN Level" else 2
                    column = next(
                        cell.column for cell in sheet[header_row] if cell.value == column_name
                    )
                    cell = sheet.cell(row, column, '="cached identifier"')
                    coordinate = cell.coordinate
                    workbook.save(path)
                    workbook.close()
                    with closing(sqlite3.connect(":memory:")) as connection:
                        connection.row_factory = sqlite3.Row
                        with self.assertRaises(EdctLoadError) as raised:
                            run_edct_analysis(path, connection=connection)
                        self.assertIn(f"{sheet_name}!{coordinate}", str(raised.exception))
                        self.assertIn("Excel", str(raised.exception))
                        self.assertEqual(list_runs(connection, "eDCT")[0]["status"], "failed")
                        part = "sheet5" if sheet_name == "PN Level" else "sheet1"
                        set_formula_caches(
                            path,
                            {f"xl/worksheets/{part}.xml": {coordinate: "#REF!"}},
                            data_type="e",
                        )
                        with self.assertRaises(EdctLoadError) as cached_error:
                            run_edct_analysis(path, connection=connection)
                        self.assertIn(f"{sheet_name}!{coordinate}", str(cached_error.exception))
                        self.assertIn("#REF!", str(cached_error.exception))
                        set_formula_caches(
                            path, {f"xl/worksheets/{part}.xml": {coordinate: cached_value}}
                        )
                        result = run_edct_analysis(path, connection=connection)
                        self.addCleanup(result.workbook.close)
                        self.assertEqual(result.row_results[("PN Level", 2)].check, 0)
                        self.assertNotIn(("PN Level", 3), result.row_results)
                        self.assertNotIn(("PN Level", 4), result.row_results)
                        self.assertEqual(
                            {run["status"] for run in list_runs(connection, "eDCT")},
                            {"failed", "ok"},
                        )

    def test_pn_blank_scope_and_normalization_boundaries(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(
                directory,
                row_overrides={
                    3: {
                        "Supplier Punch code": " P1 ",
                        "Triplet COFOR": " A, B/C ",
                        "OPEN TASK": "",
                    },
                    4: {"Supplier Punch code": "p1", "Triplet COFOR": "a, b/c"},
                },
                pn_rows=(
                    ("p1", "a, b/c"),
                    (" P1 ", " A, B/C "),
                    ("P1", "A"),
                    ("P1", "a,b/c"),
                    ("P1", "a, b/c"),
                    ("P1", "bad"),
                    ("P1", "bad"),
                    ("P1", None),
                    (None, "a, b/c"),
                    ("unknown", "=uncached()"),
                    ("EMPTY", "A"),
                    ("EMPTY", None),
                    ("001", "001"),
                    ("001", "1"),
                    ("1", "001"),
                    ("P 1", "A"),
                    ("P-1", "A"),
                    ("EXCLUDED", "A"),
                ),
                optional_columns=("Line",),
            )
            workbook = load_workbook(path)
            supplier = workbook["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            additions = (
                ("Metz_03", "EMPTY", None, None),
                ("Metz_04", "", "A", None),
                ("Metz_05", "001", "001", None),
                (None, "P1", "bad", "fallback-must-not-apply"),
                (None, "EXCLUDED", "A", "fallback-must-not-apply"),
                ("Metz_06", "P1", "unused", None),
            )
            for row, (index, punch, triplet, line) in enumerate(additions, 5):
                for column, value in (
                    ("Index", index),
                    ("Supplier Punch code", punch),
                    ("Triplet COFOR", triplet),
                    ("Line", line),
                ):
                    supplier.cell(row, headers.index(column) + 1).value = value
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)
                self.addCleanup(result.workbook.close)
                pn_checks = {
                    row: result.row_results[(sheet, row)].check
                    for sheet, row in result.assessed_rows
                    if sheet == "PN Level"
                }
                self.assertEqual(
                    pn_checks,
                    {2: 0, 3: 0, 4: 1, 5: 1, 6: 0, 7: 1, 8: 1, 9: 1, 12: 1, 13: 1, 14: 0, 15: 1},
                )
                self.assertEqual(result.rule_totals[("pn_triplet_cofor", "Triplet COFOR")], 8)
                self.assertNotIn(("Supplier Level", 8), result.assessed_rows)
                self.assertNotIn(("Supplier Level", 9), result.assessed_rows)
                self.assertIn("blank", result.row_results[("PN Level", 9)].comment.lower())
                self.assertIn("no nonblank", result.row_results[("PN Level", 12)].comment.lower())
                comment = result.row_results[("PN Level", 4)].comment
                self.assertEqual(comment.count("'a, b/c'"), 1)
                self.assertLess(comment.index("'a, b/c'"), comment.index("'unused'"))

    def test_pn_formula_caches_tables_and_archive_parts_survive_export(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(
                directory,
                row_overrides={
                    3: {"Supplier Punch code": "P1", "OPEN TASK": ""},
                    4: {"Supplier Punch code": "P1"},
                },
                pn_rows=(("P1", "=A2"), ("P1", "=A3"), ("P1", "=A4")),
            )
            workbook = load_workbook(path)
            pn = workbook["PN Level"]
            pn.add_table(Table(displayName="NativePnData", ref="A1:B4"))
            pn["B2"].fill = PatternFill("solid", fgColor="AABBCC")
            pn["A2"].hyperlink = "https://example.com/supplier"
            supplier_headers = [cell.value for cell in workbook["Supplier Level"][2]]
            triplet_column = get_column_letter(supplier_headers.index("Triplet COFOR") + 1)
            workbook.save(path)
            workbook.close()
            add_worksheet_extension(path, "{PN-PRESERVATION}", "xl/worksheets/sheet5.xml")
            set_formula_caches(
                path,
                {
                    "xl/worksheets/sheet1.xml": {
                        f"{triplet_column}3": "A",
                        f"{triplet_column}4": "B",
                    },
                    "xl/worksheets/sheet5.xml": {"B2": " a ", "B3": "C", "B4": ""},
                },
            )
            remap_table_relationship(path, "rId99")
            remap_archive_parts(path)
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)
                self.addCleanup(result.workbook.close)
                self.assertEqual(result.rows_failed, 2)
                self.assertEqual(result.pn_values[2], ("P1", " a "))
                self.assertEqual(result.row_results[("PN Level", 4)].check, 1)
                output = export_edct_result(result, directory / "out")
                rerun = run_edct_analysis(output, connection=connection)
                self.addCleanup(rerun.workbook.close)
                self.assertEqual(rerun.row_results, result.row_results)
            with ZipFile(path) as source, ZipFile(output) as exported:
                self.assertEqual(set(source.namelist()), set(exported.namelist()))
                changed = {
                    "xl/worksheets/sheet1.xml",
                    "xl/worksheets/sheet42.xml",
                    "xl/tables/table27.xml",
                    "xl/styles.xml",
                }
                for name in source.namelist():
                    if name not in changed:
                        self.assertEqual(source.read(name), exported.read(name), name)
                self.assertIn(b"{PN-PRESERVATION}", exported.read("xl/worksheets/sheet42.xml"))
                worksheet_xml = exported.read("xl/worksheets/sheet42.xml")
                root_start = worksheet_xml[
                    worksheet_xml.index(
                        b"<", worksheet_xml.index(b"<?xml") + 5
                    ) : worksheet_xml.index(b">", worksheet_xml.index(b"<?xml") + 5) + 1
                ]
                ignorable = re.search(rb'Ignorable="([^"]+)"', root_start)
                if ignorable is not None:
                    for prefix in ignorable.group(1).split():
                        self.assertIn(b"xmlns:" + prefix + b"=", root_start)
            exported = load_workbook(output, data_only=False)
            self.addCleanup(exported.close)
            for sheet_name, header_row, source_row, expected_check in (
                ("Supplier Level", 2, 3, 0),
                ("PN Level", 1, 3, 1),
            ):
                sheet = exported[sheet_name]
                headers = [cell.value for cell in sheet[header_row]]
                self.assertIn("Check", headers)
                self.assertEqual(
                    sheet.cell(source_row, headers.index("Check") + 1).value, expected_check
                )
            self.assertEqual(exported["PN Level"]["B2"].value, "=A2")
            self.assertEqual(exported["PN Level"]["B2"].fill.fgColor.rgb, "00AABBCC")
            self.assertEqual(
                exported["PN Level"]["A2"].hyperlink.target, "https://example.com/supplier"
            )
            self.assertEqual(exported["PN Level"].tables["NativePnData"].ref, "A1:B4")
            cached = load_workbook(output, data_only=True)
            self.addCleanup(cached.close)
            self.assertEqual(cached["PN Level"]["B2"].value, " a ")

    def test_pn_membership_retains_sheet_identity_and_combined_history(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(
                directory,
                row_overrides={
                    3: {"Supplier Punch code": "P1", "Triplet COFOR": "A", "OPEN TASK": ""},
                    4: {"Supplier Punch code": "P1", "Triplet COFOR": "B"},
                },
                pn_rows=((" p1 ", " a "), ("P1", "C"), ("unknown", "C"), ("", "A")),
            )
            original_bytes = path.read_bytes()
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(
                    path, connection=connection, analysis_date=date(2026, 7, 15)
                )
                self.addCleanup(result.workbook.close)
                self.assertEqual(
                    result.assessed_rows,
                    (
                        ("Supplier Level", 3),
                        ("Supplier Level", 4),
                        ("PN Level", 2),
                        ("PN Level", 3),
                    ),
                )
                self.assertEqual(result.row_results[("PN Level", 2)].check, 0)
                rejected = result.row_results[("PN Level", 3)]
                self.assertEqual(rejected.check, 1)
                for detail in ("Punch seller", "P1", "Triplet COFOR", "C", "a", "b"):
                    self.assertIn(detail, rejected.comment)
                self.assertEqual(result.rows_failed, 3)
                self.assertEqual(sum(row.check for row in result.row_results.values()), 3)
                self.assertEqual(result.rule_totals[("pn_triplet_cofor", "Triplet COFOR")], 1)
                self.assertEqual(
                    connection.execute(
                        "SELECT rows_total, rows_in_scope, rows_failed, status, exported_file FROM runs"
                    ).fetchone(),
                    (4, 4, 3, "ok", None),
                )
                output = export_edct_result(result, directory / "out")
                self.assertEqual(
                    connection.execute("SELECT exported_file FROM runs").fetchone(), (str(output),)
                )
            exported = load_workbook(output, data_only=False)
            self.addCleanup(exported.close)
            for sheet_name, row in result.assessed_rows:
                sheet = exported[sheet_name]
                headers = [cell.value for cell in sheet[1 if sheet_name == "PN Level" else 2]]
                self.assertEqual(
                    sheet.cell(row, headers.index("Check") + 1).value,
                    result.row_results[(sheet_name, row)].check,
                )
            self.assertEqual(exported["PN Level"]["A2"].value, " p1 ")
            self.assertEqual(exported["PN Level"]["B2"].value, " a ")
            self.assertIsNone(exported["PN Level"]["C4"].value)
            self.assertEqual(path.read_bytes(), original_bytes)

    def test_missing_path_and_combined_structure_errors_are_recorded(self) -> None:
        with closing(sqlite3.connect(":memory:")) as connection:
            with self.assertRaisesRegex(EdctLoadError, "Input file not found"):
                run_edct_analysis("missing.xlsx", connection=connection)

        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "broken.xlsx"
            workbook = Workbook()
            workbook.active.title = "Supplier Level"
            workbook.active.append(["metadata"])
            workbook.active.append(["Wrong header"])
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                with self.assertRaises(EdctLoadError) as raised:
                    run_edct_analysis(path, connection=connection)
                for missing in ("Open Task", "Template-Cofor-Creation", "PN Level", "Index"):
                    self.assertIn(missing, str(raised.exception))

    def test_line_header_can_define_assessed_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    3: {"Supplier Punch code": "P1", "Triplet COFOR": "A", "OPEN TASK": ""},
                    4: {"Supplier Punch code": "P1", "Triplet COFOR": "B"},
                },
                pn_rows=(("P1", "B"), ("P1", "C")),
            )
            workbook = load_workbook(path)
            supplier = workbook["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            supplier.cell(2, headers.index("Index") + 1).value = "Line"
            workbook.save(path)
            workbook.close()

            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)

            self.assertEqual(
                result.assessed_rows,
                (("Supplier Level", 3), ("Supplier Level", 4), ("PN Level", 2), ("PN Level", 3)),
            )
            self.assertEqual(result.row_results[("PN Level", 2)].check, 0)
            self.assertEqual(result.row_results[("PN Level", 3)].check, 1)
            result.workbook.close()

    def test_open_task_header_and_invalid_overseas_are_reported(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(directory, row_overrides={3: {"Overseas": "MAYBE"}})
            workbook = load_workbook(path)
            workbook["Open Task"]["A2"] = "Wrong"
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                with self.assertRaisesRegex(EdctLoadError, "Open Task.Punch Code"):
                    run_edct_analysis(path, connection=connection)

            path = build_edct_workbook(directory, row_overrides={3: {"Overseas": "MAYBE"}})
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)
            self.assertEqual(result.rule_totals[("overseas", "Overseas")], 1)
            result.workbook.close()

    def test_columns_only_structure_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "columns.xlsx"
            workbook = Workbook()
            workbook.active.title = "Supplier Level"
            workbook.active.append(["metadata"])
            workbook.active.append(["Index"])
            open_task = workbook.create_sheet("Open Task")
            open_task.append(["metadata"])
            open_task.append(["Punch Code"])
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                with self.assertRaisesRegex(EdctLoadError, "columns:"):
                    run_edct_analysis(path, connection=connection)

    def test_default_history_records_export(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            with patch.object(edct_analysis, "DB_PATH", directory / "history.db"):
                result = run_edct_analysis(input_path, analysis_date=date(2026, 7, 15))
                output = export_edct_result(result, directory / "out")
            with closing(sqlite3.connect(directory / "history.db")) as connection:
                self.assertEqual(
                    connection.execute("SELECT exported_file, status FROM runs").fetchone(),
                    (str(output), "ok"),
                )
            result.workbook.close()

    def test_empty_assessment_and_missing_supplier(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(
                directory,
                row_overrides={3: {"Index": ""}, 4: {"Index": ""}},
                pn_rows=(("P1", "A"),),
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)
                self.assertEqual(result.assessed_rows, ())
                self.assertEqual(result.rows_failed, 0)
                self.assertEqual(
                    connection.execute("SELECT rows_in_scope FROM runs").fetchone(), (0,)
                )
                result.workbook.close()
            workbook = load_workbook(path)
            del workbook["Supplier Level"]
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                with self.assertRaisesRegex(EdctLoadError, "Supplier Level"):
                    run_edct_analysis(path, connection=connection)

    def test_rule_documentation_lists_every_required_edct_column(self) -> None:
        documentation = (
            Path(__file__).resolve().parents[2] / "docs" / "EDCT_QUALITY_CHECKER.md"
        ).read_text(encoding="utf-8")
        inventory = documentation.split("## Configured column inventory", 1)[1].split(
            "## Formula rules", 1
        )[0]
        documented = set(re.findall(r"^- `([^`]+)`$", inventory, re.MULTILINE))
        self.assertEqual(documented, set(EDCT_REQUIRED_COLUMNS) | set(EDCT_PN_REQUIRED_COLUMNS))

    def test_rule_documentation_matches_formula_and_unchecked_configuration(self) -> None:
        documentation = (
            Path(__file__).resolve().parents[2] / "docs" / "EDCT_QUALITY_CHECKER.md"
        ).read_text(encoding="utf-8")
        formula_section = documentation.split("## Formula rules", 1)[1].split("## Field rules", 1)[
            0
        ]
        unchecked_section = documentation.split("## Explicitly unchecked fields", 1)[1].split(
            "## Verified sample smoke", 1
        )[0]
        self.assertEqual(
            set(re.findall(r"^- `([^`]+)`$", formula_section, re.MULTILINE)),
            set(EDCT_FORMULA_COLUMNS),
        )
        self.assertEqual(
            set(re.findall(r"^- `([^`]+)`$", unchecked_section, re.MULTILINE)),
            set(EDCT_UNCHECKED_COLUMNS),
        )
        for configured_value in (
            *EDCT_TRIPLE_STATUS_VALUES,
            *EDCT_PORTAL_VALUES,
            *EDCT_EDI_MODE_VALUES,
        ):
            self.assertIn(f"`{configured_value}`", documentation)
        field_rules = documentation.split("## Field rules", 1)[1].split(
            "## Explicitly unchecked fields", 1
        )[0]
        documented_rows = tuple(line for line in field_rules.splitlines() if line.startswith("| `"))

        def cells(line: str) -> tuple[str, ...]:
            return tuple(cell.strip() for cell in line.strip("|").split("|"))

        self.assertEqual(
            tuple(map(cells, documented_rows)),
            tuple(map(cells, EDCT_RULE_CATALOGUE_ROWS)),
        )

    def test_complete_workbook_is_annotated_and_preserved(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            with closing(sqlite3.connect(":memory:")) as connection:
                connection.row_factory = sqlite3.Row
                result = run_edct_analysis(input_path, connection=connection)
                output_path = export_edct_result(result, directory)
                runs = connection.execute("SELECT project, rows_total FROM runs").fetchall()

            exported = load_workbook(output_path, data_only=False)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            self.assertEqual(result.assessed_rows, (("Supplier Level", 3), ("Supplier Level", 4)))
            self.assertEqual([row["project"] for row in runs], ["eDCT"])
            self.assertEqual([row["rows_total"] for row in runs], [2])
            self.assertEqual(
                exported.sheetnames,
                [
                    "Supplier Level",
                    "Open Task",
                    "Template-Cofor-Creation",
                    "Other Sheet",
                    "PN Level",
                ],
            )
            self.assertEqual(exported["Other Sheet"]["A1"].value, "keep me")
            self.assertEqual(supplier["A3"].fill.fgColor.rgb, "00FF0000")
            self.assertIn("Check", headers)
            self.assertIn("Comment", headers)
            self.assertEqual(supplier.cell(3, headers.index("Check") + 1).value, 0)
            self.assertTrue(
                supplier.tables["Tabella2"].ref.endswith(get_column_letter(len(headers)) + "5")
            )
            self.assertRegex(output_path.name, r"input_eDCT_checked_\d{8}_\d{6}\.xlsx")
            original = load_workbook(input_path)
            self.assertEqual(original["Other Sheet"]["A1"].value, "keep me")
            original.close()
            exported.close()

    def test_export_copies_unannotated_sheets_from_source(self) -> None:
        class UnserializableValue:
            def __str__(self) -> str:
                raise RuntimeError("unannotated sheet was serialized")

        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(
                    input_path,
                    connection=connection,
                    settings=edct_analysis.EdctHeaderSettings.default(),
                )
                self.addCleanup(result.workbook.close)
                unrelated_cell = result.workbook["Other Sheet"]["A1"]
                unrelated_cell._value = UnserializableValue()
                unrelated_cell.data_type = "s"
                output_path = export_edct_result(result, directory)
                self.assertIn("Other Sheet", result.workbook.sheetnames)
                self.assertIsInstance(
                    result.workbook["Other Sheet"]["A1"].value, UnserializableValue
                )

            exported = load_workbook(output_path, read_only=True)
            self.assertEqual(exported["Other Sheet"]["A1"].value, "keep me")
            exported.close()
            result.workbook.close()

    def test_optional_phone_and_shipping_columns_are_ignored_and_preserved(self) -> None:
        optional_columns = ("Phone", "Phone2", "Shipping location")
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(
                directory,
                optional_columns=optional_columns,
                row_overrides={
                    3: {
                        "Phone": "not a phone number",
                        "Phone2": 123,
                        "Shipping location": "arbitrary value",
                    }
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(input_path, connection=connection)
                output_path = export_edct_result(result, directory)

            self.assertEqual(result.row_results[("Supplier Level", 3)].check, 0)
            self.assertFalse(any(rule == "phone" for rule, _ in result.rule_totals))
            self.assertFalse(any(rule == "shipping_location" for rule, _ in result.rule_totals))
            exported = load_workbook(output_path)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            for column, expected in (
                ("Phone", "not a phone number"),
                ("Phone2", 123),
                ("Shipping location", "arbitrary value"),
            ):
                self.assertEqual(supplier.cell(3, headers.index(column) + 1).value, expected)
            exported.close()

    def test_export_formats_native_dates_as_day_month_year(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(
                directory,
                row_overrides={3: {"First communication sent": date(2026, 9, 2)}},
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(input_path, connection=connection)
                output_path = export_edct_result(result, directory)

            exported = load_workbook(output_path)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            date_cell = supplier.cell(3, headers.index("First communication sent") + 1)
            self.assertEqual(date_cell.value, datetime(2026, 9, 2))
            self.assertEqual(date_cell.number_format, "DD/MM/YYYY")
            exported.close()

    def test_export_includes_styles_created_during_annotation(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(input_path, connection=connection)
                worksheet = result.workbook["Supplier Level"]
                style_source = worksheet.max_column
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row, style_source).fill = PatternFill(
                        "solid",
                        fgColor="00FF00",
                    )
                output_path = export_edct_result(result, directory)

            exported = load_workbook(output_path)
            exported.close()

    def test_export_preserves_source_table_relationship_ids(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            remap_table_relationship(input_path, "rId99")

            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(input_path, connection=connection)
                output_path = export_edct_result(result, directory)

            exported = load_workbook(output_path)
            self.assertIn("Tabella2", exported["Supplier Level"].tables)
            exported.close()

    def test_rerun_clears_results_for_rows_that_are_no_longer_assessed(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(
                directory,
                row_overrides={3: {"Phone": "invalid"}},
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                first_result = run_edct_analysis(input_path, connection=connection)
                first_output = export_edct_result(first_result, directory)

                workbook = load_workbook(first_output)
                supplier = workbook["Supplier Level"]
                headers = [cell.value for cell in supplier[2]]
                supplier.cell(3, headers.index("Index") + 1).value = None
                rerun_input = directory / "rerun.xlsx"
                workbook.save(rerun_input)
                workbook.close()

                second_result = run_edct_analysis(rerun_input, connection=connection)
                second_output = export_edct_result(second_result, directory)

            exported = load_workbook(second_output)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            self.assertIsNone(supplier.cell(3, headers.index("Check") + 1).value)
            self.assertIsNone(supplier.cell(3, headers.index("Comment") + 1).value)
            exported.close()

    def test_existing_result_columns_are_brought_inside_the_result_table(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            workbook = load_workbook(input_path)
            supplier = workbook["Supplier Level"]
            supplier.cell(2, supplier.max_column + 1, "Check")
            supplier.cell(2, supplier.max_column + 1, "Comment")
            workbook.save(input_path)
            workbook.close()

            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(input_path, connection=connection)
                output = export_edct_result(result, directory)

            exported = load_workbook(output)
            table = exported["Supplier Level"].tables["Tabella2"]
            min_col, _, max_col, _ = range_boundaries(table.ref)
            self.assertEqual(len(table.tableColumns), max_col - min_col + 1)
            self.assertEqual(
                [column.name for column in table.tableColumns[-2:]],
                ["Check", "Comment"],
            )
            exported.close()

    def test_export_preserves_unsupported_worksheet_extensions(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            marker = "{EDCT-PRESERVATION-TEST}"
            add_worksheet_extension(input_path, marker)

            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(input_path, connection=connection)
                output = export_edct_result(result, directory)

            with ZipFile(input_path) as source, ZipFile(output) as exported:
                self.assertEqual(set(exported.namelist()), set(source.namelist()))
                for name in source.namelist():
                    if name not in {
                        "xl/worksheets/sheet1.xml",
                        "xl/worksheets/sheet5.xml",
                        "xl/tables/table1.xml",
                        "xl/styles.xml",
                    }:
                        self.assertEqual(exported.read(name), source.read(name), name)
                source_xml = source.read("xl/worksheets/sheet1.xml")
                exported_xml = exported.read("xl/worksheets/sheet1.xml")

            def extension_payload(xml: bytes) -> tuple:
                root = ElementTree.fromstring(xml)
                extension = next(
                    element
                    for element in root.iter()
                    if element.tag.rsplit("}", 1)[-1] == "ext"
                    and element.attrib.get("uri") == marker
                )

                def signature(element) -> tuple:
                    return (
                        element.tag,
                        tuple(sorted(element.attrib.items())),
                        (element.text or "").strip(),
                        tuple(signature(child) for child in element),
                    )

                return signature(extension)

            self.assertEqual(extension_payload(exported_xml), extension_payload(source_xml))

    def test_missing_required_sheet_stops_analysis(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(Path(temp), include_open_task=False)
            workbook = load_workbook(path)
            supplier = workbook["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            supplier.cell(2, headers.index("EDI Mode") + 1).value = None
            workbook.save(path)
            workbook.close()
            connection = sqlite3.connect(":memory:")
            try:
                with self.assertRaisesRegex(EdctLoadError, "Open Task.*EDI Mode"):
                    run_edct_analysis(path, connection=connection)
            finally:
                connection.close()

    def test_export_failure_marks_the_history_run_failed(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(directory)
            blocked_output = directory / "not-a-directory"
            blocked_output.write_text("file", encoding="utf-8")
            connection = sqlite3.connect(":memory:")
            connection.row_factory = sqlite3.Row
            try:
                result = run_edct_analysis(path, connection=connection)
                with self.assertRaises(OSError):
                    export_edct_result(result, blocked_output)
                run = connection.execute(
                    "SELECT status, error_message FROM runs WHERE id = ?",
                    (result.run_id,),
                ).fetchone()
            finally:
                connection.close()

        self.assertEqual(run["status"], "failed")
        self.assertTrue(run["error_message"])

    def test_corrupt_analysis_workbook_is_removed_and_run_is_failed(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(directory)
            connection = sqlite3.connect(":memory:")
            connection.row_factory = sqlite3.Row
            try:
                result = run_edct_analysis(
                    path,
                    connection=connection,
                    settings=edct_analysis.EdctHeaderSettings.default(),
                )
                original_merge = edct_analysis._merge_annotated_parts

                def merge_then_corrupt(*args, **kwargs) -> None:
                    original_merge(*args, **kwargs)
                    corrupt_first_cell_style(args[2])

                with (
                    patch.object(edct_analysis, "_merge_annotated_parts", merge_then_corrupt),
                    self.assertRaisesRegex(
                        EdctLoadError,
                        "Analysis workbook validation failed",
                    ),
                ):
                    export_edct_result(result, directory)
                run = connection.execute(
                    "SELECT status, error_message FROM runs WHERE id = ?",
                    (result.run_id,),
                ).fetchone()
            finally:
                connection.close()

            self.assertEqual(list(directory.glob("*_eDCT_checked_*.xlsx")), [])
        self.assertEqual(run["status"], "failed")
        self.assertIn("Analysis workbook validation failed", run["error_message"])

    def test_field_formats_are_reported_with_exact_column_names(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    3: {
                        "Effective kick-off date": "30/07/2026",
                        "eSupplierConnect": "NOT",
                        "B2B": "NOT",
                        "New supplier portal": "NOT",
                        "SPM": "NOT",
                        "iTMS": "NOT",
                        "OPEN TASK": "YES",
                        "Sales contact": "first@example.com, second@example.com",
                        "Seller COFOR": "ABC",
                        "First communication sent": "30/07/2026",
                        "Readiness Comments": "30/07/2026: contacted",
                        "Participants": "Name - person@example.com",
                    }
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(
                    path, connection=connection, analysis_date=date(2026, 7, 30)
                )

        self.assertEqual(result.row_results[("Supplier Level", 3)].check, 3)
        for column in (
            "Sales contact",
            "Seller COFOR",
            "Participants",
        ):
            self.assertIn(column, result.row_results[("Supplier Level", 3)].comment)
        self.assertIn(
            "first@example.com, second@example.com",
            result.row_results[("Supplier Level", 3)].comment,
        )
        self.assertIn(
            "Name - person@example.com", result.row_results[("Supplier Level", 3)].comment
        )

    def test_creation_of_cofors_request_date_accepts_only_supported_dates(self) -> None:
        scenarios = (
            ("empty", "", "unmatched", 0),
            ("native date", date(2026, 8, 10), "9999", 0),
            ("strict text", "10/08/2026", "9999", 0),
            ("future date", "10/08/2030", "9999", 0),
            ("dot date", "10.08.2026", "9999", 1),
            ("impossible date", "31/02/2026", "9999", 1),
            ("timestamp text", "10/08/2026 12:00", "9999", 1),
        )
        for label, request_date, supplier_punch, expected_check in scenarios:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(
                    Path(temp),
                    row_overrides={
                        3: {
                            "Supplier Punch code": supplier_punch,
                            "OPEN TASK": "",
                            "Creation of Cofors request date": request_date,
                        }
                    },
                )
                with closing(sqlite3.connect(":memory:")) as connection:
                    result = run_edct_analysis(path, connection=connection)

                self.assertEqual(result.row_results[("Supplier Level", 3)].check, expected_check)

    def test_all_input_sheets_resolve_columns_by_header_name(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    3: {"Triplet COFOR": "9999"},
                    4: {"Triplet COFOR": "8888"},
                },
                pn_rows=(("1003", "9999"),),
            )
            workbook = load_workbook(path)

            supplier = workbook["Supplier Level"]
            supplier_headers = [cell.value for cell in supplier[2]]
            punch_column = supplier_headers.index("Supplier Punch code") + 1
            name_column = supplier_headers.index("Supplier name") + 1
            for row in range(1, supplier.max_row + 1):
                punch_value = supplier.cell(row, punch_column).value
                supplier.cell(row, punch_column).value = supplier.cell(row, name_column).value
                supplier.cell(row, name_column).value = punch_value

            open_task = workbook["Open Task"]
            open_task["G2"] = open_task["A2"].value
            open_task["G3"] = open_task["A3"].value
            open_task["A2"] = None
            open_task["A3"] = None

            template = workbook["Template-Cofor-Creation"]
            template["H2"] = template["D2"].value
            template["H3"] = template["D3"].value
            template["D2"] = None
            template["D3"] = None

            pn = workbook["PN Level"]
            pn["F1"] = pn["A1"].value
            pn["F2"] = pn["A2"].value
            pn["H1"] = pn["B1"].value
            pn["H2"] = pn["B2"].value
            pn["A1"] = pn["A2"] = None
            pn["B1"] = pn["B2"] = None

            workbook.save(path)
            workbook.close()

            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)

            self.assertIn(("Supplier Level", 3), result.row_results)
            self.assertIn(("PN Level", 2), result.row_results)
            result.workbook.close()

    def test_cofor_template_uses_second_row_headers_without_fixed_column(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    3: {
                        "Supplier Punch code": "1003",
                        "Creation of Cofors request date": "10/08/2026",
                    }
                },
            )
            workbook = load_workbook(path)
            template = workbook["Template-Cofor-Creation"]
            template["D1"] = None
            template["D2"] = None
            template["D3"] = None
            template["E2"] = "Punch Code"
            template["E3"] = "1003"
            workbook.save(path)
            workbook.close()

            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)

            self.assertEqual(result.row_results[("Supplier Level", 3)].check, 0)
            result.workbook.close()

    def test_request_date_requires_supplier_punch_in_cofor_template(self) -> None:
        scenarios = (
            ("matching punch", " 1003 ", 0),
            ("absent punch", "9999", 1),
        )
        for label, template_punch, expected_check in scenarios:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(
                    Path(temp),
                    row_overrides={
                        3: {
                            "Supplier Punch code": "1003",
                            "Creation of Cofors request date": "10/08/2026",
                        }
                    },
                )
                workbook = load_workbook(path)
                workbook["Template-Cofor-Creation"]["D3"] = template_punch
                workbook.save(path)
                workbook.close()

                with closing(sqlite3.connect(":memory:")) as connection:
                    result = run_edct_analysis(path, connection=connection)

                self.assertEqual(result.row_results[("Supplier Level", 3)].check, expected_check)
                if expected_check:
                    self.assertEqual(
                        result.rule_totals[("cofor_template", "Supplier Punch code")],
                        1,
                    )

    def test_matching_cofor_template_punch_requires_request_date(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(Path(temp))
            workbook = load_workbook(path)
            workbook["Template-Cofor-Creation"]["D3"] = " 1003 "
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)

            self.assertEqual(result.row_results[("Supplier Level", 3)].check, 1)
            self.assertEqual(
                result.rule_totals[("date_required", "Creation of Cofors request date")],
                1,
            )
            self.assertEqual(result.row_results[("Supplier Level", 4)].check, 0)

    def test_cofor_template_matching_is_normalized_but_not_fuzzy(self) -> None:
        scenarios = (
            ("case and outer whitespace", " Ab- 01 ", "aB- 01", 1),
            ("partial", "ABC123", "ABC", 0),
            ("leading zero", "00123", "123", 0),
            ("internal space", "AB 123", "AB123", 0),
            ("punctuation", "AB-123", "AB123", 0),
            ("numeric representation", "123.0", "123", 0),
        )
        for label, reference_punch, supplier_punch, expected_check in scenarios:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(
                    Path(temp),
                    row_overrides={
                        3: {
                            "Supplier Punch code": supplier_punch,
                            "OPEN TASK": "",
                        }
                    },
                )
                workbook = load_workbook(path)
                workbook["Template-Cofor-Creation"]["D3"] = reference_punch
                workbook.save(path)
                workbook.close()

                with closing(sqlite3.connect(":memory:")) as connection:
                    result = run_edct_analysis(path, connection=connection)

                self.assertEqual(result.row_results[("Supplier Level", 3)].check, expected_check)

    def test_missing_or_malformed_cofor_template_stops_analysis(self) -> None:
        scenarios = (
            ("missing sheet", False, None, "Template-Cofor-Creation"),
            (
                "wrong row 2 header",
                True,
                "Wrong header",
                r"Template-Cofor-Creation\.Punch Code",
            ),
        )
        for label, include_sheet, header, expected_error in scenarios:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(
                    Path(temp),
                    include_cofor_template=include_sheet,
                )
                if header is not None:
                    workbook = load_workbook(path)
                    workbook["Template-Cofor-Creation"]["D2"] = header
                    workbook.save(path)
                    workbook.close()

                with (
                    closing(sqlite3.connect(":memory:")) as connection,
                    self.assertRaisesRegex(EdctLoadError, expected_error),
                ):
                    run_edct_analysis(path, connection=connection)

    def test_email_and_dated_comment_boundaries(self) -> None:
        scenarios = (
            ("single email", {"Participants": "one@example.com"}, 0),
            ("email list", {"Participants": "one@example.com; two@example.org"}, 0),
            ("trailing separators", {"Participants": "one@example.com;;;"}, 0),
            ("spaced separators", {"Participants": "one@example.com ; two@example.org ;"}, 0),
            ("duplicate emails", {"Participants": "one@example.com;one@example.com;"}, 0),
            ("comma-separated email", {"Participants": "one@example.com,two@example.org"}, 1),
            ("display-name email", {"Participants": "Name <one@example.com>"}, 1),
            ("internal empty email", {"Participants": "one@example.com;;two@example.org"}, 1),
            ("separator-only email", {"Participants": ";;"}, 1),
            ("slash general comment", {"Comments": "30/07/2026: ready"}, 0),
            ("dot general comment", {"Comments": "30.07.2026: ready"}, 1),
            ("dot readiness comment", {"Readiness Comments": "30.07.2026: ready"}, 1),
            ("slash readiness comment", {"Readiness Comments": "30/07/2026: ready"}, 0),
        )
        for label, overrides, expected_check in scenarios:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(Path(temp), row_overrides={3: overrides})
                with closing(sqlite3.connect(":memory:")) as connection:
                    result = run_edct_analysis(
                        path,
                        connection=connection,
                        analysis_date=date(2026, 7, 30),
                    )
                self.assertEqual(result.row_results[("Supplier Level", 3)].check, expected_check)

    def test_catalogued_applicability_empty_and_accepted_value_policies(self) -> None:
        valid_portals = {
            "eSupplierConnect": "YES",
            "B2B": "NOT",
            "New supplier portal": "YES",
            "SPM": "NOT",
            "iTMS": "YES",
        }
        scenarios = (
            ("conditional email before trigger", {"Sales contact": "invalid"}, 0),
            (
                "conditional email after trigger",
                {
                    "Effective kick-off date": "30/07/2026",
                    "Sales contact": "invalid",
                    **valid_portals,
                },
                1,
            ),
            ("conditional cofor before trigger", {"Seller COFOR": "invalid"}, 0),
            (
                "conditional cofor after trigger",
                {
                    "Effective kick-off date": "30/07/2026",
                    "Seller COFOR": "invalid",
                    **valid_portals,
                },
                1,
            ),
            (
                "future effective date",
                {
                    "Effective kick-off date": "31/07/2026",
                    **valid_portals,
                },
                1,
            ),
            ("triple status before trigger", {"Triple Status": ""}, 0),
            (
                "triple status after trigger",
                {
                    "Cofor created date": "30/07/2026",
                    "Triple Status": "",
                    "EDI Mode": "WEB EDI",
                },
                1,
            ),
            ("overseas empty", {"Overseas": ""}, 0),
            ("overseas exact YES", {"Overseas": " YES "}, 0),
            ("overseas exact NOT", {"Overseas": "NOT"}, 0),
            ("overseas lowercase", {"Overseas": "yes"}, 1),
            ("overseas NO", {"Overseas": "NO"}, 1),
            ("supplier confirmation optional", {"Supplier Confimation": ""}, 0),
            ("supplier confirmation invalid", {"Supplier Confimation": "NO"}, 1),
            (
                "portals required after trigger",
                {"Effective kick-off date": "30/07/2026"},
                len(EDCT_PORTAL_COLUMNS),
            ),
            ("edi mode before trigger", {"EDI Mode": ""}, 0),
            (
                "edi mode after trigger",
                {
                    "Cofor created date": "30/07/2026",
                    "Triple Status": "Valid",
                    "EDI Mode": "",
                },
                1,
            ),
        )
        for label, overrides, expected_check in scenarios:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(Path(temp), row_overrides={3: overrides})
                with closing(sqlite3.connect(":memory:")) as connection:
                    result = run_edct_analysis(
                        path,
                        connection=connection,
                        analysis_date=date(2026, 7, 30),
                    )
                self.assertEqual(result.row_results[("Supplier Level", 3)].check, expected_check)

    def test_every_configured_field_rule_reaches_exported_check_and_comment(self) -> None:
        invalid_values = {
            **{column: "invalid email" for column in EDCT_EMAIL_COLUMNS},
            **{column: "invalid cofor" for column in EDCT_COFOR_COLUMNS},
            **{column: "invalid date" for column in EDCT_DATE_COLUMNS},
            **{column: "invalid comment" for column in EDCT_DATED_COMMENT_COLUMNS},
            "Triple Status": "Valid",
            "Overseas": "YES",
            "Supplier Confimation": "YES",
            "eSupplierConnect": "YES",
            "B2B": "YES",
            "New supplier portal": "YES",
            "SPM": "YES",
            "iTMS": "YES",
            "EDI Mode": "WEB EDI",
            "OPEN TASK": "YES",
        }
        expected_columns = (
            *EDCT_EMAIL_COLUMNS,
            *EDCT_COFOR_COLUMNS,
            *EDCT_DATE_COLUMNS,
            *EDCT_DATED_COMMENT_COLUMNS,
            "Supplier Punch code",
        )
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(directory, row_overrides={3: invalid_values})
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(
                    path,
                    connection=connection,
                    analysis_date=date(2026, 7, 30),
                )
                output = export_edct_result(result, directory)

            exported = load_workbook(output)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            check = supplier.cell(3, headers.index("Check") + 1).value
            comment = supplier.cell(3, headers.index("Comment") + 1).value
            exported.close()

        self.assertEqual(check, len(expected_columns))
        for column in expected_columns:
            self.assertIn(f"{column} =", comment)

    def test_lifecycle_and_open_task_conditions_are_reported(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(
                directory,
                row_overrides={
                    3: {
                        "Cofor created date": "invalid date",
                        "Triple Status": "",
                        "EDI Mode": "",
                        "Effective kick-off date": "30/07/2026",
                        "Overseas": "YES",
                        "Supplier Confimation": "NO",
                        "OPEN TASK": "",
                    },
                    4: {"OPEN TASK": "NO"},
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                connection.row_factory = sqlite3.Row
                result = run_edct_analysis(
                    path, connection=connection, analysis_date=date(2026, 7, 30)
                )
                totals = connection.execute(
                    "SELECT column_name, fail_count FROM run_columns WHERE run_id = ?",
                    (result.run_id,),
                ).fetchall()
                output = export_edct_result(result, directory)

            exported = load_workbook(output)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            exported_checks = [
                supplier.cell(row, headers.index("Check") + 1).value for row in (3, 4)
            ]
            exported_comments = [
                supplier.cell(row, headers.index("Comment") + 1).value for row in (3, 4)
            ]
            exported.close()

        self.assertEqual(result.row_results[("Supplier Level", 3)].check, 10)
        self.assertEqual(result.row_results[("Supplier Level", 4)].check, 1)
        self.assertEqual(exported_checks, [10, 1])
        for column in (
            "Triple Status",
            "EDI Mode",
            "Supplier Confimation",
            "OPEN TASK",
            "eSupplierConnect",
            "B2B",
            "New supplier portal",
            "SPM",
            "iTMS",
        ):
            self.assertIn(column, result.row_results[("Supplier Level", 3)].comment)
            self.assertIn(column, exported_comments[0])
        self.assertIn("OPEN TASK", exported_comments[1])
        self.assertEqual(sum(row["fail_count"] for row in totals), 11)

    def test_formula_structure_uses_formula_reference_row(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    4: {
                        "Starting date": '=IF(B4="","",B4)',
                        "Triplet COFOR": '= + if ( a4 = "" , "" , a4 )',
                    }
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(
                    path, connection=connection, analysis_date=date(2026, 7, 30)
                )

        self.assertEqual(result.row_results[("Supplier Level", 3)].check, 0)
        self.assertEqual(result.row_results[("Supplier Level", 4)].check, 1)
        self.assertIn("Starting date", result.row_results[("Supplier Level", 4)].comment)
        self.assertNotIn("Triplet COFOR", result.row_results[("Supplier Level", 4)].comment)

    def test_formula_comparison_rejects_every_material_change(self) -> None:
        scenarios = (
            ("constant", "manual"),
            ("missing formula", None),
            ("changed function", '=AND(A4="")'),
            ("changed operator", '=IF(A4<>"","",A4)'),
            ("changed reference", '=IF(B4="","",B4)'),
            ("changed condition", '=IF(A4="x","",A4)'),
            ("changed quoted value", '=IF(A4="","changed",A4)'),
        )
        for label, formula in scenarios:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(
                    Path(temp),
                    row_overrides={4: {"Starting date": formula}},
                )
                with closing(sqlite3.connect(":memory:")) as connection:
                    result = run_edct_analysis(
                        path,
                        connection=connection,
                        analysis_date=date(2026, 7, 30),
                    )
                self.assertEqual(result.row_results[("Supplier Level", 4)].check, 1)
                self.assertIn("Starting date", result.row_results[("Supplier Level", 4)].comment)

    def test_every_formula_rule_reaches_exported_check_and_comment(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(
                directory,
                row_overrides={
                    4: {column: "=1" for column in EDCT_FORMULA_COLUMNS},
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)
                output = export_edct_result(result, directory)

            exported = load_workbook(output)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            check = supplier.cell(4, headers.index("Check") + 1).value
            comment = supplier.cell(4, headers.index("Comment") + 1).value
            exported.close()

        self.assertEqual(check, len(EDCT_FORMULA_COLUMNS))
        for column in EDCT_FORMULA_COLUMNS:
            self.assertIn(f"{column} =", comment)

    def test_missing_formula_reference_flags_every_assessed_row_and_continues(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={3: {"Starting date": None}},
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(
                    path, connection=connection, analysis_date=date(2026, 7, 30)
                )

        self.assertEqual(result.row_results[("Supplier Level", 3)].check, 1)
        self.assertEqual(result.row_results[("Supplier Level", 4)].check, 1)

    def test_run_analysis_with_custom_header_settings(self) -> None:
        from quality_checker.checkers.edct.settings import EdctHeaderSettings

        settings = EdctHeaderSettings.default()
        settings.supplier_level["Supplier Punch code"] = "Code Fournisseur"
        settings.supplier_level["Supplier name"] = "Nom Fournisseur"
        settings.supplier_level["Sales contact"] = "Contact Commercial"
        settings.open_task["Punch Code"] = "Code Tache"
        settings.pn_level["Punch seller"] = "Vendeur"

        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            # Create workbook with renamed headers
            workbook = Workbook()
            all_columns = [
                settings.get_header("Supplier Level", col) for col in EDCT_REQUIRED_COLUMNS
            ]
            supplier = workbook.active
            supplier.title = "Supplier Level"
            supplier.append(["metadata"] * len(all_columns))
            supplier.append(all_columns)

            base: dict[str, object] = {column: "" for column in all_columns}
            for row_number, index in ((3, "Metz_01"), (4, "Metz_02")):
                values = dict(base)
                values.update(
                    {
                        "Index": index,
                        "Code Fournisseur": 1000 + row_number,
                        "Nom Fournisseur": f"Supplier {row_number}",
                        "Effective kick-off date": "01/01/2026",
                        "Contact Commercial": "invalid-email-format",
                        "Cofor created date": "",
                        "Overseas": "",
                        "OPEN TASK": "YES" if row_number == 3 else "",
                    }
                )
                for col in EDCT_FORMULA_COLUMNS:
                    values[col] = f'=IF(A{row_number}="","",A{row_number})'
                supplier.append([values[col] for col in all_columns])

            table = Table(
                displayName="Tabella2",
                ref=f"A2:{get_column_letter(len(all_columns))}4",
            )
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            supplier.add_table(table)

            open_task = workbook.create_sheet("Open Task")
            open_task.append(["metadata"])
            open_task.append(["Code Tache"])
            open_task.append([1003])

            cofor_template = workbook.create_sheet("Template-Cofor-Creation")
            cofor_template["D2"] = "Punch Code"
            cofor_template["D3"] = "9999"

            pn = workbook.create_sheet("PN Level")
            pn.append(["Vendeur", "Triplet COFOR"])
            pn.append(["1003", "123456  78"])

            path = directory / "custom_edct.xlsx"
            workbook.save(path)
            set_formula_caches(
                path,
                {"xl/worksheets/sheet1.xml": {"G3": "123456  78", "G4": "123456  78"}},
            )

            with closing(sqlite3.connect(":memory:")) as connection:
                # Running without settings should fail with missing columns
                with self.assertRaises(EdctLoadError) as err:
                    run_edct_analysis(
                        path, connection=connection, settings=EdctHeaderSettings.default()
                    )
                self.assertIn("Supplier Punch code", str(err.exception))

                # Running with custom settings should succeed
                result = run_edct_analysis(
                    path,
                    connection=connection,
                    analysis_date=date(2026, 7, 30),
                    settings=settings,
                )

            self.assertEqual(len(result.assessed_rows), 3)
            row_3 = result.row_results[("Supplier Level", 3)]
            self.assertGreater(row_3.check, 0)
            self.assertIn("Contact Commercial = invalid-email-format", row_3.comment)
            self.assertEqual(result.rule_totals[("email", "Contact Commercial")], 2)
            pn_row = result.row_results[("PN Level", 2)]
            self.assertEqual(pn_row.check, 0)


if __name__ == "__main__":
    unittest.main()
