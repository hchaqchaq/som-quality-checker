from __future__ import annotations

from collections import Counter
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
import re
import sqlite3
from time import perf_counter

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import TableColumn

from ..config import DB_PATH
from ..db.repository import (
    ColumnRecord,
    RunRecord,
    initialize_schema,
    insert_run,
    open_connection,
    update_run_exported_file,
    update_run_status,
)
from ..edct_config import (
    EDCT_COFOR_COLUMNS,
    EDCT_DATED_COMMENT_COLUMNS,
    EDCT_DATE_COLUMNS,
    EDCT_EMAIL_COLUMNS,
    EDCT_FORMULA_COLUMNS,
    EDCT_HEADER_ROW,
    EDCT_INDEX_COLUMN,
    EDCT_PHONE_COLUMNS,
    EDCT_PORTAL_COLUMNS,
    EDCT_PROJECT,
    EDCT_REQUIRED_COLUMNS,
    EDCT_REQUIRED_SHEETS,
)

EMAIL_REGEX = re.compile(r"^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$")
COFOR_REGEX = re.compile(r"^[A-Za-z0-9]{6} {2}[A-Za-z0-9]{2}$")


class EdctLoadError(Exception):
    """Raised when an eDCT workbook does not contain the required structure."""


@dataclass(frozen=True, slots=True)
class EdctRowResult:
    check: int
    comment: str


@dataclass(slots=True)
class EdctRunResult:
    run_id: int
    input_file: Path
    started_at: datetime
    finished_at: datetime
    duration_s: float
    workbook: object
    processed_rows: tuple[int, ...]
    row_results: dict[int, EdctRowResult]
    rule_totals: Counter[tuple[str, str]]
    uses_default_db: bool
    history_connection: sqlite3.Connection | None

    @property
    def rows_failed(self) -> int:
        return sum(result.check > 0 for result in self.row_results.values())


def _normalized_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _normalized_choice(value: object) -> str:
    return _normalized_text(value).casefold()


def _is_valid_email(value: object) -> bool:
    text = _normalized_text(value)
    if not text:
        return True
    parts = [part.strip() for part in text.split(";")]
    return bool(parts) and all(part and EMAIL_REGEX.fullmatch(part) for part in parts)


def _is_valid_phone(value: object) -> bool:
    text = _normalized_text(value)
    if not text:
        return True
    digits = re.sub(r"[ +().-]", "", text)
    return digits.isdigit() and 7 <= len(digits) <= 20


def _parse_date(value: object, *, allow_slash: bool = False) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _normalized_text(value)
    if not text:
        return None
    formats = ("%d.%m.%Y", "%d/%m/%Y") if allow_slash else ("%d.%m.%Y",)
    for date_format in formats:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    return None


def _is_valid_dated_comment(value: object, *, allow_slash: bool) -> bool:
    text = _normalized_text(value)
    if not text:
        return True
    date_text, separator, comment = text.partition(":")
    return bool(separator and comment.strip() and _parse_date(date_text.strip(), allow_slash=allow_slash))


def _normalized_punch(value: object) -> str:
    text = _normalized_text(value)
    return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


def _formula_key(formula: object) -> str | None:
    if not isinstance(formula, str) or not formula.startswith("="):
        return None
    normalized: list[str] = []
    quoted = False
    for character in formula:
        if character == '"':
            quoted = not quoted
            normalized.append(character)
        elif character.isspace() and not quoted:
            continue
        else:
            normalized.append(character if quoted else character.casefold())
    result = "".join(normalized)
    return "=" + result[2:] if result.startswith("=+") else result


def _open_task_punches(workbook) -> set[str]:
    worksheet = workbook["Open Task"]
    headers = _header_map(worksheet)
    column = headers.get("Punch Code")
    if column is None:
        raise EdctLoadError("Missing required columns in Open Task: Punch Code")
    return {
        punch
        for row in range(EDCT_HEADER_ROW + 1, worksheet.max_row + 1)
        if (punch := _normalized_punch(worksheet.cell(row, column).value))
    }


def _evaluate_business_rules(
    workbook,
    headers: dict[str, int],
    processed_rows: tuple[int, ...],
    analysis_date: date,
) -> tuple[dict[int, EdctRowResult], Counter[tuple[str, str]]]:
    worksheet = workbook["Supplier Level"]
    failures: dict[int, list[tuple[str, str]]] = {row: [] for row in processed_rows}
    totals: Counter[tuple[str, str]] = Counter()

    def value(row: int, column: str) -> object:
        return worksheet.cell(row, headers[column]).value

    def fail(row: int, rule: str, column: str, reason: str) -> None:
        failures[row].append((reason, column))
        totals[(rule, column)] += 1

    open_task_punches = _open_task_punches(workbook)
    if processed_rows:
        reference_row = processed_rows[0]
        for column in EDCT_FORMULA_COLUMNS:
            reference_cell = worksheet.cell(reference_row, headers[column])
            reference_formula = _formula_key(reference_cell.value)
            if reference_formula is None:
                reason = (
                    "Formula validation failed: first processed row is missing "
                    f"the reference formula for {column}"
                )
                for row in processed_rows:
                    fail(row, "formula", column, reason)
                continue
            for row in processed_rows[1:]:
                target_cell = worksheet.cell(row, headers[column])
                expected = Translator(
                    str(reference_cell.value),
                    origin=reference_cell.coordinate,
                ).translate_formula(target_cell.coordinate)
                if _formula_key(target_cell.value) != _formula_key(expected):
                    fail(row, "formula", column, "Formula does not match first processed row")

    for row in processed_rows:
        effective_date_value = value(row, "Effective kick-off date")
        effective_date = _parse_date(effective_date_value)
        cofor_date_value = value(row, "Cofor created date")
        cofor_date = _parse_date(cofor_date_value)
        effective_date_filled = bool(_normalized_text(effective_date_value))
        cofor_date_filled = bool(_normalized_text(cofor_date_value))

        for column in EDCT_EMAIL_COLUMNS:
            if column in ("Sales contact", "Logistic contact") and not effective_date_filled:
                continue
            if not _is_valid_email(value(row, column)):
                fail(row, "email", column, "Invalid email")

        if effective_date_filled:
            for column in EDCT_COFOR_COLUMNS:
                text = _normalized_text(value(row, column))
                if text and COFOR_REGEX.fullmatch(text) is None:
                    fail(row, "cofor", column, "Invalid COFOR format")

        for column in EDCT_PHONE_COLUMNS:
            if not _is_valid_phone(value(row, column)):
                fail(row, "phone", column, "Invalid phone number")

        for column in EDCT_DATE_COLUMNS:
            raw = value(row, column)
            if _normalized_text(raw) and _parse_date(raw) is None:
                fail(row, "date", column, "Invalid date")
        if effective_date is not None and effective_date > analysis_date:
            fail(row, "date_future", "Effective kick-off date", "Date cannot be in the future")

        for column in EDCT_DATED_COMMENT_COLUMNS:
            if not _is_valid_dated_comment(
                value(row, column),
                allow_slash=column in ("Comments", "Kick-off comments"),
            ):
                fail(row, "dated_comment", column, "Invalid dated comment")

        triple_status = _normalized_choice(value(row, "Triple Status"))
        if cofor_date_filled and triple_status not in {"valid", "no valid"}:
            fail(row, "triple_status", "Triple Status", "Required value must be Valid or No Valid")

        overseas = _normalized_choice(value(row, "Overseas"))
        if overseas not in {"", "yes", "no"}:
            fail(row, "overseas", "Overseas", "Invalid value")
        if overseas == "yes" and _normalized_choice(value(row, "Shipping location")) not in {"yes", "no"}:
            fail(row, "shipping_location", "Shipping location", "Required value must be Yes or No")

        supplier_confirmation = _normalized_choice(value(row, "Supplier Confimation"))
        if supplier_confirmation not in {"", "yes"}:
            fail(row, "supplier_confirmation", "Supplier Confimation", "Invalid value")

        for column in EDCT_PORTAL_COLUMNS:
            portal_value = _normalized_choice(value(row, column))
            allowed = {"yes", "not"} if effective_date_filled else {"", "yes", "not"}
            if portal_value not in allowed:
                fail(row, "portal", column, "Required value must be YES or NOT")

        edi_mode = _normalized_choice(value(row, "EDI Mode"))
        allowed_edi_modes = {"web edi", "standard edi"} if cofor_date_filled else {
            "",
            "web edi",
            "standard edi",
        }
        if edi_mode not in allowed_edi_modes:
            fail(row, "edi_mode", "EDI Mode", "Required value must be WEB EDI or Standard EDI")

        punch_in_open_task = _normalized_punch(value(row, "Supplier Punch code")) in open_task_punches
        open_task_value = _normalized_choice(value(row, "OPEN TASK"))
        if punch_in_open_task and open_task_value != "yes":
            fail(row, "open_task", "OPEN TASK", "OPEN TASK must be YES for a matching punch code")
        elif not punch_in_open_task and open_task_value:
            fail(row, "open_task", "OPEN TASK", "OPEN TASK must be empty when the punch code is absent")

    row_results: dict[int, EdctRowResult] = {}
    for row, row_failures in failures.items():
        if not row_failures:
            row_results[row] = EdctRowResult(0, "Quality check passed")
            continue
        grouped: dict[str, list[str]] = {}
        for reason, column in row_failures:
            grouped.setdefault(reason, []).append(column)
        comment = " | ".join(f"{reason}: {', '.join(columns)}" for reason, columns in grouped.items())
        row_results[row] = EdctRowResult(len(row_failures), comment)
    return row_results, totals


def _header_map(worksheet, row: int = EDCT_HEADER_ROW) -> dict[str, int]:
    return {
        _normalized_text(cell.value): cell.column
        for cell in worksheet[row]
        if _normalized_text(cell.value)
    }


def _open_history(connection: sqlite3.Connection | None) -> tuple[sqlite3.Connection, bool]:
    if connection is not None:
        initialize_schema(connection)
        return connection, False
    own_connection = open_connection(DB_PATH)
    initialize_schema(own_connection)
    return own_connection, True


def _insert_history(
    connection: sqlite3.Connection,
    *,
    started_at: datetime,
    finished_at: datetime,
    duration_s: float,
    input_file: Path,
    rows_total: int,
    rows_failed: int,
    status: str,
    error_message: str | None,
    rule_totals: Counter[tuple[str, str]] | None = None,
) -> int:
    columns = [
        ColumnRecord(rule_name=rule, column_name=column, fail_count=count)
        for (rule, column), count in (rule_totals or {}).items()
    ]
    return insert_run(
        connection,
        RunRecord(
            project=EDCT_PROJECT,
            started_at=started_at.isoformat(),
            finished_at=finished_at.isoformat(),
            duration_s=duration_s,
            input_file=str(input_file),
            exported_file=None,
            rows_total=rows_total,
            rows_in_scope=rows_total,
            rows_failed=rows_failed,
            status=status,
            error_message=error_message,
        ),
        columns,
    )


def run_edct_analysis(
    input_path: Path | str,
    connection: sqlite3.Connection | None = None,
    *,
    analysis_date: date | None = None,
) -> EdctRunResult:
    resolved_input = Path(input_path)
    started_at = datetime.now(timezone.utc)
    started_perf = perf_counter()
    history, own_connection = _open_history(connection)

    try:
        if not resolved_input.exists():
            raise EdctLoadError(f"Input file not found: {resolved_input}")
        workbook = load_workbook(resolved_input, data_only=False)
        missing_sheets = [sheet for sheet in EDCT_REQUIRED_SHEETS if sheet not in workbook.sheetnames]
        missing_columns: list[str] = []
        headers: dict[str, int] = {}
        if "Supplier Level" in workbook.sheetnames:
            headers = _header_map(workbook["Supplier Level"])
            missing_columns.extend(column for column in EDCT_REQUIRED_COLUMNS if column not in headers)
        if "Open Task" in workbook.sheetnames and "Punch Code" not in _header_map(workbook["Open Task"]):
            missing_columns.append("Open Task.Punch Code")
        if missing_sheets or missing_columns:
            workbook.close()
            parts = []
            if missing_sheets:
                parts.append(f"sheets: {', '.join(missing_sheets)}")
            if missing_columns:
                parts.append(f"columns: {', '.join(missing_columns)}")
            raise EdctLoadError(f"Missing required structure: {'; '.join(parts)}")

        supplier = workbook["Supplier Level"]
        processed_rows = tuple(
            row
            for row in range(EDCT_HEADER_ROW + 1, supplier.max_row + 1)
            if _normalized_text(supplier.cell(row, headers[EDCT_INDEX_COLUMN]).value)
        )
        row_results, rule_totals = _evaluate_business_rules(
            workbook,
            headers,
            processed_rows,
            analysis_date or date.today(),
        )
        finished_at = datetime.now(timezone.utc)
        duration_s = perf_counter() - started_perf
        run_id = _insert_history(
            history,
            started_at=started_at,
            finished_at=finished_at,
            duration_s=duration_s,
            input_file=resolved_input,
            rows_total=len(processed_rows),
            rows_failed=sum(result.check > 0 for result in row_results.values()),
            status="ok",
            error_message=None,
            rule_totals=rule_totals,
        )
        return EdctRunResult(
            run_id=run_id,
            input_file=resolved_input,
            started_at=started_at,
            finished_at=finished_at,
            duration_s=duration_s,
            workbook=workbook,
            processed_rows=processed_rows,
            row_results=row_results,
            rule_totals=rule_totals,
            uses_default_db=own_connection,
            history_connection=None if own_connection else history,
        )
    except Exception as exc:
        finished_at = datetime.now(timezone.utc)
        _insert_history(
            history,
            started_at=started_at,
            finished_at=finished_at,
            duration_s=perf_counter() - started_perf,
            input_file=resolved_input,
            rows_total=0,
            rows_failed=0,
            status="failed",
            error_message=str(exc),
        )
        raise
    finally:
        if own_connection:
            history.close()


def _result_columns(worksheet) -> tuple[int, int]:
    headers = _header_map(worksheet)
    table = worksheet.tables.get("Tabella2")
    if table is None:
        raise EdctLoadError("Missing required table: Tabella2")

    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    result_columns: list[int] = []
    for name in ("Check", "Comment"):
        column = headers.get(name)
        if column is None:
            max_col += 1
            column = max_col
            worksheet.cell(EDCT_HEADER_ROW, column, name)
            table.tableColumns.append(TableColumn(id=len(table.tableColumns) + 1, name=name))
        else:
            max_col = max(max_col, column)
        result_columns.append(column)

    table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
    if table.autoFilter is not None:
        table.autoFilter.ref = table.ref
    return result_columns[0], result_columns[1]


def _update_result_history(
    result: EdctRunResult,
    action,
) -> None:
    if result.run_id < 0:
        return
    if result.history_connection is not None:
        action(result.history_connection)
        return
    if result.uses_default_db:
        connection = open_connection(DB_PATH)
        try:
            initialize_schema(connection)
            action(connection)
        finally:
            connection.close()


def export_edct_result(result: EdctRunResult, output_dir: Path | str) -> Path:
    try:
        return _export_edct_result(result, output_dir)
    except Exception as exc:
        _update_result_history(
            result,
            lambda connection: update_run_status(connection, result.run_id, "failed", str(exc)),
        )
        raise


def _export_edct_result(result: EdctRunResult, output_dir: Path | str) -> Path:
    target_dir = Path(output_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    worksheet = result.workbook["Supplier Level"]
    check_column, comment_column = _result_columns(worksheet)
    style_source = max(1, min(check_column, comment_column) - 1)

    for column in (check_column, comment_column):
        for row in range(EDCT_HEADER_ROW, worksheet.max_row + 1):
            source = worksheet.cell(row, style_source)
            target = worksheet.cell(row, column)
            target._style = copy(source._style)
            if source.has_style:
                target.number_format = source.number_format

    for row in result.processed_rows:
        row_result = result.row_results[row]
        worksheet.cell(row, check_column, row_result.check)
        worksheet.cell(row, comment_column, row_result.comment)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_stem = re.sub(r"[^A-Za-z0-9._-]+", "_", result.input_file.stem).strip("._") or "analysis"
    target = target_dir / f"{safe_stem}_eDCT_checked_{timestamp}.xlsx"
    result.workbook.save(target)

    _update_result_history(
        result,
        lambda connection: update_run_exported_file(connection, result.run_id, str(target)),
    )
    return target
