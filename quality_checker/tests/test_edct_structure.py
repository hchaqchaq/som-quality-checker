from __future__ import annotations

import sqlite3
import tempfile
from contextlib import closing
from datetime import date
from pathlib import Path

from edct_support import (
    EdctTestCase,
    build_edct_workbook,
    set_formula_caches,
)
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from quality_checker.checkers.edct.config import (
    EDCT_FORMULA_COLUMNS,
    EDCT_REQUIRED_COLUMNS,
)
from quality_checker.checkers.edct.export import export_edct_result
from quality_checker.checkers.edct.models import EdctLoadError
from quality_checker.checkers.edct.runner import run_edct_analysis
from quality_checker.checkers.edct.settings import EdctHeaderSettings
from quality_checker.checkers.edct.workbook import inspect_edct_workbook


def test_header_inspection_accepts_normalized_names_and_approved_aliases() -> None:
    with tempfile.TemporaryDirectory() as temp:
        path = build_edct_workbook(Path(temp))
        workbook = load_workbook(path)
        supplier = workbook["Supplier Level"]
        headers = [cell.value for cell in supplier[2]]
        supplier.cell(
            2, headers.index("Supplier Confimation") + 1
        ).value = "  supplier confirmation  "
        pn = workbook["PN Level"]
        pn["A1"] = "  PUNCH SELLER "
        workbook.save(path)
        workbook.close()

        stages: list[str] = []
        result = inspect_edct_workbook(
            path,
            EdctHeaderSettings.default(),
            progress=stages.append,
        )

        assert result.ready
        assert result.resolved_headers["Supplier Level"]["Supplier Confimation"] == (
            "supplier confirmation"
        )
        assert result.resolved_headers["PN Level"]["Punch seller"] == "PUNCH SELLER"
        assert ("Supplier Level", "Supplier Confimation") in result.alias_matches
        assert stages == ["Checking worksheet structure…", "Checking required headers…"]

        with closing(sqlite3.connect(":memory:")) as connection:
            analysis = run_edct_analysis(
                path,
                connection=connection,
                settings=result.resolved_settings(),
            )
        analysis.workbook.close()


class EdctStructureTests(EdctTestCase):
    def test_missing_path_and_combined_structure_errors_are_recorded(self) -> None:
        with closing(sqlite3.connect(":memory:")) as connection:
            with self.assertRaisesRegex(EdctLoadError, "Input file not found"):
                run_edct_analysis("missing.xlsx", connection=connection)

        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "broken.xlsx"
            workbook = Workbook()
            workbook.active.title = "Supplier Level"
            workbook.active.append(["metadata"])
            workbook.active.append(["Wrong header"])
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                with self.assertRaises(EdctLoadError) as raised:
                    run_edct_analysis(path, connection=connection)
                for missing in ("Open Task", "Template-Cofor-Creation", "PN Level", "Index"):
                    self.assertIn(missing, str(raised.exception))

    def test_line_header_can_define_assessed_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    3: {"Supplier Punch code": "P1", "Triplet COFOR": "A", "OPEN TASK": ""},
                    4: {"Supplier Punch code": "P1", "Triplet COFOR": "B"},
                },
                pn_rows=(("P1", "B"), ("P1", "C")),
            )
            workbook = load_workbook(path)
            supplier = workbook["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            supplier.cell(2, headers.index("Index") + 1).value = "Line"
            workbook.save(path)
            workbook.close()

            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)

            self.assertEqual(
                result.assessed_rows,
                (("Supplier Level", 3), ("Supplier Level", 4), ("PN Level", 2), ("PN Level", 3)),
            )
            self.assertEqual(result.row_results[("PN Level", 2)].check, 0)
            self.assertEqual(result.row_results[("PN Level", 3)].check, 1)
            result.workbook.close()

    def test_open_task_header_and_invalid_overseas_are_reported(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(directory, row_overrides={3: {"Overseas": "MAYBE"}})
            workbook = load_workbook(path)
            workbook["Open Task"]["A2"] = "Wrong"
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                with self.assertRaisesRegex(EdctLoadError, "Open Task.Punch Code"):
                    run_edct_analysis(path, connection=connection)

            path = build_edct_workbook(directory, row_overrides={3: {"Overseas": "MAYBE"}})
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)
            self.assertEqual(result.rule_totals[("overseas", "Overseas")], 1)
            result.workbook.close()

    def test_columns_only_structure_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "columns.xlsx"
            workbook = Workbook()
            workbook.active.title = "Supplier Level"
            workbook.active.append(["metadata"])
            workbook.active.append(["Index"])
            open_task = workbook.create_sheet("Open Task")
            open_task.append(["metadata"])
            open_task.append(["Punch Code"])
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                with self.assertRaisesRegex(EdctLoadError, "columns:"):
                    run_edct_analysis(path, connection=connection)

    def test_empty_assessment_and_missing_supplier(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(
                directory,
                row_overrides={3: {"Index": ""}, 4: {"Index": ""}},
                pn_rows=(("P1", "A"),),
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)
                self.assertEqual(result.assessed_rows, ())
                self.assertEqual(result.rows_failed, 0)
                self.assertEqual(
                    connection.execute("SELECT rows_in_scope FROM runs").fetchone(), (0,)
                )
                result.workbook.close()
            workbook = load_workbook(path)
            del workbook["Supplier Level"]
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                with self.assertRaisesRegex(EdctLoadError, "Supplier Level"):
                    run_edct_analysis(path, connection=connection)

    def test_missing_required_sheet_stops_analysis(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(Path(temp), include_open_task=False)
            workbook = load_workbook(path)
            supplier = workbook["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            supplier.cell(2, headers.index("EDI Mode") + 1).value = None
            workbook.save(path)
            workbook.close()
            connection = sqlite3.connect(":memory:")
            try:
                with self.assertRaisesRegex(EdctLoadError, "Open Task.*EDI Mode"):
                    run_edct_analysis(path, connection=connection)
            finally:
                connection.close()

    def test_all_input_sheets_resolve_columns_by_header_name(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    3: {"Triplet COFOR": "9999"},
                    4: {"Triplet COFOR": "8888"},
                },
                pn_rows=(("1003", "9999"),),
            )
            workbook = load_workbook(path)

            supplier = workbook["Supplier Level"]
            supplier_headers = [cell.value for cell in supplier[2]]
            punch_column = supplier_headers.index("Supplier Punch code") + 1
            name_column = supplier_headers.index("Supplier name") + 1
            for row in range(1, supplier.max_row + 1):
                punch_value = supplier.cell(row, punch_column).value
                supplier.cell(row, punch_column).value = supplier.cell(row, name_column).value
                supplier.cell(row, name_column).value = punch_value

            open_task = workbook["Open Task"]
            open_task["G2"] = open_task["A2"].value
            open_task["G3"] = open_task["A3"].value
            open_task["A2"] = None
            open_task["A3"] = None

            template = workbook["Template-Cofor-Creation"]
            template["H2"] = template["D2"].value
            template["H3"] = template["D3"].value
            template["D2"] = None
            template["D3"] = None

            pn = workbook["PN Level"]
            pn["F1"] = pn["A1"].value
            pn["F2"] = pn["A2"].value
            pn["H1"] = pn["B1"].value
            pn["H2"] = pn["B2"].value
            pn["A1"] = pn["A2"] = None
            pn["B1"] = pn["B2"] = None

            workbook.save(path)
            workbook.close()

            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)

            self.assertIn(("Supplier Level", 3), result.row_results)
            self.assertIn(("PN Level", 2), result.row_results)
            result.workbook.close()

    def test_cofor_template_uses_second_row_headers_without_fixed_column(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    3: {
                        "Supplier Punch code": "1003",
                        "Creation of Cofors request date": "10/08/2026",
                    }
                },
            )
            workbook = load_workbook(path)
            template = workbook["Template-Cofor-Creation"]
            template["D1"] = None
            template["D2"] = None
            template["D3"] = None
            template["E2"] = "Punch Code"
            template["E3"] = "1003"
            workbook.save(path)
            workbook.close()

            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)

            self.assertEqual(result.row_results[("Supplier Level", 3)].check, 0)
            result.workbook.close()

    def test_missing_or_malformed_cofor_template_stops_analysis(self) -> None:
        scenarios = (
            ("missing sheet", False, None, "Template-Cofor-Creation"),
            (
                "wrong row 2 header",
                True,
                "Wrong header",
                r"Template-Cofor-Creation\.Punch Code",
            ),
        )
        for label, include_sheet, header, expected_error in scenarios:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(
                    Path(temp),
                    include_cofor_template=include_sheet,
                )
                if header is not None:
                    workbook = load_workbook(path)
                    workbook["Template-Cofor-Creation"]["D2"] = header
                    workbook.save(path)
                    workbook.close()

                with (
                    closing(sqlite3.connect(":memory:")) as connection,
                    self.assertRaisesRegex(EdctLoadError, expected_error),
                ):
                    run_edct_analysis(path, connection=connection)

    def test_run_analysis_with_custom_header_settings(self) -> None:
        from quality_checker.checkers.edct.settings import EdctHeaderSettings

        settings = EdctHeaderSettings.default()
        settings.supplier_level["Supplier Punch code"] = "Code Fournisseur"
        settings.supplier_level["Supplier name"] = "Nom Fournisseur"
        settings.supplier_level["Effective kick-off date"] = "Date lancement"
        settings.supplier_level["Sales contact"] = "Contact Commercial"
        settings.open_task["Punch Code"] = "Code Tache"
        settings.pn_level["Punch seller"] = "Vendeur"

        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            # Create workbook with renamed headers
            workbook = Workbook()
            all_columns = [
                settings.get_header("Supplier Level", col) for col in EDCT_REQUIRED_COLUMNS
            ]
            supplier = workbook.active
            supplier.title = "Supplier Level"
            supplier.append(["metadata"] * len(all_columns))
            supplier.append(all_columns)

            base: dict[str, object] = {column: "" for column in all_columns}
            for row_number, index in ((3, "Metz_01"), (4, "Metz_02")):
                values = dict(base)
                values.update(
                    {
                        "Index": index,
                        "Code Fournisseur": 1000 + row_number,
                        "Nom Fournisseur": f"Supplier {row_number}",
                        "Date lancement": date(2026, 1, 1),
                        "Contact Commercial": "invalid-email-format",
                        "Cofor created date": "",
                        "Overseas": "",
                        "OPEN TASK": "YES" if row_number == 3 else "",
                    }
                )
                for col in EDCT_FORMULA_COLUMNS:
                    values[col] = f'=IF(A{row_number}="","",A{row_number})'
                supplier.append([values[col] for col in all_columns])

            table = Table(
                displayName="Tabella2",
                ref=f"A2:{get_column_letter(len(all_columns))}4",
            )
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            supplier.add_table(table)

            open_task = workbook.create_sheet("Open Task")
            open_task.append(["metadata"])
            open_task.append(["Code Tache"])
            open_task.append([1003])

            cofor_template = workbook.create_sheet("Template-Cofor-Creation")
            cofor_template["D2"] = "Punch Code"
            cofor_template["D3"] = "9999"

            pn = workbook.create_sheet("PN Level")
            pn.append(["Vendeur", "Triplet COFOR"])
            pn.append(["1003", "123456  78"])

            path = directory / "custom_edct.xlsx"
            workbook.save(path)
            set_formula_caches(
                path,
                {"xl/worksheets/sheet1.xml": {"G3": "123456  78", "G4": "123456  78"}},
            )

            with closing(sqlite3.connect(":memory:")) as connection:
                # Running without settings should fail with missing columns
                with self.assertRaises(EdctLoadError) as err:
                    run_edct_analysis(
                        path, connection=connection, settings=EdctHeaderSettings.default()
                    )
                self.assertIn("Supplier Punch code", str(err.exception))

                # Running with custom settings should succeed
                result = run_edct_analysis(
                    path,
                    connection=connection,
                    analysis_date=date(2026, 7, 30),
                    settings=settings,
                )
                output = export_edct_result(result, directory / "out")

            exported = load_workbook(output)
            self.addCleanup(exported.close)
            exported_headers = {cell.value: cell.column for cell in exported["Supplier Level"][2]}
            self.assertEqual(
                exported["Supplier Level"]
                .cell(3, exported_headers["Date lancement"])
                .number_format,
                "DD/MM/YYYY",
            )
            result.workbook.close()

            self.assertEqual(len(result.assessed_rows), 3)
            row_3 = result.row_results[("Supplier Level", 3)]
            self.assertGreater(row_3.check, 0)
            self.assertIn("Contact Commercial = invalid-email-format", row_3.comment)
            self.assertEqual(result.rule_totals[("email", "Contact Commercial")], 2)
            pn_row = result.row_results[("PN Level", 2)]
            self.assertEqual(pn_row.check, 0)
