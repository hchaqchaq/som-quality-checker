from __future__ import annotations

import re
import sqlite3
import tempfile
import unittest
from contextlib import closing
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from som_analyzer.analysis.edct import EdctLoadError, export_edct_result, run_edct_analysis
from som_analyzer.edct_config import EDCT_FORMULA_COLUMNS, EDCT_REQUIRED_COLUMNS


def build_edct_workbook(
    directory: Path,
    *,
    include_open_task: bool = True,
    row_overrides: dict[int, dict[str, object]] | None = None,
) -> Path:
    workbook = Workbook()
    supplier = workbook.active
    supplier.title = "Supplier Level"
    supplier.append(["metadata"] * len(EDCT_REQUIRED_COLUMNS))
    supplier.append(list(EDCT_REQUIRED_COLUMNS))

    base = {column: "" for column in EDCT_REQUIRED_COLUMNS}
    for row_number, index in ((3, "Metz_01"), (4, "Metz_02")):
        values = dict(base)
        values.update(
            {
                "Index": index,
                "Supplier Punch code": 1000 + row_number,
                "Supplier name": f"Supplier {row_number}",
                "Effective kick-off date": "",
                "Cofor created date": "",
                "Overseas": "",
                "OPEN TASK": "YES" if row_number == 3 else "",
            }
        )
        for column in EDCT_FORMULA_COLUMNS:
            values[column] = f"=IF(A{row_number}=\"\",\"\",A{row_number})"
        values.update((row_overrides or {}).get(row_number, {}))
        supplier.append([values[column] for column in EDCT_REQUIRED_COLUMNS])

    supplier.append([""] * len(EDCT_REQUIRED_COLUMNS))
    supplier["A3"].fill = PatternFill("solid", fgColor="FF0000")
    table = Table(
        displayName="Tabella2",
        ref=f"A2:{get_column_letter(len(EDCT_REQUIRED_COLUMNS))}5",
    )
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    supplier.add_table(table)

    if include_open_task:
        open_task = workbook.create_sheet("Open Task")
        open_task.append(["metadata"])
        open_task.append(["Punch Code"])
        open_task.append([1003])
    workbook.create_sheet("Other Sheet")["A1"] = "keep me"

    path = directory / "input.xlsx"
    workbook.save(path)
    workbook.close()
    return path


class EdctWorkbookTests(unittest.TestCase):
    def test_rule_documentation_lists_every_required_edct_column(self) -> None:
        documentation = (
            Path(__file__).resolve().parents[2] / "docs" / "EDCT_VALIDATION_RULES.md"
        ).read_text(encoding="utf-8")
        inventory = documentation.split("## Configured column inventory", 1)[1].split(
            "## Formula rules", 1
        )[0]
        documented = set(re.findall(r"^- `([^`]+)`$", inventory, re.MULTILINE))
        self.assertEqual(documented, set(EDCT_REQUIRED_COLUMNS))

    def test_complete_workbook_is_annotated_and_preserved(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            with closing(sqlite3.connect(":memory:")) as connection:
                connection.row_factory = sqlite3.Row
                result = run_edct_analysis(input_path, connection=connection)
                output_path = export_edct_result(result, directory)
                runs = connection.execute("SELECT project, rows_total FROM runs").fetchall()

            exported = load_workbook(output_path, data_only=False)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            self.assertEqual(result.processed_rows, (3, 4))
            self.assertEqual([row["project"] for row in runs], ["eDCT"])
            self.assertEqual([row["rows_total"] for row in runs], [2])
            self.assertEqual(exported.sheetnames, ["Supplier Level", "Open Task", "Other Sheet"])
            self.assertEqual(exported["Other Sheet"]["A1"].value, "keep me")
            self.assertEqual(supplier["A3"].fill.fgColor.rgb, "00FF0000")
            self.assertIn("Check", headers)
            self.assertIn("Comment", headers)
            self.assertEqual(supplier.cell(3, headers.index("Check") + 1).value, 0)
            self.assertEqual(supplier.cell(3, headers.index("Comment") + 1).value, "Quality check passed")
            self.assertTrue(supplier.tables["Tabella2"].ref.endswith(get_column_letter(len(headers)) + "5"))
            self.assertRegex(output_path.name, r"input_eDCT_checked_\d{8}_\d{6}\.xlsx")
            self.assertEqual(load_workbook(input_path)["Other Sheet"]["A1"].value, "keep me")
            exported.close()

    def test_missing_required_sheet_stops_analysis(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(Path(temp), include_open_task=False)
            workbook = load_workbook(path)
            supplier = workbook["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            supplier.cell(2, headers.index("EDI Mode") + 1).value = None
            workbook.save(path)
            workbook.close()
            connection = sqlite3.connect(":memory:")
            try:
                with self.assertRaisesRegex(EdctLoadError, "Open Task.*EDI Mode"):
                    run_edct_analysis(path, connection=connection)
            finally:
                connection.close()

    def test_export_failure_marks_the_history_run_failed(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(directory)
            blocked_output = directory / "not-a-directory"
            blocked_output.write_text("file", encoding="utf-8")
            connection = sqlite3.connect(":memory:")
            connection.row_factory = sqlite3.Row
            try:
                result = run_edct_analysis(path, connection=connection)
                with self.assertRaises(OSError):
                    export_edct_result(result, blocked_output)
                run = connection.execute(
                    "SELECT status, error_message FROM runs WHERE id = ?",
                    (result.run_id,),
                ).fetchone()
            finally:
                connection.close()

        self.assertEqual(run["status"], "failed")
        self.assertTrue(run["error_message"])

    def test_field_formats_are_reported_with_exact_column_names(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    3: {
                        "Effective kick-off date": "30.07.2026",
                        "eSupplierConnect": "NOT",
                        "B2B": "NOT",
                        "New supplier portal": "NOT",
                        "SPM": "NOT",
                        "iTMS": "NOT",
                        "OPEN TASK": "YES",
                        "Sales contact": "first@example.com, second@example.com",
                        "Seller COFOR": "ABC",
                        "Phone": "call me",
                        "First communication sent": "30/07/2026",
                        "Readiness Comments": "30/07/2026: contacted",
                        "Participants": "Name - person@example.com",
                    }
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection, analysis_date=date(2026, 7, 30))

        self.assertEqual(result.row_results[3].check, 6)
        for column in (
            "Sales contact",
            "Seller COFOR",
            "Phone",
            "First communication sent",
            "Readiness Comments",
            "Participants",
        ):
            self.assertIn(column, result.row_results[3].comment)

    def test_lifecycle_and_open_task_conditions_are_reported(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    3: {
                        "Cofor created date": "invalid date",
                        "Triple Status": "",
                        "EDI Mode": "",
                        "Effective kick-off date": "30.07.2026",
                        "Overseas": "YES",
                        "Shipping location": "",
                        "Supplier Confimation": "NO",
                        "OPEN TASK": "",
                    },
                    4: {"OPEN TASK": "NO"},
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                connection.row_factory = sqlite3.Row
                result = run_edct_analysis(path, connection=connection, analysis_date=date(2026, 7, 30))
                totals = connection.execute(
                    "SELECT column_name, fail_count FROM run_columns WHERE run_id = ?",
                    (result.run_id,),
                ).fetchall()

        self.assertEqual(result.row_results[3].check, 11)
        self.assertEqual(result.row_results[4].check, 1)
        for column in (
            "Triple Status",
            "EDI Mode",
            "Shipping location",
            "Supplier Confimation",
            "OPEN TASK",
            "eSupplierConnect",
            "B2B",
            "New supplier portal",
            "SPM",
            "iTMS",
        ):
            self.assertIn(column, result.row_results[3].comment)
        self.assertEqual(sum(row["fail_count"] for row in totals), 12)

    def test_formula_structure_uses_first_processed_row_as_reference(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    4: {
                        "Starting date": '=IF(B4="","",B4)',
                        "Triplet COFOR": '= + if ( a4 = "" , "" , a4 )',
                    }
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection, analysis_date=date(2026, 7, 30))

        self.assertEqual(result.row_results[3].check, 0)
        self.assertEqual(result.row_results[4].check, 1)
        self.assertIn("Starting date", result.row_results[4].comment)
        self.assertNotIn("Triplet COFOR", result.row_results[4].comment)

    def test_missing_first_formula_flags_every_processed_row_and_continues(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={3: {"Starting date": None}},
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection, analysis_date=date(2026, 7, 30))

        expected = (
            "Formula validation failed: first processed row is missing "
            "the reference formula for Starting date"
        )
        self.assertEqual(result.row_results[3].check, 1)
        self.assertEqual(result.row_results[4].check, 1)
        self.assertIn(expected, result.row_results[3].comment)
        self.assertIn(expected, result.row_results[4].comment)


if __name__ == "__main__":
    unittest.main()
