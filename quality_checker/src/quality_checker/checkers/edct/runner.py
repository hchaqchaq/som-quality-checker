from __future__ import annotations

import re
import sqlite3
from collections import Counter, defaultdict
from collections.abc import Callable
from copy import copy
from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path
from posixpath import basename, dirname, normpath
from tempfile import NamedTemporaryFile
from time import perf_counter
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import TableColumn
from openpyxl.worksheet.worksheet import Worksheet

from ...application import DB_PATH
from ...db.repository import (
    ColumnRecord,
    RunRecord,
    initialize_schema,
    insert_run,
    open_connection,
    update_run_exported_file,
    update_run_status,
)
from .config import (
    EDCT_COFOR_COLUMNS,
    EDCT_COFOR_REQUEST_DATE_COLUMN,
    EDCT_COFOR_TEMPLATE_HEADER_ROW,
    EDCT_COFOR_TEMPLATE_PUNCH_HEADER,
    EDCT_COFOR_TEMPLATE_SHEET,
    EDCT_DATE_COLUMNS,
    EDCT_DATED_COMMENT_COLUMNS,
    EDCT_EDI_MODE_VALUES,
    EDCT_EMAIL_COLUMNS,
    EDCT_FORMULA_COLUMNS,
    EDCT_HEADER_ROW,
    EDCT_INDEX_COLUMN,
    EDCT_INDEX_COLUMNS,
    EDCT_PN_HEADER_ROW,
    EDCT_PN_REQUIRED_COLUMNS,
    EDCT_PN_SELLER_COLUMN,
    EDCT_PN_SHEET,
    EDCT_PORTAL_COLUMNS,
    EDCT_PORTAL_VALUES,
    EDCT_PROJECT,
    EDCT_REQUIRED_COLUMNS,
    EDCT_REQUIRED_SHEETS,
    EDCT_TRIPLE_STATUS_VALUES,
)
from .settings import EdctHeaderSettings, load_edct_settings

EMAIL_REGEX = re.compile(r"^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$")
COFOR_REGEX = re.compile(r"^[A-Za-z0-9]{6} {2}[A-Za-z0-9]{2}$")


class EdctLoadError(Exception):
    """Raised when an eDCT workbook does not contain the required structure."""


@dataclass(frozen=True, slots=True)
class EdctRowResult:
    check: int
    comment: str


@dataclass(slots=True)
class EdctRunResult:
    run_id: int
    input_file: Path
    started_at: datetime
    finished_at: datetime
    duration_s: float
    workbook: Workbook
    assessed_rows: tuple[tuple[str, int], ...]
    row_results: dict[tuple[str, int], EdctRowResult]
    rule_totals: Counter[tuple[str, str]]
    uses_default_db: bool
    history_connection: sqlite3.Connection | None
    pn_values: dict[int, tuple[object, object]]

    @property
    def rows_failed(self) -> int:
        return sum(result.check > 0 for result in self.row_results.values())


def _normalized_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _normalized_choice(value: object) -> str:
    return _normalized_text(value).casefold()


def _is_valid_email(value: object) -> bool:
    text = _normalized_text(value)
    if not text:
        return True
    text = re.sub(r"(?:;\s*)+$", "", text)
    parts = [part.strip() for part in text.split(";")]
    return bool(parts) and all(part and EMAIL_REGEX.fullmatch(part) for part in parts)


def _parse_date(
    value: object,
    *,
    formats: tuple[str, ...] = ("%d/%m/%Y",),
) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _normalized_text(value)
    if not text:
        return None
    for date_format in formats:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    return None


def _is_valid_dated_comment(value: object) -> bool:
    text = _normalized_text(value)
    if not text:
        return True
    date_text, separator, comment = text.partition(":")
    return bool(separator and comment.strip() and _parse_date(date_text.strip()))


def _normalized_punch(value: object) -> str:
    text = _normalized_text(value)
    return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


def _formula_key(formula: object) -> str | None:
    if not isinstance(formula, str) or not formula.startswith("="):
        return None
    normalized: list[str] = []
    quoted = False
    for character in formula:
        if character == '"':
            quoted = not quoted
            normalized.append(character)
        elif character.isspace() and not quoted:
            continue
        else:
            normalized.append(character if quoted else character.casefold())
    result = "".join(normalized)
    return "=" + result[2:] if result.startswith("=+") else result


def _open_task_punches(workbook, settings: EdctHeaderSettings) -> set[str]:
    worksheet = workbook["Open Task"]
    headers = _header_map(worksheet)
    punch_col_name = settings.get_header("Open Task", "Punch Code")
    column = headers.get(punch_col_name)
    if column is None:
        raise EdctLoadError(f"Missing required columns in Open Task: {punch_col_name}")
    return {
        punch
        for row in range(EDCT_HEADER_ROW + 1, worksheet.max_row + 1)
        if (punch := _normalized_punch(worksheet.cell(row, column).value))
    }


def _cofor_template_punches(workbook, column: int) -> set[str]:
    worksheet = workbook[EDCT_COFOR_TEMPLATE_SHEET]
    return {
        punch
        for row in range(EDCT_COFOR_TEMPLATE_HEADER_ROW + 1, worksheet.max_row + 1)
        if (punch := _normalized_choice(worksheet.cell(row, column).value))
    }


def _evaluate_business_rules(
    workbook,
    headers: dict[str, int],
    assessed_rows: tuple[int, ...],
    cofor_template_column: int,
    analysis_date: date,
    settings: EdctHeaderSettings,
) -> tuple[dict[int, EdctRowResult], Counter[tuple[str, str]]]:
    worksheet = workbook["Supplier Level"]
    failures: dict[int, list[tuple[str, str, object]]] = {row: [] for row in assessed_rows}
    totals: Counter[tuple[str, str]] = Counter()

    def get_supplier_header(canonical: str) -> str:
        return settings.get_header("Supplier Level", canonical)

    def value(row: int, column: str) -> object:
        header_name = get_supplier_header(column)
        return worksheet.cell(row, headers[header_name]).value

    def fail(row: int, rule: str, column: str, reason: str) -> None:
        header_name = get_supplier_header(column)
        failures[row].append((reason, header_name, value(row, column)))
        totals[(rule, header_name)] += 1

    open_task_punches = _open_task_punches(workbook, settings)
    triple_status_values = {item.casefold() for item in EDCT_TRIPLE_STATUS_VALUES}
    portal_values = {item.casefold() for item in EDCT_PORTAL_VALUES}
    edi_mode_values = {item.casefold() for item in EDCT_EDI_MODE_VALUES}
    cofor_template_punches = _cofor_template_punches(workbook, cofor_template_column)
    if assessed_rows:
        reference_row = assessed_rows[0]
        for column in EDCT_FORMULA_COLUMNS:
            header_name = get_supplier_header(column)
            reference_cell = worksheet.cell(reference_row, headers[header_name])
            reference_formula = _formula_key(reference_cell.value)
            if reference_formula is None:
                reason = (
                    "Formula validation failed: formula reference row is missing "
                    f"the reference formula for {header_name}"
                )
                for row in assessed_rows:
                    fail(row, "formula", column, reason)
                continue
            for row in assessed_rows[1:]:
                target_cell = worksheet.cell(row, headers[header_name])
                expected = Translator(
                    str(reference_cell.value),
                    origin=reference_cell.coordinate,
                ).translate_formula(target_cell.coordinate)
                if _formula_key(target_cell.value) != _formula_key(expected):
                    fail(row, "formula", column, "Formula does not match formula reference row")

    for row in assessed_rows:
        effective_date_value = value(row, "Effective kick-off date")
        effective_date = _parse_date(effective_date_value)
        cofor_date_value = value(row, "Cofor created date")
        effective_date_filled = bool(_normalized_text(effective_date_value))
        cofor_date_filled = bool(_normalized_text(cofor_date_value))

        for column in EDCT_EMAIL_COLUMNS:
            if column in ("Sales contact", "Logistic contact") and not effective_date_filled:
                continue
            if not _is_valid_email(value(row, column)):
                fail(row, "email", column, "Invalid email")

        if effective_date_filled:
            for column in EDCT_COFOR_COLUMNS:
                text = _normalized_text(value(row, column))
                if text and COFOR_REGEX.fullmatch(text) is None:
                    fail(row, "cofor", column, "Invalid COFOR format")

        for column in EDCT_DATE_COLUMNS:
            raw = value(row, column)
            if _normalized_text(raw) and _parse_date(raw) is None:
                fail(
                    row,
                    "date",
                    column,
                    "Invalid date format, expected DD/MM/YYYY",
                )
        if effective_date is not None and effective_date > analysis_date:
            fail(row, "date_future", "Effective kick-off date", "Date cannot be in the future")

        for column in EDCT_DATED_COMMENT_COLUMNS:
            if not _is_valid_dated_comment(value(row, column)):
                fail(row, "dated_comment", column, "Invalid dated comment")
        triple_status = _normalized_choice(value(row, "Triple Status"))
        if cofor_date_filled and triple_status not in triple_status_values:
            fail(row, "triple_status", "Triple Status", "Required value must be Valid or No Valid")

        overseas = _normalized_text(value(row, "Overseas"))
        if overseas not in {"", "YES", "NOT"}:
            fail(
                row,
                "overseas",
                "Overseas",
                "Invalid value, expected YES, NOT, or empty",
            )

        supplier_confirmation = _normalized_choice(value(row, "Supplier Confimation"))
        if supplier_confirmation not in {"", "yes"}:
            fail(row, "supplier_confirmation", "Supplier Confimation", "Invalid value")

        for column in EDCT_PORTAL_COLUMNS:
            portal_value = _normalized_choice(value(row, column))
            allowed = portal_values if effective_date_filled else {"", *portal_values}
            if portal_value not in allowed:
                fail(row, "portal", column, "Required value must be YES or NOT")

        edi_mode = _normalized_choice(value(row, "EDI Mode"))
        allowed_edi_modes = edi_mode_values if cofor_date_filled else {"", *edi_mode_values}
        if edi_mode not in allowed_edi_modes:
            fail(row, "edi_mode", "EDI Mode", "Required value must be WEB EDI or Standard EDI")

        request_date_filled = bool(_normalized_text(value(row, EDCT_COFOR_REQUEST_DATE_COLUMN)))
        punch_in_cofor_template = (
            _normalized_choice(value(row, "Supplier Punch code")) in cofor_template_punches
        )
        request_date_col_name = get_supplier_header(EDCT_COFOR_REQUEST_DATE_COLUMN)
        supplier_punch_col_name = get_supplier_header("Supplier Punch code")
        if punch_in_cofor_template and not request_date_filled:
            fail(
                row,
                "date_required",
                EDCT_COFOR_REQUEST_DATE_COLUMN,
                f"{request_date_col_name} is required because the Punch Code exists "
                "in Template-Cofor-Creation",
            )
        if request_date_filled and not punch_in_cofor_template:
            fail(
                row,
                "cofor_template",
                "Supplier Punch code",
                f"{supplier_punch_col_name} must exist in Template-Cofor-Creation when "
                f"{request_date_col_name} is populated",
            )
        punch_in_open_task = (
            _normalized_punch(value(row, "Supplier Punch code")) in open_task_punches
        )
        open_task_value = _normalized_choice(value(row, "OPEN TASK"))
        if punch_in_open_task and open_task_value != "yes":
            fail(row, "open_task", "OPEN TASK", "OPEN TASK must be YES for a matching punch code")
        elif not punch_in_open_task and open_task_value:
            fail(
                row,
                "open_task",
                "OPEN TASK",
                "OPEN TASK must be empty when the punch code is absent",
            )

    row_results: dict[int, EdctRowResult] = {}
    for row, row_failures in failures.items():
        if not row_failures:
            row_results[row] = EdctRowResult(0, "Quality check passed")
            continue
        grouped: dict[str, list[str]] = {}
        for reason, column, invalid_value in row_failures:
            display_value = _normalized_text(invalid_value) or "<empty>"
            grouped.setdefault(reason, []).append(f"{column} = {display_value}")
        comment = " | ".join(
            f"{reason}: {', '.join(columns)}" for reason, columns in grouped.items()
        )
        row_results[row] = EdctRowResult(len(row_failures), comment)
    return row_results, totals


def _header_map(worksheet, row: int = EDCT_HEADER_ROW) -> dict[str, int]:
    return {
        _normalized_text(cell.value): cell.column
        for cell in worksheet[row]
        if _normalized_text(cell.value)
    }


def _worksheet_parts(archive: ZipFile) -> dict[str, str]:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    relationship_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    relationships = {
        element.attrib["Id"]: element.attrib["Target"]
        for element in ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    }
    root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    return {
        sheet.attrib["name"]: normpath("xl/" + relationships[sheet.attrib[relationship_id]]).lstrip(
            "/"
        )
        if not relationships[sheet.attrib[relationship_id]].startswith("/")
        else relationships[sheet.attrib[relationship_id]].lstrip("/")
        for sheet in root.findall(f"{{{namespace}}}sheets/{{{namespace}}}sheet")
    }


def _table_part(archive: ZipFile, worksheet_part: str, table_name: str) -> str:
    relationships_part = f"{dirname(worksheet_part)}/_rels/{basename(worksheet_part)}.rels"
    relationships = ElementTree.fromstring(archive.read(relationships_part))
    for relationship in relationships:
        if not relationship.attrib["Type"].endswith("/table"):
            continue
        target = relationship.attrib["Target"]
        part = (
            target.lstrip("/")
            if target.startswith("/")
            else normpath(f"{dirname(worksheet_part)}/{target}")
        )
        table = ElementTree.fromstring(archive.read(part))
        if table.attrib.get("displayName") == table_name:
            return part
    raise EdctLoadError(f"Missing source table part: {table_name}")


def _evaluate_pn_triplets(
    workbook: Workbook,
    input_path: Path,
    supplier_headers: dict[str, int],
    supplier_rows: tuple[int, ...],
    pn_headers: dict[str, int],
    settings: EdctHeaderSettings,
) -> tuple[dict[tuple[str, int], EdctRowResult], dict[int, tuple[object, object]]]:
    pn = workbook[EDCT_PN_SHEET]
    configured_pn_seller = settings.get_header(EDCT_PN_SHEET, EDCT_PN_SELLER_COLUMN)
    configured_pn_triplet = settings.get_header(EDCT_PN_SHEET, "Triplet COFOR")
    configured_supplier_punch = settings.get_header("Supplier Level", "Supplier Punch code")
    configured_supplier_triplet = settings.get_header("Supplier Level", "Triplet COFOR")

    seller_column = pn_headers[configured_pn_seller]
    pn_triplet_column = pn_headers[configured_pn_triplet]
    supplier_punch_column = supplier_headers[configured_supplier_punch]
    supplier_triplet_column = supplier_headers[configured_supplier_triplet]

    candidate_rows = [
        row
        for row in range(EDCT_PN_HEADER_ROW + 1, pn.max_row + 1)
        if _normalized_text(pn.cell(row, seller_column).value)
    ]
    if not candidate_rows or not supplier_rows:
        return {}, {}

    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    empty_cached_formulas: set[tuple[str, str]] = set()
    with ZipFile(input_path) as archive:
        parts = _worksheet_parts(archive)
        for name in ("Supplier Level", EDCT_PN_SHEET):
            root = ElementTree.fromstring(archive.read(parts[name]))
            for cell in root.iter(f"{{{namespace}}}c"):
                cached_value = cell.find(f"{{{namespace}}}v")
                if (
                    cell.attrib.get("t") == "str"
                    and cell.find(f"{{{namespace}}}f") is not None
                    and cached_value is not None
                    and cached_value.text is None
                ):
                    empty_cached_formulas.add((name, cell.attrib["r"]))

    cached = load_workbook(input_path, data_only=True)
    try:

        def resolved_value(sheet_name: str, row: int, column: int) -> object:
            cell = workbook[sheet_name].cell(row, column)
            if cell.data_type != "f":
                return cell.value
            cached_cell = cached[sheet_name].cell(row, column)
            value = cached_cell.value
            if value is None or cached_cell.data_type == "e":
                if value is None and (sheet_name, cell.coordinate) in empty_cached_formulas:
                    return ""
                reason = (
                    f"Cached formula error {value}"
                    if cached_cell.data_type == "e"
                    else "Missing cached formula value"
                )
                raise EdctLoadError(
                    f"{reason}: {sheet_name}!{cell.coordinate}. "
                    "Correct formula errors, recalculate and save the workbook in Excel "
                    "before analysis."
                )
            return value

        sellers = {row: resolved_value(EDCT_PN_SHEET, row, seller_column) for row in candidate_rows}
        if not any(_normalized_choice(seller) for seller in sellers.values()):
            return {}, {}
        allowed_triplets: dict[str, set[str]] = defaultdict(set)
        for row in supplier_rows:
            punch = _normalized_choice(resolved_value("Supplier Level", row, supplier_punch_column))
            if not punch:
                continue
            allowed = allowed_triplets[punch]
            triplet = _normalized_choice(
                resolved_value("Supplier Level", row, supplier_triplet_column)
            )
            if triplet:
                allowed.add(triplet)

        row_results: dict[tuple[str, int], EdctRowResult] = {}
        pn_values: dict[int, tuple[object, object]] = {}
        for row, seller in sellers.items():
            punch = _normalized_choice(seller)
            if not punch or punch not in allowed_triplets:
                continue
            triplet_value = resolved_value(EDCT_PN_SHEET, row, pn_triplet_column)
            pn_values[row] = (seller, triplet_value)
            triplet = _normalized_choice(triplet_value)
            allowed = allowed_triplets[punch]
            if triplet and triplet in allowed:
                row_result = EdctRowResult(0, "Quality check passed")
            else:
                allowed_display = ", ".join(repr(value) for value in sorted(allowed))
                allowed_display = allowed_display or "<no nonblank allowed triplets>"
                rejected = _normalized_text(triplet_value) or "<blank>"
                row_result = EdctRowResult(
                    1,
                    f"{configured_pn_triplet} = {rejected} is not allowed for "
                    f"{configured_pn_seller} = {_normalized_text(seller)}; "
                    f"allowed {configured_pn_triplet} values: {allowed_display}",
                )
            row_results[(EDCT_PN_SHEET, row)] = row_result
        return row_results, pn_values
    finally:
        cached.close()


def _open_history(connection: sqlite3.Connection | None) -> tuple[sqlite3.Connection, bool]:
    if connection is not None:
        initialize_schema(connection)
        return connection, False
    own_connection = open_connection(DB_PATH)
    initialize_schema(own_connection)
    return own_connection, True


def _insert_history(
    connection: sqlite3.Connection,
    *,
    started_at: datetime,
    finished_at: datetime,
    duration_s: float,
    input_file: Path,
    rows_total: int,
    rows_failed: int,
    status: str,
    error_message: str | None,
    rule_totals: Counter[tuple[str, str]] | None = None,
) -> int:
    columns = [
        ColumnRecord(rule_name=rule, column_name=column, fail_count=count)
        for (rule, column), count in (rule_totals or {}).items()
    ]
    return insert_run(
        connection,
        RunRecord(
            project=EDCT_PROJECT,
            started_at=started_at.isoformat(),
            finished_at=finished_at.isoformat(),
            duration_s=duration_s,
            input_file=str(input_file),
            exported_file=None,
            rows_total=rows_total,
            rows_in_scope=rows_total,
            rows_failed=rows_failed,
            status=status,
            error_message=error_message,
        ),
        columns,
    )


def run_edct_analysis(
    input_path: Path | str,
    connection: sqlite3.Connection | None = None,
    *,
    analysis_date: date | None = None,
    settings: EdctHeaderSettings | None = None,
) -> EdctRunResult:
    active_settings = settings if settings is not None else load_edct_settings()
    resolved_input = Path(input_path)
    started_at = datetime.now(UTC)
    started_perf = perf_counter()
    history, own_connection = _open_history(connection)

    try:
        if not resolved_input.exists():
            raise EdctLoadError(f"Input file not found: {resolved_input}")
        workbook = load_workbook(resolved_input, data_only=False)
        missing_sheets = [
            sheet for sheet in EDCT_REQUIRED_SHEETS if sheet not in workbook.sheetnames
        ]
        missing_columns: list[str] = []
        headers: dict[str, int] = {}
        pn_headers: dict[str, int] = {}
        supplier_index_header = active_settings.get_header("Supplier Level", EDCT_INDEX_COLUMN)
        if "Supplier Level" in workbook.sheetnames:
            headers = _header_map(workbook["Supplier Level"])
            for column in EDCT_REQUIRED_COLUMNS:
                if column == EDCT_INDEX_COLUMN:
                    continue
                configured_header = active_settings.get_header("Supplier Level", column)
                if configured_header not in headers:
                    missing_columns.append(configured_header)
            if supplier_index_header == EDCT_INDEX_COLUMN:
                if not any(column in headers for column in EDCT_INDEX_COLUMNS):
                    missing_columns.append("Index or Line")
            else:
                if supplier_index_header not in headers:
                    missing_columns.append(supplier_index_header)

        if EDCT_PN_SHEET in workbook.sheetnames:
            pn_headers = _header_map(workbook[EDCT_PN_SHEET], EDCT_PN_HEADER_ROW)
            for column in EDCT_PN_REQUIRED_COLUMNS:
                configured_header = active_settings.get_header(EDCT_PN_SHEET, column)
                if configured_header not in pn_headers:
                    missing_columns.append(f"{EDCT_PN_SHEET}.{configured_header}")

        open_task_punch_header = active_settings.get_header("Open Task", "Punch Code")
        if "Open Task" in workbook.sheetnames and open_task_punch_header not in _header_map(
            workbook["Open Task"]
        ):
            missing_columns.append(f"Open Task.{open_task_punch_header}")

        expected_cofor_header = active_settings.get_header(
            EDCT_COFOR_TEMPLATE_SHEET, EDCT_COFOR_TEMPLATE_PUNCH_HEADER
        )
        cofor_template_column: int | None = None
        if EDCT_COFOR_TEMPLATE_SHEET in workbook.sheetnames:
            cofor_template_headers = _header_map(
                workbook[EDCT_COFOR_TEMPLATE_SHEET],
                EDCT_COFOR_TEMPLATE_HEADER_ROW,
            )
            cofor_template_column = cofor_template_headers.get(expected_cofor_header)
            if cofor_template_column is None:
                missing_columns.append(f"{EDCT_COFOR_TEMPLATE_SHEET}.{expected_cofor_header}")
        if missing_sheets or missing_columns:
            workbook.close()
            parts = []
            if missing_sheets:
                parts.append(f"sheets: {', '.join(missing_sheets)}")
            if missing_columns:
                parts.append(f"columns: {', '.join(missing_columns)}")
            raise EdctLoadError(f"Missing required structure: {'; '.join(parts)}")
        assert cofor_template_column is not None

        supplier = workbook["Supplier Level"]
        if supplier_index_header in headers:
            index_column = supplier_index_header
        elif supplier_index_header == EDCT_INDEX_COLUMN and "Line" in headers:
            index_column = "Line"
        else:
            index_column = supplier_index_header

        supplier_rows = tuple(
            row
            for row in range(EDCT_HEADER_ROW + 1, supplier.max_row + 1)
            if _normalized_text(supplier.cell(row, headers[index_column]).value)
        )
        pn_results, pn_values = _evaluate_pn_triplets(
            workbook, resolved_input, headers, supplier_rows, pn_headers, active_settings
        )
        supplier_results, rule_totals = _evaluate_business_rules(
            workbook,
            headers,
            supplier_rows,
            cofor_template_column,
            analysis_date or date.today(),
            active_settings,
        )
        row_results = {
            ("Supplier Level", row): row_result for row, row_result in supplier_results.items()
        }
        row_results.update(pn_results)
        pn_failures = sum(row_result.check for row_result in pn_results.values())
        if pn_failures:
            configured_pn_triplet = active_settings.get_header(EDCT_PN_SHEET, "Triplet COFOR")
            rule_totals[("pn_triplet_cofor", configured_pn_triplet)] = pn_failures
        assessed_rows = tuple(row_results)
        finished_at = datetime.now(UTC)
        duration_s = perf_counter() - started_perf
        run_id = _insert_history(
            history,
            started_at=started_at,
            finished_at=finished_at,
            duration_s=duration_s,
            input_file=resolved_input,
            rows_total=len(assessed_rows),
            rows_failed=sum(result.check > 0 for result in row_results.values()),
            status="ok",
            error_message=None,
            rule_totals=rule_totals,
        )
        return EdctRunResult(
            run_id=run_id,
            input_file=resolved_input,
            started_at=started_at,
            finished_at=finished_at,
            duration_s=duration_s,
            workbook=workbook,
            assessed_rows=assessed_rows,
            row_results=row_results,
            rule_totals=rule_totals,
            uses_default_db=own_connection,
            history_connection=None if own_connection else history,
            pn_values=pn_values,
        )
    except Exception as exc:
        if "workbook" in locals():
            workbook.close()
        finished_at = datetime.now(UTC)
        _insert_history(
            history,
            started_at=started_at,
            finished_at=finished_at,
            duration_s=perf_counter() - started_perf,
            input_file=resolved_input,
            rows_total=0,
            rows_failed=0,
            status="failed",
            error_message=str(exc),
        )
        raise
    finally:
        if own_connection:
            history.close()


def _result_columns(worksheet, header_row: int = EDCT_HEADER_ROW) -> tuple[int, int]:
    headers = _header_map(worksheet, header_row)
    if worksheet.title == EDCT_PN_SHEET:
        columns: list[int] = []
        for name in ("Check", "Comment"):
            column = headers.get(name)
            if column is None:
                column = worksheet.max_column + 1
                worksheet.cell(header_row, column, name)
            columns.append(column)
        return columns[0], columns[1]
    table = worksheet.tables.get("Tabella2")
    if table is None:
        raise EdctLoadError("Missing required table: Tabella2")

    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    result_columns: list[int] = []
    for name in ("Check", "Comment"):
        column = headers.get(name)
        if column is None:
            column = max(max_col, worksheet.max_column) + 1
            worksheet.cell(EDCT_HEADER_ROW, column, name)
        if column > max_col:
            for added_column in range(max_col + 1, column + 1):
                header = _normalized_text(worksheet.cell(EDCT_HEADER_ROW, added_column).value)
                if not header:
                    header = f"Column {added_column}"
                    worksheet.cell(EDCT_HEADER_ROW, added_column, header)
                table.tableColumns.append(TableColumn(id=len(table.tableColumns) + 1, name=header))
            max_col = column
        result_columns.append(column)

    table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
    if table.autoFilter is not None:
        table.autoFilter.ref = table.ref
    return result_columns[0], result_columns[1]


def _update_result_history(
    result: EdctRunResult,
    action: Callable[[sqlite3.Connection], None],
) -> None:
    if result.run_id < 0:
        return
    if result.history_connection is not None:
        action(result.history_connection)
        return
    if result.uses_default_db:
        connection = open_connection(DB_PATH)
        try:
            initialize_schema(connection)
            action(connection)
        finally:
            connection.close()


def _preserve_ooxml_extensions(
    source: Path,
    target: Path,
    replacement_parts: dict[str, str],
) -> None:
    def local_name(element) -> str:
        return element.tag.rsplit("}", 1)[-1]

    def restore_extensions(source_element, target_element) -> None:
        source_children: dict[str, list] = {}
        target_children: dict[str, list] = {}
        for child in source_element:
            source_children.setdefault(local_name(child), []).append(child)
        for child in target_element:
            target_children.setdefault(local_name(child), []).append(child)

        for name, children in source_children.items():
            if name == "extLst":
                for existing in target_children.get(name, []):
                    target_element.remove(existing)
                target_element.extend(children)
                continue
            for source_child, target_child in zip(
                children, target_children.get(name, []), strict=False
            ):
                restore_extensions(source_child, target_child)

    def restore_identified_extensions(source_root, target_root) -> None:
        target_elements = list(target_root.iter())
        for source_element in source_root.iter():
            extensions = [child for child in source_element if local_name(child) == "extLst"]
            identity = {
                attribute.rsplit("}", 1)[-1]: value
                for attribute, value in source_element.attrib.items()
                if attribute.rsplit("}", 1)[-1] in {"id", "name"}
            }
            if not extensions or not identity:
                continue
            matches = [
                element
                for element in target_elements
                if local_name(element) == local_name(source_element)
                and all(
                    any(
                        attribute.rsplit("}", 1)[-1] == name and value == expected
                        for attribute, value in element.attrib.items()
                    )
                    for name, expected in identity.items()
                )
            ]
            if len(matches) == 1:
                for existing in list(matches[0]):
                    if local_name(existing) == "extLst":
                        matches[0].remove(existing)
                matches[0].extend(extensions)

    def restore_formula_values(source_root, target_root) -> None:
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        source_cells = {
            cell.attrib["r"]: cell
            for cell in source_root.iter(f"{{{namespace}}}c")
            if cell.find(f"{{{namespace}}}f") is not None
        }
        for cell in target_root.iter(f"{{{namespace}}}c"):
            original = source_cells.get(cell.attrib["r"])
            if original is None or cell.find(f"{{{namespace}}}f") is None:
                continue
            if "t" in original.attrib:
                cell.set("t", original.attrib["t"])
            else:
                cell.attrib.pop("t", None)
            for child in list(cell):
                if local_name(child) in {"f", "v", "is"}:
                    cell.remove(child)
            for child in original:
                if local_name(child) in {"f", "v", "is"}:
                    cell.append(copy(child))

    def preserve_namespace_compatibility(source_xml: bytes, source_root) -> None:
        root_match = re.search(rb"<(?:[A-Za-z_][\w.-]*:)?worksheet\b", source_xml)
        if root_match is None:
            return
        root_end = source_xml.index(b">", root_match.start())
        root_start = source_xml[root_match.start() : root_end + 1].decode("utf-8")
        namespaces = {
            (match.group(1) or ""): match.group(2)
            for match in re.finditer(
                r'xmlns(?::([A-Za-z_][\w.-]*))?="([^"]+)"',
                root_start,
            )
        }
        used_namespaces = {
            name[1:].split("}", 1)[0]
            for element in source_root.iter()
            for name in (element.tag, *element.attrib)
            if name.startswith("{")
        }
        for prefix, namespace in namespaces.items():
            if namespace in used_namespaces and not re.fullmatch(r"ns\d+", prefix):
                ElementTree.register_namespace(prefix, namespace)

        compatibility = "http://schemas.openxmlformats.org/markup-compatibility/2006"
        ignorable_key = f"{{{compatibility}}}Ignorable"
        ignorable = source_root.attrib.get(ignorable_key)
        if ignorable:
            retained = [
                prefix for prefix in ignorable.split() if namespaces.get(prefix) in used_namespaces
            ]
            if retained:
                source_root.set(ignorable_key, " ".join(retained))
            else:
                source_root.attrib.pop(ignorable_key)

    with NamedTemporaryFile(dir=target.parent, suffix=".xlsx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with (
            ZipFile(source) as source_archive,
            ZipFile(target) as target_archive,
            ZipFile(temporary_path, "w", ZIP_DEFLATED) as output_archive,
        ):
            original_parts = {
                generated: original for original, generated in replacement_parts.items()
            }
            for item in target_archive.infolist():
                content = target_archive.read(item.filename)
                if item.filename in original_parts:
                    source_xml = source_archive.read(original_parts[item.filename])
                    source_root = ElementTree.fromstring(source_xml)
                    has_extensions = any(
                        local_name(element) == "extLst" for element in source_root.iter()
                    )
                    is_worksheet = local_name(source_root) == "worksheet"
                    if is_worksheet:
                        preserve_namespace_compatibility(source_xml, source_root)
                    if has_extensions or is_worksheet or local_name(source_root) == "table":
                        target_root = ElementTree.fromstring(content)
                        if has_extensions:
                            restore_extensions(source_root, target_root)
                            restore_identified_extensions(source_root, target_root)
                        if is_worksheet:
                            restore_formula_values(source_root, target_root)
                            annotated_children = {local_name(child): child for child in target_root}
                            for index, child in enumerate(list(source_root)):
                                name = local_name(child)
                                if (
                                    name in {"dimension", "sheetData"}
                                    and name in annotated_children
                                ):
                                    source_root.remove(child)
                                    source_root.insert(index, annotated_children[name])
                            target_root = source_root
                        elif local_name(source_root) == "table":
                            target_root.set("id", source_root.attrib["id"])
                        content = ElementTree.tostring(
                            target_root,
                            encoding="utf-8",
                            xml_declaration=True,
                        )
                output_archive.writestr(item, content)
        temporary_path.replace(target)
    finally:
        temporary_path.unlink(missing_ok=True)


def _merge_annotated_parts(
    source: Path,
    annotated: Path,
    target: Path,
    replacement_parts: dict[str, str],
) -> None:
    with (
        ZipFile(source) as source_archive,
        ZipFile(annotated) as annotated_archive,
        ZipFile(target, "w", ZIP_DEFLATED) as output_archive,
    ):
        missing = set(replacement_parts.values()).difference(annotated_archive.namelist())
        if missing:
            raise EdctLoadError(f"Missing generated workbook parts: {', '.join(sorted(missing))}")
        for item in source_archive.infolist():
            replacement = replacement_parts.get(item.filename)
            content = (
                annotated_archive.read(replacement)
                if replacement is not None
                else source_archive.read(item.filename)
            )
            output_archive.writestr(item, content)


def _save_annotated_worksheets(
    workbook: Workbook,
    worksheets: tuple[Worksheet, ...],
    path: Path,
) -> None:
    original_sheets = workbook._sheets
    original_active_sheet_index = workbook._active_sheet_index
    workbook._sheets = list(worksheets)
    workbook._active_sheet_index = 0
    try:
        workbook.save(path)
    finally:
        workbook._sheets = original_sheets
        workbook._active_sheet_index = original_active_sheet_index


def _validate_analysis_workbook(path: Path) -> None:
    try:
        with path.open("rb") as exported_file:
            workbook = load_workbook(exported_file, read_only=True, data_only=False)
            try:
                for sheet_name in ("Supplier Level", EDCT_PN_SHEET):
                    for row in workbook[sheet_name].iter_rows():
                        for cell in row:
                            _ = cell.value, cell.number_format
            finally:
                workbook.close()
    except Exception as exc:
        path.unlink(missing_ok=True)
        raise EdctLoadError(f"Analysis workbook validation failed: {exc}") from exc


def export_edct_result(result: EdctRunResult, output_dir: Path | str) -> Path:
    try:
        return _export_edct_result(result, output_dir)
    except Exception as error:
        error_message = str(error)

        def mark_export_failed(connection: sqlite3.Connection) -> None:
            update_run_status(connection, result.run_id, "failed", error_message)

        _update_result_history(result, mark_export_failed)
        raise


def _export_edct_result(result: EdctRunResult, output_dir: Path | str) -> Path:
    target_dir = Path(output_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    worksheets = (
        (result.workbook["Supplier Level"], EDCT_HEADER_ROW),
        (result.workbook[EDCT_PN_SHEET], EDCT_PN_HEADER_ROW),
    )
    result_columns: dict[str, tuple[int, int]] = {}
    for worksheet, header_row in worksheets:
        check_column, comment_column = _result_columns(worksheet, header_row)
        result_columns[worksheet.title] = (check_column, comment_column)
        style_source = max(1, min(check_column, comment_column) - 1)
        for column in (check_column, comment_column):
            for row in range(header_row, worksheet.max_row + 1):
                source = worksheet.cell(row, style_source)
                target = worksheet.cell(row, column)
                target._style = copy(source._style)
        for row in range(header_row + 1, worksheet.max_row + 1):
            worksheet.cell(row, check_column).value = None
            worksheet.cell(row, comment_column).value = None
        if worksheet.title == "Supplier Level":
            headers = _header_map(worksheet)
            for column in EDCT_DATE_COLUMNS:
                column_index = headers[column]
                for row in range(header_row + 1, worksheet.max_row + 1):
                    cell = worksheet.cell(row, column_index)
                    if isinstance(cell.value, (date, datetime)):
                        cell.number_format = "DD/MM/YYYY"

    for (sheet_name, row), row_result in result.row_results.items():
        worksheet = result.workbook[sheet_name]
        check_column, comment_column = result_columns[sheet_name]
        worksheet.cell(row, check_column, row_result.check)
        worksheet.cell(row, comment_column, row_result.comment)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_stem = re.sub(r"[^A-Za-z0-9._-]+", "_", result.input_file.stem).strip("._") or "analysis"
    target = target_dir / f"{safe_stem}_eDCT_checked_{timestamp}.xlsx"
    supplier = result.workbook["Supplier Level"]
    table = supplier.tables["Tabella2"]
    with NamedTemporaryFile(dir=target_dir, suffix=".xlsx", delete=False) as temporary:
        annotated = Path(temporary.name)
    try:
        _save_annotated_worksheets(
            result.workbook,
            tuple(worksheet for worksheet, _ in worksheets),
            annotated,
        )
        with ZipFile(result.input_file) as source_archive:
            source_sheets = _worksheet_parts(source_archive)
            source_table = _table_part(
                source_archive, source_sheets["Supplier Level"], table.displayName
            )
            replacement_parts = {
                source_sheets[worksheet.title]: worksheet.path.lstrip("/")
                for worksheet, _ in worksheets
            }
            replacement_parts[source_table] = table.path.lstrip("/")
            replacement_parts["xl/styles.xml"] = "xl/styles.xml"
        _preserve_ooxml_extensions(result.input_file, annotated, replacement_parts)
        _merge_annotated_parts(result.input_file, annotated, target, replacement_parts)
        _validate_analysis_workbook(target)
    except Exception:
        target.unlink(missing_ok=True)
        raise
    finally:
        annotated.unlink(missing_ok=True)

    _update_result_history(
        result,
        lambda connection: update_run_exported_file(connection, result.run_id, str(target)),
    )
    return target
