from __future__ import annotations

from collections import Counter, defaultdict
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook

from .config import EDCT_PN_COFOR_COLUMNS, EDCT_PN_HEADER_ROW, EDCT_PN_SELLER_COLUMN, EDCT_PN_SHEET
from .models import EdctLoadError, EdctRowResult
from .ooxml import worksheet_parts
from .validator import COFOR_REGEX
from .workbook import LoadedEdctWorkbook, normalized_choice, normalized_text


def _normalized_text(value: object) -> str:
    return normalized_text(value)


def _normalized_choice(value: object) -> str:
    return normalized_choice(value)


def _normalized_triplet(value: object) -> str:
    return _normalized_choice(value).replace("\u00a0", " ")


def validate_pn(
    source: LoadedEdctWorkbook,
) -> tuple[
    dict[tuple[str, int], EdctRowResult], dict[int, tuple[object, object]], Counter[tuple[str, str]]
]:
    workbook = source.workbook
    input_path = source.path
    supplier_headers = source.supplier_headers
    supplier_rows = source.supplier_rows
    pn_headers = source.pn_headers
    settings = source.settings
    pn = workbook[EDCT_PN_SHEET]
    configured_pn_seller = settings.get_header(EDCT_PN_SHEET, EDCT_PN_SELLER_COLUMN)
    configured_pn_triplet = settings.get_header(EDCT_PN_SHEET, "Triplet COFOR")
    configured_supplier_punch = settings.get_header("Supplier Level", "Supplier Punch code")
    configured_supplier_triplet = settings.get_header("Supplier Level", "Triplet COFOR")

    seller_column = pn_headers[configured_pn_seller]
    pn_triplet_column = pn_headers[configured_pn_triplet]
    cofor_columns: dict[str, int] = {}
    for name in EDCT_PN_COFOR_COLUMNS:
        header = settings.get_header(EDCT_PN_SHEET, name)
        if header in pn_headers:
            cofor_columns[header] = pn_headers[header]
    supplier_punch_column = supplier_headers[configured_supplier_punch]
    supplier_triplet_column = supplier_headers[configured_supplier_triplet]

    candidate_rows = [
        row
        for row in range(EDCT_PN_HEADER_ROW + 1, pn.max_row + 1)
        if _normalized_text(pn.cell(row, seller_column).value)
        or any(_normalized_text(pn.cell(row, column).value) for column in cofor_columns.values())
    ]
    if not candidate_rows:
        return {}, {}, Counter()

    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    empty_cached_formulas: set[tuple[str, str]] = set()
    with ZipFile(input_path) as archive:
        parts = worksheet_parts(archive)
        for name in ("Supplier Level", EDCT_PN_SHEET):
            root = ElementTree.fromstring(archive.read(parts[name]))
            for cell in root.iter(f"{{{namespace}}}c"):
                cached_value = cell.find(f"{{{namespace}}}v")
                if (
                    cell.attrib.get("t") == "str"
                    and cell.find(f"{{{namespace}}}f") is not None
                    and cached_value is not None
                    and cached_value.text is None
                ):
                    empty_cached_formulas.add((name, cell.attrib["r"]))

    cached = load_workbook(input_path, data_only=True)
    try:

        def resolved_value(sheet_name: str, row: int, column: int) -> object:
            cell = workbook[sheet_name].cell(row, column)
            if cell.data_type != "f":
                return cell.value
            cached_cell = cached[sheet_name].cell(row, column)
            value = cached_cell.value
            if value is None or cached_cell.data_type == "e":
                if value is None and (sheet_name, cell.coordinate) in empty_cached_formulas:
                    return ""
                reason = (
                    f"Cached formula error {value}"
                    if cached_cell.data_type == "e"
                    else "Missing cached formula value"
                )
                raise EdctLoadError(
                    f"{reason}: {sheet_name}!{cell.coordinate}. "
                    "Correct formula errors, recalculate and save the workbook in Excel "
                    "before analysis."
                )
            return value

        sellers = {row: resolved_value(EDCT_PN_SHEET, row, seller_column) for row in candidate_rows}
        allowed_triplets: dict[str, set[str]] = defaultdict(set)
        for row in supplier_rows:
            punch = _normalized_choice(resolved_value("Supplier Level", row, supplier_punch_column))
            if not punch:
                continue
            allowed = allowed_triplets[punch]
            triplet = _normalized_triplet(
                resolved_value("Supplier Level", row, supplier_triplet_column)
            )
            if triplet:
                allowed.add(triplet)

        row_results: dict[tuple[str, int], EdctRowResult] = {}
        pn_values: dict[int, tuple[object, object]] = {}
        totals: Counter[tuple[str, str]] = Counter()
        for row, seller in sellers.items():
            punch = _normalized_choice(seller)
            known_seller = bool(punch and punch in allowed_triplets)
            failures: list[str] = []
            has_cofor_value = False
            for name, column in cofor_columns.items():
                raw = resolved_value(EDCT_PN_SHEET, row, column)
                text = _normalized_text(raw).replace("\u00a0", " ")
                has_cofor_value |= bool(text)
                if text and COFOR_REGEX.fullmatch(text) is None:
                    failures.append(f"{name} = {_normalized_text(raw)}: Invalid COFOR format")
                    totals[("pn_cofor", name)] += 1
            if not has_cofor_value and not known_seller:
                continue
            triplet_cell = pn.cell(row, pn_triplet_column)
            if known_seller:
                triplet_value = resolved_value(EDCT_PN_SHEET, row, pn_triplet_column)
            elif triplet_cell.data_type == "f":
                triplet_value = cached[EDCT_PN_SHEET].cell(row, pn_triplet_column).value
            else:
                triplet_value = triplet_cell.value
            pn_values[row] = (seller, triplet_value)
            if known_seller:
                triplet = _normalized_triplet(triplet_value)
                allowed = allowed_triplets[punch]
                if not triplet or triplet not in allowed:
                    allowed_display = ", ".join(repr(value) for value in sorted(allowed))
                    allowed_display = allowed_display or "<no nonblank allowed triplets>"
                    rejected = _normalized_text(triplet_value) or "<blank>"
                    failures.append(
                        f"{configured_pn_triplet} = {rejected} is not allowed for "
                        f"{configured_pn_seller} = {_normalized_text(seller)}; "
                        f"allowed {configured_pn_triplet} values: {allowed_display}"
                    )
                    totals[("pn_triplet_cofor", configured_pn_triplet)] += 1
            row_results[(EDCT_PN_SHEET, row)] = EdctRowResult(
                len(failures), " | ".join(failures) if failures else "Quality check passed"
            )
        return row_results, pn_values, totals
    finally:
        cached.close()
