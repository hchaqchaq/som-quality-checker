from __future__ import annotations

import json
from collections.abc import Mapping
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from ...application import DATA_DIR, ensure_data_dir
from .config import (
    EDCT_COFOR_TEMPLATE_HEADER_ROW,
    EDCT_COFOR_TEMPLATE_PUNCH_HEADER,
    EDCT_COFOR_TEMPLATE_SHEET,
    EDCT_HEADER_ROW,
    EDCT_PN_HEADER_ROW,
    EDCT_PN_REQUIRED_COLUMNS,
    EDCT_PN_SHEET,
    EDCT_REQUIRED_COLUMNS,
)

SETTINGS_FILE_NAME = "edct_header_mappings.json"
SUPPLIER_LEVEL_SHEET = "Supplier Level"
OPEN_TASK_SHEET = "Open Task"
OPEN_TASK_PUNCH_HEADER = "Punch Code"
HEADER_ALIASES = {
    (SUPPLIER_LEVEL_SHEET, "Index"): ("Line",),
    (SUPPLIER_LEVEL_SHEET, "Supplier Confimation"): ("Supplier Confirmation",),
}


def get_default_supplier_level_headers() -> dict[str, str]:
    return {column: column for column in EDCT_REQUIRED_COLUMNS}


def get_default_pn_level_headers() -> dict[str, str]:
    return {column: column for column in EDCT_PN_REQUIRED_COLUMNS}


def get_default_open_task_headers() -> dict[str, str]:
    return {OPEN_TASK_PUNCH_HEADER: OPEN_TASK_PUNCH_HEADER}


def get_default_cofor_template_headers() -> dict[str, str]:
    return {EDCT_COFOR_TEMPLATE_PUNCH_HEADER: EDCT_COFOR_TEMPLATE_PUNCH_HEADER}


@dataclass(slots=True)
class EdctHeaderSettings:
    supplier_level: dict[str, str] = field(default_factory=get_default_supplier_level_headers)
    pn_level: dict[str, str] = field(default_factory=get_default_pn_level_headers)
    open_task: dict[str, str] = field(default_factory=get_default_open_task_headers)
    cofor_template: dict[str, str] = field(default_factory=get_default_cofor_template_headers)

    @classmethod
    def default(cls) -> EdctHeaderSettings:
        return cls(
            supplier_level=get_default_supplier_level_headers(),
            pn_level=get_default_pn_level_headers(),
            open_task=get_default_open_task_headers(),
            cofor_template=get_default_cofor_template_headers(),
        )

    def get_sheet_mapping(self, sheet: str) -> dict[str, str]:
        if sheet == SUPPLIER_LEVEL_SHEET:
            return self.supplier_level
        if sheet == EDCT_PN_SHEET:
            return self.pn_level
        if sheet == OPEN_TASK_SHEET:
            return self.open_task
        if sheet == EDCT_COFOR_TEMPLATE_SHEET:
            return self.cofor_template
        return {}

    def get_header(self, sheet: str, canonical_name: str) -> str:
        mapping = self.get_sheet_mapping(sheet)
        return mapping.get(canonical_name, canonical_name)

    def to_dict(self) -> dict[str, dict[str, str]]:
        return {
            SUPPLIER_LEVEL_SHEET: dict(self.supplier_level),
            EDCT_PN_SHEET: dict(self.pn_level),
            OPEN_TASK_SHEET: dict(self.open_task),
            EDCT_COFOR_TEMPLATE_SHEET: dict(self.cofor_template),
        }

    @classmethod
    def from_dict(cls, data: Mapping[str, object]) -> EdctHeaderSettings:
        default = cls.default()

        def merge_mapping(sheet_key: str, default_mapping: dict[str, str]) -> dict[str, str]:
            sheet_data = data.get(sheet_key)
            if not isinstance(sheet_data, dict):
                return dict(default_mapping)
            merged = dict(default_mapping)
            for key, val in sheet_data.items():
                if isinstance(key, str) and isinstance(val, str):
                    merged[key] = val
            return merged

        return cls(
            supplier_level=merge_mapping(SUPPLIER_LEVEL_SHEET, default.supplier_level),
            pn_level=merge_mapping(EDCT_PN_SHEET, default.pn_level),
            open_task=merge_mapping(OPEN_TASK_SHEET, default.open_task),
            cofor_template=merge_mapping(EDCT_COFOR_TEMPLATE_SHEET, default.cofor_template),
        )

    def validate(self) -> list[str]:
        errors: list[str] = []
        sheets = (
            (SUPPLIER_LEVEL_SHEET, self.supplier_level),
            (EDCT_PN_SHEET, self.pn_level),
            (OPEN_TASK_SHEET, self.open_task),
            (EDCT_COFOR_TEMPLATE_SHEET, self.cofor_template),
        )
        for sheet_name, mapping in sheets:
            seen_headers: dict[str, str] = {}
            for canonical, configured in mapping.items():
                stripped = configured.strip()
                if not stripped:
                    errors.append(f"{sheet_name}: header for '{canonical}' cannot be empty.")
                    continue
                accepted = {
                    canonical.strip().casefold(),
                    stripped.casefold(),
                    *(
                        alias.strip().casefold()
                        for alias in HEADER_ALIASES.get((sheet_name, canonical), ())
                    ),
                }
                collisions = accepted.intersection(seen_headers)
                if collisions:
                    normalized = sorted(collisions)[0]
                    first_field = seen_headers[normalized]
                    errors.append(
                        f"{sheet_name}: duplicate or ambiguous header mapping for "
                        f"'{first_field}' and '{canonical}'."
                    )
                for normalized in accepted:
                    seen_headers.setdefault(normalized, canonical)
        return errors


def get_settings_path() -> Path:
    return DATA_DIR / SETTINGS_FILE_NAME


def load_edct_settings(path: Path | None = None) -> EdctHeaderSettings:
    target_path = path or get_settings_path()
    if not target_path.exists():
        return EdctHeaderSettings.default()
    try:
        content = target_path.read_text(encoding="utf-8")
        data = json.loads(content)
        if isinstance(data, dict):
            return EdctHeaderSettings.from_dict(data)
    except Exception:
        pass
    return EdctHeaderSettings.default()


def save_edct_settings(settings: EdctHeaderSettings, path: Path | None = None) -> None:
    target_path = path or get_settings_path()
    ensure_data_dir()
    data = settings.to_dict()
    target_path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


def inspect_workbook_headers(path: Path | str) -> dict[str, list[str]]:
    resolved = Path(path)
    if not resolved.exists():
        raise FileNotFoundError(f"Workbook not found: {resolved}")
    workbook = load_workbook(resolved, read_only=True, data_only=True)
    try:
        results: dict[str, list[str]] = {}
        sheet_header_rows = {
            SUPPLIER_LEVEL_SHEET: (EDCT_HEADER_ROW,),
            OPEN_TASK_SHEET: (EDCT_HEADER_ROW,),
            EDCT_PN_SHEET: (EDCT_PN_HEADER_ROW,),
            EDCT_COFOR_TEMPLATE_SHEET: (EDCT_COFOR_TEMPLATE_HEADER_ROW,),
        }
        for sheet_name, header_rows in sheet_header_rows.items():
            if sheet_name in workbook.sheetnames:
                ws = workbook[sheet_name]
                headers: list[str] = []
                for row_idx in header_rows:
                    for cell in ws[row_idx]:
                        if cell.value is not None:
                            text = str(cell.value).strip()
                            if text and text not in headers:
                                headers.append(text)
                results[sheet_name] = headers
            else:
                results[sheet_name] = []
        return results
    finally:
        workbook.close()
