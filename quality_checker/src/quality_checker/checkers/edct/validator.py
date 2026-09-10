from __future__ import annotations

import re
from collections import Counter
from datetime import date, datetime

from openpyxl.formula.translate import Translator

from .config import (
    EDCT_COFOR_COLUMNS,
    EDCT_COFOR_REQUEST_DATE_COLUMN,
    EDCT_DATE_COLUMNS,
    EDCT_DATED_COMMENT_COLUMNS,
    EDCT_EDI_MODE_VALUES,
    EDCT_EMAIL_COLUMNS,
    EDCT_FORMULA_COLUMNS,
    EDCT_PORTAL_COLUMNS,
    EDCT_PORTAL_VALUES,
    EDCT_TRIPLE_STATUS_VALUES,
)
from .models import EdctRowResult
from .workbook import LoadedEdctWorkbook, normalized_punch

EMAIL_REGEX = re.compile(r"^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$")
COFOR_REGEX = re.compile(r"^[A-Za-z0-9]{6} {2}[A-Za-z0-9]{2}$")


def _normalized_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _normalized_choice(value: object) -> str:
    return _normalized_text(value).casefold()


def _is_valid_email(value: object) -> bool:
    text = _normalized_text(value)
    if not text:
        return True
    text = re.sub(r"(?:;\s*)+$", "", text)
    parts = [part.strip() for part in text.split(";")]
    return bool(parts) and all(part and EMAIL_REGEX.fullmatch(part) for part in parts)


def _parse_date(
    value: object,
    *,
    formats: tuple[str, ...] = ("%d/%m/%Y",),
) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _normalized_text(value)
    if not text:
        return None
    for date_format in formats:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    return None


def _is_valid_dated_comment(value: object) -> bool:
    text = _normalized_text(value)
    if not text:
        return True
    date_text, separator, comment = text.partition(":")
    return bool(separator and comment.strip() and _parse_date(date_text.strip()))


def _formula_key(formula: object) -> str | None:
    if not isinstance(formula, str) or not formula.startswith("="):
        return None
    normalized: list[str] = []
    quoted = False
    for character in formula:
        if character == '"':
            quoted = not quoted
            normalized.append(character)
        elif character.isspace() and not quoted:
            continue
        else:
            normalized.append(character if quoted else character.casefold())
    result = "".join(normalized)
    return "=" + result[2:] if result.startswith("=+") else result


def validate_suppliers(
    source: LoadedEdctWorkbook,
    analysis_date: date,
) -> tuple[dict[int, EdctRowResult], Counter[tuple[str, str]]]:
    workbook = source.workbook
    headers = source.supplier_headers
    assessed_rows = source.supplier_rows
    settings = source.settings
    worksheet = workbook["Supplier Level"]
    failures: dict[int, list[tuple[str, str, object]]] = {row: [] for row in assessed_rows}
    totals: Counter[tuple[str, str]] = Counter()

    def get_supplier_header(canonical: str) -> str:
        return settings.get_header("Supplier Level", canonical)

    def value(row: int, column: str) -> object:
        header_name = get_supplier_header(column)
        return worksheet.cell(row, headers[header_name]).value

    def fail(row: int, rule: str, column: str, reason: str) -> None:
        header_name = get_supplier_header(column)
        failures[row].append((reason, header_name, value(row, column)))
        totals[(rule, header_name)] += 1

    open_task_punches = source.open_task_punches()
    triple_status_values = {item.casefold() for item in EDCT_TRIPLE_STATUS_VALUES}
    portal_values = {item.casefold() for item in EDCT_PORTAL_VALUES}
    edi_mode_values = {item.casefold() for item in EDCT_EDI_MODE_VALUES}
    cofor_template_punches = source.cofor_template_punches()
    if assessed_rows:
        reference_row = assessed_rows[0]
        for column in EDCT_FORMULA_COLUMNS:
            header_name = get_supplier_header(column)
            reference_cell = worksheet.cell(reference_row, headers[header_name])
            reference_formula = _formula_key(reference_cell.value)
            if reference_formula is None:
                reason = (
                    "Formula validation failed: formula reference row is missing "
                    f"the reference formula for {header_name}"
                )
                for row in assessed_rows:
                    fail(row, "formula", column, reason)
                continue
            for row in assessed_rows[1:]:
                target_cell = worksheet.cell(row, headers[header_name])
                expected = Translator(
                    str(reference_cell.value),
                    origin=reference_cell.coordinate,
                ).translate_formula(target_cell.coordinate)
                if _formula_key(target_cell.value) != _formula_key(expected):
                    fail(row, "formula", column, "Formula does not match formula reference row")

    for row in assessed_rows:
        effective_date_value = value(row, "Effective kick-off date")
        effective_date = _parse_date(effective_date_value)
        cofor_date_value = value(row, "Cofor created date")
        effective_date_filled = bool(_normalized_text(effective_date_value))
        cofor_date_filled = bool(_normalized_text(cofor_date_value))

        for column in EDCT_EMAIL_COLUMNS:
            if column in ("Sales contact", "Logistic contact") and not effective_date_filled:
                continue
            if not _is_valid_email(value(row, column)):
                fail(row, "email", column, "Invalid email")

        if effective_date_filled:
            for column in EDCT_COFOR_COLUMNS:
                text = _normalized_text(value(row, column))
                if text and COFOR_REGEX.fullmatch(text) is None:
                    fail(row, "cofor", column, "Invalid COFOR format")

        for column in EDCT_DATE_COLUMNS:
            raw = value(row, column)
            if _normalized_text(raw) and _parse_date(raw) is None:
                fail(
                    row,
                    "date",
                    column,
                    "Invalid date format, expected DD/MM/YYYY",
                )
        if effective_date is not None and effective_date > analysis_date:
            fail(row, "date_future", "Effective kick-off date", "Date cannot be in the future")

        for column in EDCT_DATED_COMMENT_COLUMNS:
            if not _is_valid_dated_comment(value(row, column)):
                fail(row, "dated_comment", column, "Invalid dated comment")
        triple_status = _normalized_choice(value(row, "Triple Status"))
        if cofor_date_filled and triple_status not in triple_status_values:
            fail(row, "triple_status", "Triple Status", "Required value must be Valid or No Valid")

        overseas = _normalized_text(value(row, "Overseas"))
        if overseas not in {"", "YES", "NOT"}:
            fail(
                row,
                "overseas",
                "Overseas",
                "Invalid value, expected YES, NOT, or empty",
            )

        supplier_confirmation = _normalized_choice(value(row, "Supplier Confimation"))
        if supplier_confirmation not in {"", "yes"}:
            fail(row, "supplier_confirmation", "Supplier Confimation", "Invalid value")

        for column in EDCT_PORTAL_COLUMNS:
            portal_value = _normalized_choice(value(row, column))
            allowed = portal_values if effective_date_filled else {"", *portal_values}
            if portal_value not in allowed:
                fail(row, "portal", column, "Required value must be YES or NOT")

        edi_mode = _normalized_choice(value(row, "EDI Mode"))
        allowed_edi_modes = edi_mode_values if cofor_date_filled else {"", *edi_mode_values}
        if edi_mode not in allowed_edi_modes:
            fail(row, "edi_mode", "EDI Mode", "Required value must be WEB EDI or Standard EDI")

        request_date_filled = bool(_normalized_text(value(row, EDCT_COFOR_REQUEST_DATE_COLUMN)))
        punch_in_cofor_template = (
            _normalized_choice(value(row, "Supplier Punch code")) in cofor_template_punches
        )
        request_date_col_name = get_supplier_header(EDCT_COFOR_REQUEST_DATE_COLUMN)
        supplier_punch_col_name = get_supplier_header("Supplier Punch code")
        if punch_in_cofor_template and not request_date_filled:
            fail(
                row,
                "date_required",
                EDCT_COFOR_REQUEST_DATE_COLUMN,
                f"{request_date_col_name} is required because the Punch Code exists "
                "in Template-Cofor-Creation",
            )
        if request_date_filled and not punch_in_cofor_template:
            fail(
                row,
                "cofor_template",
                "Supplier Punch code",
                f"{supplier_punch_col_name} must exist in Template-Cofor-Creation when "
                f"{request_date_col_name} is populated",
            )
        punch_in_open_task = (
            normalized_punch(value(row, "Supplier Punch code")) in open_task_punches
        )
        open_task_value = _normalized_choice(value(row, "OPEN TASK"))
        if punch_in_open_task and open_task_value != "yes":
            fail(row, "open_task", "OPEN TASK", "OPEN TASK must be YES for a matching punch code")
        elif not punch_in_open_task and open_task_value:
            fail(
                row,
                "open_task",
                "OPEN TASK",
                "OPEN TASK must be empty when the punch code is absent",
            )

    row_results: dict[int, EdctRowResult] = {}
    for row, row_failures in failures.items():
        if not row_failures:
            row_results[row] = EdctRowResult(0, "Quality check passed")
            continue
        grouped: dict[str, list[str]] = {}
        for reason, column, invalid_value in row_failures:
            display_value = _normalized_text(invalid_value) or "<empty>"
            grouped.setdefault(reason, []).append(f"{column} = {display_value}")
        comment = " | ".join(
            f"{reason}: {', '.join(columns)}" for reason, columns in grouped.items()
        )
        row_results[row] = EdctRowResult(len(row_failures), comment)
    return row_results, totals
