from __future__ import annotations

import sqlite3
import tempfile
from contextlib import closing
from datetime import date
from pathlib import Path

from edct_support import (
    EdctTestCase,
    build_edct_workbook,
    set_formula_caches,
)
from openpyxl import load_workbook
from quality_checker.checkers.edct.export import export_edct_result
from quality_checker.checkers.edct.models import EdctLoadError
from quality_checker.checkers.edct.runner import run_edct_analysis
from quality_checker.db.repository import list_runs


class EdctPnTests(EdctTestCase):
    def test_pn_required_structure_stops_analysis_and_records_failure(self) -> None:
        for missing in ("PN Level", "Punch seller", "Triplet COFOR"):
            with self.subTest(missing=missing), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(Path(temp))
                workbook = load_workbook(path)
                if missing == "PN Level":
                    del workbook["PN Level"]
                else:
                    column = 1 if missing == "Punch seller" else 2
                    workbook["PN Level"].cell(1, column).value = "Unexpected header"
                workbook.save(path)
                workbook.close()
                with closing(sqlite3.connect(":memory:")) as connection:
                    connection.row_factory = sqlite3.Row
                    with self.assertRaises(EdctLoadError) as raised:
                        run_edct_analysis(path, connection=connection)
                    self.assertIn(missing, str(raised.exception))
                    history = list_runs(connection, "eDCT")
                    self.assertEqual(len(history), 1)
                    self.assertEqual(history[0]["status"], "failed")
                    self.assertIsNone(history[0]["exported_file"])

    def test_required_formula_caches_fail_actionably_and_allow_recovery(self) -> None:
        scenarios = (
            ("Supplier Level", "Supplier Punch code", 3, "1003"),
            ("Supplier Level", "Triplet COFOR", 3, "A"),
            ("PN Level", "Punch seller", 2, "1003"),
            ("PN Level", "Triplet COFOR", 2, "A"),
        )
        for sheet_name, column_name, row, cached_value in scenarios:
            with self.subTest(sheet=sheet_name, column=column_name):
                with tempfile.TemporaryDirectory() as temp:
                    path = build_edct_workbook(
                        Path(temp),
                        row_overrides={
                            3: {"Triplet COFOR": "A"},
                            4: {"Triplet COFOR": "A"},
                        },
                        pn_rows=((1003, "A"), ("unknown", "=A3"), ("", "=A4")),
                    )
                    workbook = load_workbook(path)
                    sheet = workbook[sheet_name]
                    header_row = 1 if sheet_name == "PN Level" else 2
                    column = next(
                        cell.column for cell in sheet[header_row] if cell.value == column_name
                    )
                    cell = sheet.cell(row, column, '="cached identifier"')
                    coordinate = cell.coordinate
                    workbook.save(path)
                    workbook.close()
                    with closing(sqlite3.connect(":memory:")) as connection:
                        connection.row_factory = sqlite3.Row
                        with self.assertRaises(EdctLoadError) as raised:
                            run_edct_analysis(path, connection=connection)
                        self.assertIn(f"{sheet_name}!{coordinate}", str(raised.exception))
                        self.assertIn("Excel", str(raised.exception))
                        self.assertEqual(list_runs(connection, "eDCT")[0]["status"], "failed")
                        part = "sheet5" if sheet_name == "PN Level" else "sheet1"
                        set_formula_caches(
                            path,
                            {f"xl/worksheets/{part}.xml": {coordinate: "#REF!"}},
                            data_type="e",
                        )
                        with self.assertRaises(EdctLoadError) as cached_error:
                            run_edct_analysis(path, connection=connection)
                        self.assertIn(f"{sheet_name}!{coordinate}", str(cached_error.exception))
                        self.assertIn("#REF!", str(cached_error.exception))
                        set_formula_caches(
                            path, {f"xl/worksheets/{part}.xml": {coordinate: cached_value}}
                        )
                        result = run_edct_analysis(path, connection=connection)
                        self.addCleanup(result.workbook.close)
                        self.assertEqual(result.row_results[("PN Level", 2)].check, 0)
                        self.assertNotIn(("PN Level", 3), result.row_results)
                        self.assertNotIn(("PN Level", 4), result.row_results)
                        self.assertEqual(
                            {run["status"] for run in list_runs(connection, "eDCT")},
                            {"failed", "ok"},
                        )

    def test_pn_blank_scope_and_normalization_boundaries(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(
                directory,
                row_overrides={
                    3: {
                        "Supplier Punch code": " P1 ",
                        "Triplet COFOR": " A, B/C ",
                        "OPEN TASK": "",
                    },
                    4: {"Supplier Punch code": "p1", "Triplet COFOR": "a, b/c"},
                },
                pn_rows=(
                    ("p1", "a, b/c"),
                    (" P1 ", " A, B/C "),
                    ("P1", "A"),
                    ("P1", "a,b/c"),
                    ("P1", "a, b/c"),
                    ("P1", "bad"),
                    ("P1", "bad"),
                    ("P1", None),
                    (None, "a, b/c"),
                    ("unknown", "=uncached()"),
                    ("EMPTY", "A"),
                    ("EMPTY", None),
                    ("001", "001"),
                    ("001", "1"),
                    ("1", "001"),
                    ("P 1", "A"),
                    ("P-1", "A"),
                    ("EXCLUDED", "A"),
                ),
                optional_columns=("Line",),
            )
            workbook = load_workbook(path)
            supplier = workbook["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            additions = (
                ("Metz_03", "EMPTY", None, None),
                ("Metz_04", "", "A", None),
                ("Metz_05", "001", "001", None),
                (None, "P1", "bad", "fallback-must-not-apply"),
                (None, "EXCLUDED", "A", "fallback-must-not-apply"),
                ("Metz_06", "P1", "unused", None),
            )
            for row, (index, punch, triplet, line) in enumerate(additions, 5):
                for column, value in (
                    ("Index", index),
                    ("Supplier Punch code", punch),
                    ("Triplet COFOR", triplet),
                    ("Line", line),
                ):
                    supplier.cell(row, headers.index(column) + 1).value = value
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)
                self.addCleanup(result.workbook.close)
                pn_checks = {
                    row: result.row_results[(sheet, row)].check
                    for sheet, row in result.assessed_rows
                    if sheet == "PN Level"
                }
                self.assertEqual(
                    pn_checks,
                    {2: 0, 3: 0, 4: 1, 5: 1, 6: 0, 7: 1, 8: 1, 9: 1, 12: 1, 13: 1, 14: 0, 15: 1},
                )
                self.assertEqual(result.rule_totals[("pn_triplet_cofor", "Triplet COFOR")], 8)
                self.assertNotIn(("Supplier Level", 8), result.assessed_rows)
                self.assertNotIn(("Supplier Level", 9), result.assessed_rows)
                self.assertIn("blank", result.row_results[("PN Level", 9)].comment.lower())
                self.assertIn("no nonblank", result.row_results[("PN Level", 12)].comment.lower())
                comment = result.row_results[("PN Level", 4)].comment
                self.assertEqual(comment.count("'a, b/c'"), 1)
                self.assertLess(comment.index("'a, b/c'"), comment.index("'unused'"))

    def test_pn_cofor_format_assesses_each_populated_field_without_seller(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    3: {"Supplier Punch code": "P1", "Triplet COFOR": "not-a-cofor"},
                    4: {"Triplet COFOR": ""},
                },
                pn_rows=(("P1", "not-a-cofor"), (None, None), (None, None)),
            )
            workbook = load_workbook(path)
            pn = workbook["PN Level"]
            pn["C1"] = "Shipper COFOR"
            pn["D1"] = "Manufacturer COFOR"
            pn["E1"] = "Seller COFOR"
            pn["F1"] = "Empty return COFOR"
            pn["C2"] = "ABC123\u00a0 45"
            pn["D2"] = "bad"
            pn["E2"] = "ABC123  45"
            pn["F2"] = "wrong"
            pn["C3"] = "bad shipper"
            pn["D3"] = "BAD"
            pn["E3"] = "bad seller"
            pn["F3"] = "A1B2C3  Z9"
            pn["C4"] = "ABC123  45"
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)
                self.addCleanup(result.workbook.close)
                self.assertEqual(result.row_results[("PN Level", 2)].check, 2)
                self.assertIn(
                    "Manufacturer COFOR = bad", result.row_results[("PN Level", 2)].comment
                )
                self.assertIn(
                    "Empty return COFOR = wrong", result.row_results[("PN Level", 2)].comment
                )
                self.assertNotIn("Triplet COFOR", result.row_results[("PN Level", 2)].comment)
                self.assertEqual(result.row_results[("PN Level", 3)].check, 3)
                for field in ("Shipper COFOR", "Manufacturer COFOR", "Seller COFOR"):
                    self.assertIn(field, result.row_results[("PN Level", 3)].comment)
                self.assertEqual(result.row_results[("PN Level", 4)].check, 0)
                self.assertEqual(result.rule_totals[("pn_cofor", "Manufacturer COFOR")], 2)
                self.assertEqual(result.rule_totals[("pn_cofor", "Empty return COFOR")], 1)
                self.assertEqual(result.rule_totals[("pn_cofor", "Shipper COFOR")], 1)
                self.assertEqual(result.rule_totals[("pn_cofor", "Seller COFOR")], 1)
                self.assertNotIn(("pn_triplet_cofor", "Triplet COFOR"), result.rule_totals)

    def test_pn_triplet_treats_excel_non_breaking_spaces_as_spaces(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    3: {
                        "Supplier Punch code": "P1",
                        "Triplet COFOR": "A00E0I\u00a0 01_A00E0I  01_A00E0I\u00a0 01",
                        "OPEN TASK": "",
                    },
                    4: {"Triplet COFOR": "OTHER"},
                },
                pn_rows=(("P1", "A00E0I  01_A00E0I\u00a0 01_A00E0I  01"),),
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)

        self.assertEqual(result.row_results[("PN Level", 2)].check, 0)

    def test_pn_membership_retains_sheet_identity_and_combined_history(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(
                directory,
                row_overrides={
                    3: {"Supplier Punch code": "P1", "Triplet COFOR": "A", "OPEN TASK": ""},
                    4: {"Supplier Punch code": "P1", "Triplet COFOR": "B"},
                },
                pn_rows=((" p1 ", " a "), ("P1", "C"), ("unknown", "C"), ("", "A")),
            )
            original_bytes = path.read_bytes()
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(
                    path, connection=connection, analysis_date=date(2026, 7, 15)
                )
                self.addCleanup(result.workbook.close)
                self.assertEqual(
                    result.assessed_rows,
                    (
                        ("Supplier Level", 3),
                        ("Supplier Level", 4),
                        ("PN Level", 2),
                        ("PN Level", 3),
                    ),
                )
                self.assertEqual(result.row_results[("PN Level", 2)].check, 0)
                rejected = result.row_results[("PN Level", 3)]
                self.assertEqual(rejected.check, 1)
                for detail in ("Punch seller", "P1", "Triplet COFOR", "C", "a", "b"):
                    self.assertIn(detail, rejected.comment)
                self.assertEqual(result.rows_failed, 3)
                self.assertEqual(sum(row.check for row in result.row_results.values()), 3)
                self.assertEqual(result.rule_totals[("pn_triplet_cofor", "Triplet COFOR")], 1)
                self.assertEqual(
                    connection.execute(
                        "SELECT rows_total, rows_in_scope, rows_failed, status, exported_file FROM runs"
                    ).fetchone(),
                    (4, 4, 3, "ok", None),
                )
                output = export_edct_result(result, directory / "out")
                self.assertEqual(
                    connection.execute("SELECT exported_file FROM runs").fetchone(), (str(output),)
                )
            exported = load_workbook(output, data_only=False)
            self.addCleanup(exported.close)
            for sheet_name, row in result.assessed_rows:
                sheet = exported[sheet_name]
                headers = [cell.value for cell in sheet[1 if sheet_name == "PN Level" else 2]]
                self.assertEqual(
                    sheet.cell(row, headers.index("Check") + 1).value,
                    result.row_results[(sheet_name, row)].check,
                )
            self.assertEqual(exported["PN Level"]["A2"].value, " p1 ")
            self.assertEqual(exported["PN Level"]["B2"].value, " a ")
            self.assertIsNone(exported["PN Level"]["C4"].value)
            self.assertEqual(path.read_bytes(), original_bytes)
