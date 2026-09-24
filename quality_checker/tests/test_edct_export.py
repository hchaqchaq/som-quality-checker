from __future__ import annotations

import re
import sqlite3
import tempfile
from contextlib import closing
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree
from zipfile import ZipFile

from edct_support import (
    EdctTestCase,
    add_worksheet_extension,
    build_edct_workbook,
    corrupt_first_cell_style,
    remap_archive_parts,
    remap_table_relationship,
    set_formula_caches,
)
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table
from quality_checker.checkers.edct import export as edct_export
from quality_checker.checkers.edct import runner as edct_analysis
from quality_checker.checkers.edct.config import (
    EDCT_COFOR_COLUMNS,
    EDCT_DATE_COLUMNS,
    EDCT_DATED_COMMENT_COLUMNS,
    EDCT_EMAIL_COLUMNS,
    EDCT_FORMULA_COLUMNS,
)
from quality_checker.checkers.edct.export import export_edct_result
from quality_checker.checkers.edct.models import EdctLoadError
from quality_checker.checkers.edct.runner import run_edct_analysis


class EdctExportTests(EdctTestCase):
    def test_pn_annotations_follow_table_despite_distant_formatted_cells(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(directory, pn_rows=(("1003", "bad"),))
            workbook = load_workbook(path)
            pn = workbook["PN Level"]
            pn.add_table(Table(displayName="NativePnData", ref="A1:B2"))
            pn["AZ1"] = "Check"
            pn["BA1"] = "Comment"
            pn["AZ2"] = 99
            pn["BA2"] = "stale"
            workbook.save(path)
            supplier_headers = [cell.value for cell in workbook["Supplier Level"][2]]
            triplet_column = get_column_letter(supplier_headers.index("Triplet COFOR") + 1)
            workbook.close()
            set_formula_caches(
                path,
                {
                    "xl/worksheets/sheet1.xml": {
                        f"{triplet_column}3": "A",
                        f"{triplet_column}4": "B",
                    }
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)
                self.addCleanup(result.workbook.close)
                output = export_edct_result(result, directory / "out")
            exported = load_workbook(output)
            self.addCleanup(exported.close)
            sheet = exported["PN Level"]
            self.assertEqual((sheet["C1"].value, sheet["D1"].value), ("Check", "Comment"))
            self.assertEqual(sheet["C2"].value, result.row_results[("PN Level", 2)].check)
            self.assertEqual(sheet["D2"].value, result.row_results[("PN Level", 2)].comment)
            self.assertIsNone(sheet["AZ1"].value)
            self.assertIsNone(sheet["BA2"].value)
            self.assertEqual(sheet.tables["NativePnData"].ref, "A1:B2")

    def test_pn_formula_caches_tables_and_archive_parts_survive_export(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(
                directory,
                row_overrides={
                    3: {"Supplier Punch code": "P1", "OPEN TASK": ""},
                    4: {"Supplier Punch code": "P1"},
                },
                pn_rows=(("P1", "=A2"), ("P1", "=A3"), ("P1", "=A4")),
            )
            workbook = load_workbook(path)
            pn = workbook["PN Level"]
            pn.add_table(Table(displayName="NativePnData", ref="A1:B4"))
            pn["B2"].fill = PatternFill("solid", fgColor="AABBCC")
            pn["A2"].hyperlink = "https://example.com/supplier"
            supplier_headers = [cell.value for cell in workbook["Supplier Level"][2]]
            triplet_column = get_column_letter(supplier_headers.index("Triplet COFOR") + 1)
            workbook.save(path)
            workbook.close()
            add_worksheet_extension(path, "{PN-PRESERVATION}", "xl/worksheets/sheet5.xml")
            set_formula_caches(
                path,
                {
                    "xl/worksheets/sheet1.xml": {
                        f"{triplet_column}3": "A",
                        f"{triplet_column}4": "B",
                    },
                    "xl/worksheets/sheet5.xml": {"B2": " a ", "B3": "C", "B4": ""},
                },
            )
            remap_table_relationship(path, "rId99")
            remap_archive_parts(path)
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)
                self.addCleanup(result.workbook.close)
                self.assertEqual(result.rows_failed, 3)
                self.assertEqual(result.pn_values[2], ("P1", " a "))
                self.assertEqual(result.row_results[("PN Level", 4)].check, 1)
                output = export_edct_result(result, directory / "out")
                rerun = run_edct_analysis(output, connection=connection)
                self.addCleanup(rerun.workbook.close)
                self.assertEqual(rerun.row_results, result.row_results)
            with ZipFile(path) as source, ZipFile(output) as exported:
                self.assertEqual(set(source.namelist()), set(exported.namelist()))
                changed = {
                    "xl/worksheets/sheet1.xml",
                    "xl/worksheets/sheet42.xml",
                    "xl/tables/table27.xml",
                    "xl/worksheets/sheet3.xml",
                    "xl/styles.xml",
                }
                for name in source.namelist():
                    if name not in changed:
                        self.assertEqual(source.read(name), exported.read(name), name)
                self.assertIn(b"{PN-PRESERVATION}", exported.read("xl/worksheets/sheet42.xml"))
                worksheet_xml = exported.read("xl/worksheets/sheet42.xml")
                root_start = worksheet_xml[
                    worksheet_xml.index(
                        b"<", worksheet_xml.index(b"<?xml") + 5
                    ) : worksheet_xml.index(b">", worksheet_xml.index(b"<?xml") + 5) + 1
                ]
                ignorable = re.search(rb'Ignorable="([^"]+)"', root_start)
                if ignorable is not None:
                    for prefix in ignorable.group(1).split():
                        self.assertIn(b"xmlns:" + prefix + b"=", root_start)
            exported = load_workbook(output, data_only=False)
            self.addCleanup(exported.close)
            for sheet_name, header_row, source_row, expected_check in (
                ("Supplier Level", 2, 3, 0),
                ("PN Level", 1, 3, 1),
            ):
                sheet = exported[sheet_name]
                headers = [cell.value for cell in sheet[header_row]]
                self.assertIn("Check", headers)
                self.assertEqual(
                    sheet.cell(source_row, headers.index("Check") + 1).value, expected_check
                )
            self.assertEqual(exported["PN Level"]["B2"].value, "=A2")
            self.assertEqual(exported["PN Level"]["B2"].fill.fgColor.rgb, "00AABBCC")
            self.assertEqual(
                exported["PN Level"]["A2"].hyperlink.target, "https://example.com/supplier"
            )
            self.assertEqual(exported["PN Level"].tables["NativePnData"].ref, "A1:B4")
            cached = load_workbook(output, data_only=True)
            self.addCleanup(cached.close)
            self.assertEqual(cached["PN Level"]["B2"].value, " a ")

    def test_default_history_records_export(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            with (
                patch.object(edct_analysis, "DB_PATH", directory / "history.db"),
                patch.object(edct_export, "DB_PATH", directory / "history.db"),
            ):
                result = run_edct_analysis(input_path, analysis_date=date(2026, 7, 15))
                output = export_edct_result(result, directory / "out")
            with closing(sqlite3.connect(directory / "history.db")) as connection:
                self.assertEqual(
                    connection.execute("SELECT exported_file, status FROM runs").fetchone(),
                    (str(output), "ok"),
                )
            result.workbook.close()

    def test_complete_workbook_is_annotated_and_preserved(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            with closing(sqlite3.connect(":memory:")) as connection:
                connection.row_factory = sqlite3.Row
                result = run_edct_analysis(input_path, connection=connection)
                output_path = export_edct_result(result, directory)
                runs = connection.execute("SELECT project, rows_total FROM runs").fetchall()

            exported = load_workbook(output_path, data_only=False)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            self.assertEqual(
                result.assessed_rows,
                (("Supplier Level", 3), ("Supplier Level", 4), ("Template-Cofor-Creation", 3)),
            )
            self.assertEqual([row["project"] for row in runs], ["eDCT"])
            self.assertEqual([row["rows_total"] for row in runs], [3])
            self.assertEqual(
                exported.sheetnames,
                [
                    "Supplier Level",
                    "Open Task",
                    "Template-Cofor-Creation",
                    "Other Sheet",
                    "PN Level",
                ],
            )
            self.assertEqual(exported["Other Sheet"]["A1"].value, "keep me")
            self.assertEqual(supplier["A3"].fill.fgColor.rgb, "00FF0000")
            self.assertIn("Check", headers)
            self.assertIn("Comment", headers)
            self.assertEqual(supplier.cell(3, headers.index("Check") + 1).value, 0)
            self.assertTrue(
                supplier.tables["Tabella2"].ref.endswith(get_column_letter(len(headers)) + "5")
            )
            self.assertRegex(output_path.name, r"input_eDCT_checked_\d{8}_\d{6}\.xlsx")
            original = load_workbook(input_path)
            self.assertEqual(original["Other Sheet"]["A1"].value, "keep me")
            original.close()
            exported.close()

    def test_export_copies_unannotated_sheets_from_source(self) -> None:
        class UnserializableValue:
            def __str__(self) -> str:
                raise RuntimeError("unannotated sheet was serialized")

        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(
                    input_path,
                    connection=connection,
                    settings=edct_analysis.EdctHeaderSettings.default(),
                )
                self.addCleanup(result.workbook.close)
                unrelated_cell = result.workbook["Other Sheet"]["A1"]
                unrelated_cell._value = UnserializableValue()
                unrelated_cell.data_type = "s"
                output_path = export_edct_result(result, directory)
                self.assertIn("Other Sheet", result.workbook.sheetnames)
                self.assertIsInstance(
                    result.workbook["Other Sheet"]["A1"].value, UnserializableValue
                )

            exported = load_workbook(output_path, read_only=True)
            self.assertEqual(exported["Other Sheet"]["A1"].value, "keep me")
            exported.close()
            result.workbook.close()

    def test_optional_phone_and_shipping_columns_are_ignored_and_preserved(self) -> None:
        optional_columns = ("Phone", "Phone2", "Shipping location")
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(
                directory,
                optional_columns=optional_columns,
                row_overrides={
                    3: {
                        "Phone": "not a phone number",
                        "Phone2": 123,
                        "Shipping location": "arbitrary value",
                    }
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(input_path, connection=connection)
                output_path = export_edct_result(result, directory)

            self.assertEqual(result.row_results[("Supplier Level", 3)].check, 0)
            self.assertFalse(any(rule == "phone" for rule, _ in result.rule_totals))
            self.assertFalse(any(rule == "shipping_location" for rule, _ in result.rule_totals))
            exported = load_workbook(output_path)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            for column, expected in (
                ("Phone", "not a phone number"),
                ("Phone2", 123),
                ("Shipping location", "arbitrary value"),
            ):
                self.assertEqual(supplier.cell(3, headers.index(column) + 1).value, expected)
            exported.close()

    def test_export_formats_native_dates_as_day_month_year(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(
                directory,
                row_overrides={3: {"First communication sent": date(2026, 9, 2)}},
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(input_path, connection=connection)
                output_path = export_edct_result(result, directory)

            exported = load_workbook(output_path)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            date_cell = supplier.cell(3, headers.index("First communication sent") + 1)
            self.assertEqual(date_cell.value, datetime(2026, 9, 2))
            self.assertEqual(date_cell.number_format, "DD/MM/YYYY")
            exported.close()

    def test_export_includes_styles_created_during_annotation(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(input_path, connection=connection)
                worksheet = result.workbook["Supplier Level"]
                style_source = worksheet.max_column
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row, style_source).fill = PatternFill(
                        "solid",
                        fgColor="00FF00",
                    )
                output_path = export_edct_result(result, directory)

            exported = load_workbook(output_path)
            exported.close()

    def test_export_preserves_source_table_relationship_ids(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            remap_table_relationship(input_path, "rId99")

            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(input_path, connection=connection)
                output_path = export_edct_result(result, directory)

            exported = load_workbook(output_path)
            self.assertIn("Tabella2", exported["Supplier Level"].tables)
            exported.close()

    def test_rerun_clears_results_for_rows_that_are_no_longer_assessed(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(
                directory,
                row_overrides={3: {"Phone": "invalid"}},
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                first_result = run_edct_analysis(input_path, connection=connection)
                first_output = export_edct_result(first_result, directory)

                workbook = load_workbook(first_output)
                supplier = workbook["Supplier Level"]
                headers = [cell.value for cell in supplier[2]]
                supplier.cell(3, headers.index("Index") + 1).value = None
                rerun_input = directory / "rerun.xlsx"
                workbook.save(rerun_input)
                workbook.close()

                second_result = run_edct_analysis(rerun_input, connection=connection)
                second_output = export_edct_result(second_result, directory)

            exported = load_workbook(second_output)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            self.assertIsNone(supplier.cell(3, headers.index("Check") + 1).value)
            self.assertIsNone(supplier.cell(3, headers.index("Comment") + 1).value)
            exported.close()

    def test_existing_result_columns_are_brought_inside_the_result_table(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            workbook = load_workbook(input_path)
            supplier = workbook["Supplier Level"]
            supplier.cell(2, supplier.max_column + 1, "Check")
            supplier.cell(2, supplier.max_column + 1, "Comment")
            workbook.save(input_path)
            workbook.close()

            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(input_path, connection=connection)
                output = export_edct_result(result, directory)

            exported = load_workbook(output)
            table = exported["Supplier Level"].tables["Tabella2"]
            min_col, _, max_col, _ = range_boundaries(table.ref)
            self.assertEqual(len(table.tableColumns), max_col - min_col + 1)
            self.assertEqual(
                [column.name for column in table.tableColumns[-2:]],
                ["Check", "Comment"],
            )
            exported.close()

    def test_export_preserves_unsupported_worksheet_extensions(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            input_path = build_edct_workbook(directory)
            marker = "{EDCT-PRESERVATION-TEST}"
            add_worksheet_extension(input_path, marker)

            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(input_path, connection=connection)
                output = export_edct_result(result, directory)

            with ZipFile(input_path) as source, ZipFile(output) as exported:
                self.assertEqual(set(exported.namelist()), set(source.namelist()))
                for name in source.namelist():
                    if name not in {
                        "xl/worksheets/sheet1.xml",
                        "xl/worksheets/sheet5.xml",
                        "xl/worksheets/sheet3.xml",
                        "xl/tables/table1.xml",
                        "xl/styles.xml",
                    }:
                        self.assertEqual(exported.read(name), source.read(name), name)
                source_xml = source.read("xl/worksheets/sheet1.xml")
                exported_xml = exported.read("xl/worksheets/sheet1.xml")

            def extension_payload(xml: bytes) -> tuple:
                root = ElementTree.fromstring(xml)
                extension = next(
                    element
                    for element in root.iter()
                    if element.tag.rsplit("}", 1)[-1] == "ext"
                    and element.attrib.get("uri") == marker
                )

                def signature(element) -> tuple:
                    return (
                        element.tag,
                        tuple(sorted(element.attrib.items())),
                        (element.text or "").strip(),
                        tuple(signature(child) for child in element),
                    )

                return signature(extension)

            self.assertEqual(extension_payload(exported_xml), extension_payload(source_xml))

    def test_export_failure_marks_the_history_run_failed(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(directory)
            blocked_output = directory / "not-a-directory"
            blocked_output.write_text("file", encoding="utf-8")
            connection = sqlite3.connect(":memory:")
            connection.row_factory = sqlite3.Row
            try:
                result = run_edct_analysis(path, connection=connection)
                with self.assertRaises(OSError):
                    export_edct_result(result, blocked_output)
                run = connection.execute(
                    "SELECT status, error_message FROM runs WHERE id = ?",
                    (result.run_id,),
                ).fetchone()
            finally:
                connection.close()

        self.assertEqual(run["status"], "failed")
        self.assertTrue(run["error_message"])

    def test_corrupt_analysis_workbook_is_removed_and_run_is_failed(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(directory)
            connection = sqlite3.connect(":memory:")
            connection.row_factory = sqlite3.Row
            try:
                result = run_edct_analysis(
                    path,
                    connection=connection,
                    settings=edct_analysis.EdctHeaderSettings.default(),
                )
                original_merge = edct_export._merge_annotated_parts

                def merge_then_corrupt(*args, **kwargs) -> None:
                    original_merge(*args, **kwargs)
                    corrupt_first_cell_style(args[2])

                with (
                    patch.object(edct_export, "_merge_annotated_parts", merge_then_corrupt),
                    self.assertRaisesRegex(
                        EdctLoadError,
                        "Analysis workbook validation failed",
                    ),
                ):
                    export_edct_result(result, directory)
                run = connection.execute(
                    "SELECT status, error_message FROM runs WHERE id = ?",
                    (result.run_id,),
                ).fetchone()
            finally:
                connection.close()

            self.assertEqual(list(directory.glob("*_eDCT_checked_*.xlsx")), [])
        self.assertEqual(run["status"], "failed")
        self.assertIn("Analysis workbook validation failed", run["error_message"])

    def test_every_configured_field_rule_reaches_exported_check_and_comment(self) -> None:
        invalid_values = {
            **{column: "invalid email" for column in EDCT_EMAIL_COLUMNS},
            **{column: "invalid cofor" for column in EDCT_COFOR_COLUMNS},
            **{column: "invalid date" for column in EDCT_DATE_COLUMNS},
            **{column: "invalid comment" for column in EDCT_DATED_COMMENT_COLUMNS},
            "Triple Status": "Valid",
            "Overseas": "YES",
            "Supplier Confimation": "YES",
            "eSupplierConnect": "YES",
            "B2B": "YES",
            "New supplier portal": "YES",
            "SPM": "YES",
            "iTMS": "YES",
            "EDI Mode": "WEB EDI",
            "OPEN TASK": "YES",
        }
        expected_columns = (
            *EDCT_EMAIL_COLUMNS,
            *EDCT_COFOR_COLUMNS,
            *EDCT_DATE_COLUMNS,
            *EDCT_DATED_COMMENT_COLUMNS,
            "Supplier Punch code",
        )
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(directory, row_overrides={3: invalid_values})
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(
                    path,
                    connection=connection,
                    analysis_date=date(2026, 7, 30),
                )
                output = export_edct_result(result, directory)

            exported = load_workbook(output)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            check = supplier.cell(3, headers.index("Check") + 1).value
            comment = supplier.cell(3, headers.index("Comment") + 1).value
            exported.close()

        self.assertEqual(check, len(expected_columns))
        for column in expected_columns:
            self.assertIn(f"{column} =", comment)

    def test_every_formula_rule_reaches_exported_check_and_comment(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(
                directory,
                row_overrides={
                    4: {column: "=1" for column in EDCT_FORMULA_COLUMNS},
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)
                output = export_edct_result(result, directory)

            exported = load_workbook(output)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            check = supplier.cell(4, headers.index("Check") + 1).value
            comment = supplier.cell(4, headers.index("Comment") + 1).value
            exported.close()

        self.assertEqual(check, len(EDCT_FORMULA_COLUMNS))
        for column in EDCT_FORMULA_COLUMNS:
            self.assertIn(f"{column} =", comment)
