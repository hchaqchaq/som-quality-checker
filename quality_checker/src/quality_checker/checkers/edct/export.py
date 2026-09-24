from __future__ import annotations

import re
import sqlite3
from collections.abc import Callable
from copy import copy
from datetime import date, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import TableColumn
from openpyxl.worksheet.worksheet import Worksheet

from ...application import DB_PATH
from ...db.repository import (
    initialize_schema,
    open_connection,
    update_run_exported_file,
    update_run_status,
)
from .config import (
    EDCT_COFOR_TEMPLATE_HEADER_ROW,
    EDCT_COFOR_TEMPLATE_SHEET,
    EDCT_DATE_COLUMNS,
    EDCT_HEADER_ROW,
    EDCT_PN_HEADER_ROW,
    EDCT_PN_SHEET,
)
from .models import EdctLoadError, EdctRunResult
from .ooxml import table_part as _table_part
from .ooxml import worksheet_parts as _worksheet_parts
from .workbook import header_map as _header_map
from .workbook import normalized_text as _normalized_text


def _result_columns(worksheet, header_row: int = EDCT_HEADER_ROW) -> tuple[int, int]:
    headers = _header_map(worksheet, header_row)
    if worksheet.title == EDCT_PN_SHEET:
        table_end = max(
            (range_boundaries(table.ref)[2] for table in worksheet.tables.values()),
            default=0,
        )
        data_end = max(
            (column for name, column in headers.items() if name not in ("Check", "Comment")),
            default=0,
        )
        first = max(table_end, data_end) + 1
        for offset, name in enumerate(("Check", "Comment")):
            old_column = headers.get(name)
            if old_column is not None and old_column != first + offset:
                for row in range(header_row, worksheet.max_row + 1):
                    worksheet.cell(row, old_column).value = None
            worksheet.cell(header_row, first + offset, name)
        return first, first + 1
    if worksheet.title == EDCT_COFOR_TEMPLATE_SHEET:
        columns: list[int] = []
        for name in ("Check", "Comment"):
            column = headers.get(name)
            if column is None:
                column = worksheet.max_column + 1
                worksheet.cell(header_row, column, name)
            columns.append(column)
        return columns[0], columns[1]
    table = worksheet.tables.get("Tabella2")
    if table is None:
        raise EdctLoadError("Missing required table: Tabella2")

    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    result_columns: list[int] = []
    for name in ("Check", "Comment"):
        column = headers.get(name)
        if column is None:
            column = max(max_col, worksheet.max_column) + 1
            worksheet.cell(EDCT_HEADER_ROW, column, name)
        if column > max_col:
            for added_column in range(max_col + 1, column + 1):
                header = _normalized_text(worksheet.cell(EDCT_HEADER_ROW, added_column).value)
                if not header:
                    header = f"Column {added_column}"
                    worksheet.cell(EDCT_HEADER_ROW, added_column, header)
                table.tableColumns.append(TableColumn(id=len(table.tableColumns) + 1, name=header))
            max_col = column
        result_columns.append(column)

    table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
    if table.autoFilter is not None:
        table.autoFilter.ref = table.ref
    return result_columns[0], result_columns[1]


def _update_result_history(
    result: EdctRunResult,
    action: Callable[[sqlite3.Connection], None],
) -> None:
    if result.run_id < 0:
        return
    if result.history_connection is not None:
        action(result.history_connection)
        return
    if result.uses_default_db:
        connection = open_connection(DB_PATH)
        try:
            initialize_schema(connection)
            action(connection)
        finally:
            connection.close()


def _preserve_ooxml_extensions(
    source: Path,
    target: Path,
    replacement_parts: dict[str, str],
) -> None:
    def local_name(element) -> str:
        return element.tag.rsplit("}", 1)[-1]

    def restore_extensions(source_element, target_element) -> None:
        source_children: dict[str, list] = {}
        target_children: dict[str, list] = {}
        for child in source_element:
            source_children.setdefault(local_name(child), []).append(child)
        for child in target_element:
            target_children.setdefault(local_name(child), []).append(child)

        for name, children in source_children.items():
            if name == "extLst":
                for existing in target_children.get(name, []):
                    target_element.remove(existing)
                target_element.extend(children)
                continue
            for source_child, target_child in zip(
                children, target_children.get(name, []), strict=False
            ):
                restore_extensions(source_child, target_child)

    def restore_identified_extensions(source_root, target_root) -> None:
        target_elements = list(target_root.iter())
        for source_element in source_root.iter():
            extensions = [child for child in source_element if local_name(child) == "extLst"]
            identity = {
                attribute.rsplit("}", 1)[-1]: value
                for attribute, value in source_element.attrib.items()
                if attribute.rsplit("}", 1)[-1] in {"id", "name"}
            }
            if not extensions or not identity:
                continue
            matches = [
                element
                for element in target_elements
                if local_name(element) == local_name(source_element)
                and all(
                    any(
                        attribute.rsplit("}", 1)[-1] == name and value == expected
                        for attribute, value in element.attrib.items()
                    )
                    for name, expected in identity.items()
                )
            ]
            if len(matches) == 1:
                for existing in list(matches[0]):
                    if local_name(existing) == "extLst":
                        matches[0].remove(existing)
                matches[0].extend(extensions)

    def restore_formula_values(source_root, target_root) -> None:
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        source_cells = {
            cell.attrib["r"]: cell
            for cell in source_root.iter(f"{{{namespace}}}c")
            if cell.find(f"{{{namespace}}}f") is not None
        }
        for cell in target_root.iter(f"{{{namespace}}}c"):
            original = source_cells.get(cell.attrib["r"])
            if original is None or cell.find(f"{{{namespace}}}f") is None:
                continue
            if "t" in original.attrib:
                cell.set("t", original.attrib["t"])
            else:
                cell.attrib.pop("t", None)
            for child in list(cell):
                if local_name(child) in {"f", "v", "is"}:
                    cell.remove(child)
            for child in original:
                if local_name(child) in {"f", "v", "is"}:
                    cell.append(copy(child))

    def preserve_namespace_compatibility(source_xml: bytes, source_root) -> None:
        root_match = re.search(rb"<(?:[A-Za-z_][\w.-]*:)?worksheet\b", source_xml)
        if root_match is None:
            return
        root_end = source_xml.index(b">", root_match.start())
        root_start = source_xml[root_match.start() : root_end + 1].decode("utf-8")
        namespaces = {
            (match.group(1) or ""): match.group(2)
            for match in re.finditer(
                r'xmlns(?::([A-Za-z_][\w.-]*))?="([^"]+)"',
                root_start,
            )
        }
        used_namespaces = {
            name[1:].split("}", 1)[0]
            for element in source_root.iter()
            for name in (element.tag, *element.attrib)
            if name.startswith("{")
        }
        for prefix, namespace in namespaces.items():
            if namespace in used_namespaces and not re.fullmatch(r"ns\d+", prefix):
                ElementTree.register_namespace(prefix, namespace)

        compatibility = "http://schemas.openxmlformats.org/markup-compatibility/2006"
        ignorable_key = f"{{{compatibility}}}Ignorable"
        ignorable = source_root.attrib.get(ignorable_key)
        if ignorable:
            retained = [
                prefix for prefix in ignorable.split() if namespaces.get(prefix) in used_namespaces
            ]
            if retained:
                source_root.set(ignorable_key, " ".join(retained))
            else:
                source_root.attrib.pop(ignorable_key)

    with NamedTemporaryFile(dir=target.parent, suffix=".xlsx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with (
            ZipFile(source) as source_archive,
            ZipFile(target) as target_archive,
            ZipFile(temporary_path, "w", ZIP_DEFLATED) as output_archive,
        ):
            original_parts = {
                generated: original for original, generated in replacement_parts.items()
            }
            for item in target_archive.infolist():
                content = target_archive.read(item.filename)
                if item.filename in original_parts:
                    source_xml = source_archive.read(original_parts[item.filename])
                    source_root = ElementTree.fromstring(source_xml)
                    has_extensions = any(
                        local_name(element) == "extLst" for element in source_root.iter()
                    )
                    is_worksheet = local_name(source_root) == "worksheet"
                    if is_worksheet:
                        preserve_namespace_compatibility(source_xml, source_root)
                    if has_extensions or is_worksheet or local_name(source_root) == "table":
                        target_root = ElementTree.fromstring(content)
                        if has_extensions:
                            restore_extensions(source_root, target_root)
                            restore_identified_extensions(source_root, target_root)
                        if is_worksheet:
                            restore_formula_values(source_root, target_root)
                            annotated_children = {local_name(child): child for child in target_root}
                            for index, child in enumerate(list(source_root)):
                                name = local_name(child)
                                if (
                                    name in {"dimension", "sheetData"}
                                    and name in annotated_children
                                ):
                                    source_root.remove(child)
                                    source_root.insert(index, annotated_children[name])
                            target_root = source_root
                        elif local_name(source_root) == "table":
                            target_root.set("id", source_root.attrib["id"])
                        content = ElementTree.tostring(
                            target_root,
                            encoding="utf-8",
                            xml_declaration=True,
                        )
                output_archive.writestr(item, content)
        temporary_path.replace(target)
    finally:
        temporary_path.unlink(missing_ok=True)


def _merge_annotated_parts(
    source: Path,
    annotated: Path,
    target: Path,
    replacement_parts: dict[str, str],
) -> None:
    with (
        ZipFile(source) as source_archive,
        ZipFile(annotated) as annotated_archive,
        ZipFile(target, "w", ZIP_DEFLATED) as output_archive,
    ):
        missing = set(replacement_parts.values()).difference(annotated_archive.namelist())
        if missing:
            raise EdctLoadError(f"Missing generated workbook parts: {', '.join(sorted(missing))}")
        for item in source_archive.infolist():
            replacement = replacement_parts.get(item.filename)
            content = (
                annotated_archive.read(replacement)
                if replacement is not None
                else source_archive.read(item.filename)
            )
            output_archive.writestr(item, content)


def _save_annotated_worksheets(
    workbook: Workbook,
    worksheets: tuple[Worksheet, ...],
    path: Path,
) -> None:
    original_sheets = workbook._sheets
    original_active_sheet_index = workbook._active_sheet_index
    workbook._sheets = list(worksheets)
    workbook._active_sheet_index = 0
    try:
        workbook.save(path)
    finally:
        workbook._sheets = original_sheets
        workbook._active_sheet_index = original_active_sheet_index


def _validate_analysis_workbook(path: Path) -> None:
    try:
        with path.open("rb") as exported_file:
            workbook = load_workbook(exported_file, read_only=True, data_only=False)
            try:
                for sheet_name in ("Supplier Level", EDCT_PN_SHEET):
                    for row in workbook[sheet_name].iter_rows():
                        for cell in row:
                            _ = cell.value, cell.number_format
            finally:
                workbook.close()
    except Exception as exc:
        path.unlink(missing_ok=True)
        raise EdctLoadError(f"Analysis workbook validation failed: {exc}") from exc


def export_edct_result(
    result: EdctRunResult,
    output_dir: Path | str,
    *,
    progress: Callable[[str], None] | None = None,
) -> Path:
    report = progress or (lambda _message: None)
    try:
        report("Exporting workbook…")
        return _export_edct_result(result, output_dir, progress=report)
    except Exception as error:
        error_message = str(error)

        def mark_export_failed(connection: sqlite3.Connection) -> None:
            update_run_status(connection, result.run_id, "failed", error_message)

        _update_result_history(result, mark_export_failed)
        raise


def _export_edct_result(
    result: EdctRunResult,
    output_dir: Path | str,
    *,
    progress: Callable[[str], None],
) -> Path:
    target_dir = Path(output_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    worksheets = (
        (result.workbook["Supplier Level"], EDCT_HEADER_ROW),
        (result.workbook[EDCT_PN_SHEET], EDCT_PN_HEADER_ROW),
        (result.workbook[EDCT_COFOR_TEMPLATE_SHEET], EDCT_COFOR_TEMPLATE_HEADER_ROW),
    )
    result_columns: dict[str, tuple[int, int]] = {}
    for worksheet, header_row in worksheets:
        check_column, comment_column = _result_columns(worksheet, header_row)
        result_columns[worksheet.title] = (check_column, comment_column)
        style_source = max(1, min(check_column, comment_column) - 1)
        for column in (check_column, comment_column):
            for row in range(header_row, worksheet.max_row + 1):
                source = worksheet.cell(row, style_source)
                target = worksheet.cell(row, column)
                target._style = copy(source._style)
        for row in range(header_row + 1, worksheet.max_row + 1):
            worksheet.cell(row, check_column).value = None
            worksheet.cell(row, comment_column).value = None
        if worksheet.title == "Supplier Level":
            headers = _header_map(worksheet)
            for canonical_column in EDCT_DATE_COLUMNS:
                configured_column = result.supplier_headers.get(canonical_column, canonical_column)
                column_index = headers[configured_column]
                for row in range(header_row + 1, worksheet.max_row + 1):
                    cell = worksheet.cell(row, column_index)
                    if isinstance(cell.value, (date, datetime)):
                        cell.number_format = "DD/MM/YYYY"

    for (sheet_name, row), row_result in result.row_results.items():
        worksheet = result.workbook[sheet_name]
        check_column, comment_column = result_columns[sheet_name]
        worksheet.cell(row, check_column, row_result.check)
        worksheet.cell(row, comment_column, row_result.comment)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_stem = re.sub(r"[^A-Za-z0-9._-]+", "_", result.input_file.stem).strip("._") or "analysis"
    target = target_dir / f"{safe_stem}_eDCT_checked_{timestamp}.xlsx"
    supplier = result.workbook["Supplier Level"]
    table = supplier.tables["Tabella2"]
    with NamedTemporaryFile(dir=target_dir, suffix=".xlsx", delete=False) as temporary:
        annotated = Path(temporary.name)
    try:
        _save_annotated_worksheets(
            result.workbook,
            tuple(worksheet for worksheet, _ in worksheets),
            annotated,
        )
        with ZipFile(result.input_file) as source_archive:
            source_sheets = _worksheet_parts(source_archive)
            source_table = _table_part(
                source_archive, source_sheets["Supplier Level"], table.displayName
            )
            replacement_parts = {
                source_sheets[worksheet.title]: worksheet.path.lstrip("/")
                for worksheet, _ in worksheets
            }
            replacement_parts[source_table] = table.path.lstrip("/")
            replacement_parts["xl/styles.xml"] = "xl/styles.xml"
        _preserve_ooxml_extensions(result.input_file, annotated, replacement_parts)
        _merge_annotated_parts(result.input_file, annotated, target, replacement_parts)
        progress("Verifying export…")
        _validate_analysis_workbook(target)
    except Exception:
        target.unlink(missing_ok=True)
        raise
    finally:
        annotated.unlink(missing_ok=True)

    _update_result_history(
        result,
        lambda connection: update_run_exported_file(connection, result.run_id, str(target)),
    )
    return target
