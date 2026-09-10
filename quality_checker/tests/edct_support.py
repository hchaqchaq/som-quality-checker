from __future__ import annotations

import unittest
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from quality_checker.checkers.edct import runner as edct_analysis
from quality_checker.checkers.edct.config import (
    EDCT_FORMULA_COLUMNS,
    EDCT_REQUIRED_COLUMNS,
)


def build_edct_workbook(
    directory: Path,
    *,
    include_open_task: bool = True,
    row_overrides: dict[int, dict[str, object]] | None = None,
    include_cofor_template: bool = True,
    optional_columns: tuple[str, ...] = (),
    pn_rows: tuple[tuple[object, object], ...] = (),
) -> Path:
    workbook = Workbook()
    all_columns = (*EDCT_REQUIRED_COLUMNS, *optional_columns)
    supplier = workbook.active
    supplier.title = "Supplier Level"
    supplier.append(["metadata"] * len(all_columns))
    supplier.append(list(all_columns))

    base: dict[str, object] = {column: "" for column in all_columns}
    for row_number, index in ((3, "Metz_01"), (4, "Metz_02")):
        values: dict[str, object] = dict(base)
        values.update(
            {
                "Index": index,
                "Supplier Punch code": 1000 + row_number,
                "Supplier name": f"Supplier {row_number}",
                "Effective kick-off date": "",
                "Cofor created date": "",
                "Overseas": "",
                "OPEN TASK": "YES" if row_number == 3 else "",
            }
        )
        for column in EDCT_FORMULA_COLUMNS:
            values[column] = f'=IF(A{row_number}="","",A{row_number})'
        values.update((row_overrides or {}).get(row_number, {}))
        supplier.append([values[column] for column in all_columns])

    supplier.append([""] * len(all_columns))
    supplier["A3"].fill = PatternFill("solid", fgColor="FF0000")
    table = Table(
        displayName="Tabella2",
        ref=f"A2:{get_column_letter(len(all_columns))}5",
    )
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    supplier.add_table(table)

    if include_open_task:
        open_task = workbook.create_sheet("Open Task")
        open_task.append(["metadata"])
        open_task.append(["Punch Code"])
        open_task.append([1003])
    if include_cofor_template:
        cofor_template = workbook.create_sheet("Template-Cofor-Creation")
        cofor_template["D2"] = "Punch Code"
        cofor_template["D3"] = "9999"
    workbook.create_sheet("Other Sheet")["A1"] = "keep me"
    pn = workbook.create_sheet("PN Level")
    pn.append(["Punch seller", "Triplet COFOR"])
    for seller, triplet in pn_rows:
        pn.append([seller, triplet])

    path = directory / "input.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def set_formula_caches(
    path: Path, caches: dict[str, dict[str, str]], *, data_type: str = "str"
) -> None:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    replacement = path.with_suffix(".zip")
    with ZipFile(path) as source, ZipFile(replacement, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename in caches:
                root = ElementTree.fromstring(content)
                for cell in root.iter(f"{{{namespace}}}c"):
                    if cell.attrib["r"] not in caches[item.filename]:
                        continue
                    cell.set("t", data_type)
                    value = cell.find(f"{{{namespace}}}v")
                    if value is None:
                        value = ElementTree.SubElement(cell, f"{{{namespace}}}v")
                    value.text = caches[item.filename][cell.attrib["r"]]
                content = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, content)
    replacement.replace(path)


def remap_archive_parts(path: Path) -> None:
    replacements = {
        "xl/tables/table1.xml": "xl/tables/table27.xml",
        "xl/worksheets/sheet5.xml": "xl/worksheets/sheet42.xml",
        "xl/worksheets/_rels/sheet5.xml.rels": "xl/worksheets/_rels/sheet42.xml.rels",
    }
    replacement = path.with_suffix(".zip")
    with ZipFile(path) as source, ZipFile(replacement, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename.endswith((".xml", ".rels")):
                content = content.replace(b"tables/table1.xml", b"tables/table27.xml")
                content = content.replace(b"worksheets/sheet5.xml", b"worksheets/sheet42.xml")
            if item.filename == "xl/tables/table1.xml":
                content = content.replace(b'id="1"', b'id="27"', 1)
            item.filename = replacements.get(item.filename, item.filename)
            target.writestr(item, content)
    replacement.replace(path)


def add_worksheet_extension(
    path: Path, marker: str, worksheet_part: str = "xl/worksheets/sheet1.xml"
) -> None:
    replacement = path.with_suffix(".zip")
    extension = (
        f'<extLst><ext uri="{marker}"><test:payload xmlns:test="urn:test">'
        "keep me</test:payload></ext></extLst>"
    ).encode()
    root_namespaces = (
        b' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
        b' mc:Ignorable="x14ac xr"'
        b' xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac"'
        b' xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision"'
        b' xr:uid="{97570D38-0785-4F6B-AB88-F8FBC149D0F1}"'
    )
    with ZipFile(path) as source, ZipFile(replacement, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == worksheet_part:
                root_end = content.index(b">", content.index(b"<worksheet"))
                content = content[:root_end] + root_namespaces + content[root_end:]
                content = content.replace(
                    b"<sheetFormatPr ",
                    b'<sheetFormatPr x14ac:dyDescent="0.25" ',
                    1,
                )
                content = content.replace(b"</worksheet>", extension + b"</worksheet>")
            target.writestr(item, content)
    replacement.replace(path)


def corrupt_first_cell_style(path: Path) -> None:
    replacement = path.with_suffix(".zip")
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with ZipFile(path) as source, ZipFile(replacement, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(content)
                cell = root.find(f".//{{{namespace}}}c")
                if cell is None:
                    raise AssertionError("Generated worksheet has no cells to corrupt")
                cell.set("s", "999999")
                content = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, content)
    replacement.replace(path)


def remap_table_relationship(path: Path, relationship_id: str) -> None:
    replacement = path.with_suffix(".zip")
    relationship_attribute = (
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    )
    with ZipFile(path) as source, ZipFile(replacement, "w", ZIP_DEFLATED) as target:
        relationship_root = ElementTree.fromstring(
            source.read("xl/worksheets/_rels/sheet1.xml.rels")
        )
        table_relationship = next(
            relationship
            for relationship in relationship_root
            if relationship.attrib["Type"].endswith("/table")
        )
        original_id = table_relationship.attrib["Id"]
        table_relationship.set("Id", relationship_id)

        worksheet_root = ElementTree.fromstring(source.read("xl/worksheets/sheet1.xml"))
        table_part = next(
            element
            for element in worksheet_root.iter()
            if element.tag.rsplit("}", 1)[-1] == "tablePart"
            and element.attrib[relationship_attribute] == original_id
        )
        table_part.set(relationship_attribute, relationship_id)

        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                content = ElementTree.tostring(
                    relationship_root,
                    encoding="utf-8",
                    xml_declaration=True,
                )
            elif item.filename == "xl/worksheets/sheet1.xml":
                content = ElementTree.tostring(
                    worksheet_root,
                    encoding="utf-8",
                    xml_declaration=True,
                )
            target.writestr(item, content)
    replacement.replace(path)


class EdctTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls._settings_patcher = patch.object(
            edct_analysis,
            "load_edct_settings",
            return_value=edct_analysis.EdctHeaderSettings.default(),
        )
        cls._settings_patcher.start()

    @classmethod
    def tearDownClass(cls) -> None:
        cls._settings_patcher.stop()
