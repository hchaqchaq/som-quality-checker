from __future__ import annotations

import re
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .config import (
    EDCT_COFOR_TEMPLATE_HEADER_ROW,
    EDCT_COFOR_TEMPLATE_PUNCH_HEADER,
    EDCT_COFOR_TEMPLATE_SHEET,
    EDCT_HEADER_ROW,
    EDCT_INDEX_COLUMN,
    EDCT_PN_HEADER_ROW,
    EDCT_PN_REQUIRED_COLUMNS,
    EDCT_PN_SHEET,
    EDCT_REQUIRED_COLUMNS,
    EDCT_REQUIRED_SHEETS,
)
from .models import EdctLoadError
from .settings import HEADER_ALIASES, EdctHeaderSettings

OPEN_TASK_SHEET = "Open Task"
SUPPLIER_LEVEL_SHEET = "Supplier Level"
OPEN_TASK_PUNCH_HEADER = "Punch Code"
_HEADER_ROWS = {
    SUPPLIER_LEVEL_SHEET: EDCT_HEADER_ROW,
    EDCT_PN_SHEET: EDCT_PN_HEADER_ROW,
    OPEN_TASK_SHEET: EDCT_HEADER_ROW,
    EDCT_COFOR_TEMPLATE_SHEET: EDCT_COFOR_TEMPLATE_HEADER_ROW,
}
_REQUIRED_BY_SHEET = {
    SUPPLIER_LEVEL_SHEET: EDCT_REQUIRED_COLUMNS,
    EDCT_PN_SHEET: EDCT_PN_REQUIRED_COLUMNS,
    OPEN_TASK_SHEET: (OPEN_TASK_PUNCH_HEADER,),
    EDCT_COFOR_TEMPLATE_SHEET: (EDCT_COFOR_TEMPLATE_PUNCH_HEADER,),
}


def normalized_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def normalized_choice(value: object) -> str:
    return normalized_text(value).casefold()


def normalized_punch(value: object) -> str:
    text = normalized_text(value)
    return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


def header_map(worksheet: Worksheet, row: int = EDCT_HEADER_ROW) -> dict[str, int]:
    return {
        normalized_text(cell.value): cell.column
        for cell in worksheet[row]
        if normalized_text(cell.value)
    }


@dataclass(frozen=True, slots=True)
class EdctInspectionResult:
    resolved_headers: dict[str, dict[str, str]]
    alias_matches: tuple[tuple[str, str], ...] = ()

    @property
    def ready(self) -> bool:
        return True

    @property
    def validated_sheets(self) -> tuple[str, ...]:
        return tuple(self.resolved_headers)

    def resolved_settings(self) -> EdctHeaderSettings:
        return EdctHeaderSettings.from_dict(self.resolved_headers)


def _row_headers(worksheet: Worksheet, row: int) -> list[tuple[str, int]]:
    return [
        (str(cell.value), cell.column)
        for cell in worksheet[row]
        if cell.value is not None and str(cell.value).strip()
    ]


def _accepted_names(settings: EdctHeaderSettings, sheet: str, canonical: str) -> set[str]:
    values = {
        canonical,
        settings.get_header(sheet, canonical),
        *HEADER_ALIASES.get((sheet, canonical), ()),
    }
    return {value.strip().casefold() for value in values if value.strip()}


def _resolve_workbook_structure(
    workbook: Workbook,
    settings: EdctHeaderSettings,
) -> EdctInspectionResult:
    settings_errors = settings.validate()
    if settings_errors:
        raise EdctLoadError(f"Workbook rejected — invalid header settings: {settings_errors[0]}")
    missing_sheets = [sheet for sheet in EDCT_REQUIRED_SHEETS if sheet not in workbook.sheetnames]
    missing_columns: list[str] = []
    ambiguous: list[str] = []
    resolved: dict[str, dict[str, str]] = {}
    alias_matches: list[tuple[str, str]] = []
    for sheet, required_columns in _REQUIRED_BY_SHEET.items():
        if sheet not in workbook.sheetnames:
            continue
        physical_headers = _row_headers(workbook[sheet], _HEADER_ROWS[sheet])
        sheet_resolved: dict[str, str] = {}
        for canonical in required_columns:
            accepted = _accepted_names(settings, sheet, canonical)
            matches = [
                (name, column)
                for name, column in physical_headers
                if name.strip().casefold() in accepted
            ]
            if (
                canonical == EDCT_INDEX_COLUMN
                and settings.get_header(sheet, canonical) == canonical
            ):
                index_matches = [
                    match for match in matches if match[0].strip().casefold() == "index"
                ]
                if index_matches:
                    matches = index_matches
            if not matches:
                missing_columns.append(f"{sheet}.{canonical} (row {_HEADER_ROWS[sheet]})")
            elif len(matches) > 1:
                names = ", ".join(repr(name) for name, _ in matches)
                ambiguous.append(f"{sheet}.{canonical} (row {_HEADER_ROWS[sheet]}) matches {names}")
            else:
                physical = matches[0][0].strip()
                sheet_resolved[canonical] = physical
                alias_names = {
                    alias.strip().casefold() for alias in HEADER_ALIASES.get((sheet, canonical), ())
                }
                canonical_names = {
                    canonical.strip().casefold(),
                    settings.get_header(sheet, canonical).strip().casefold(),
                }
                if physical.casefold() in alias_names.difference(canonical_names):
                    alias_matches.append((sheet, canonical))
        resolved[sheet] = sheet_resolved

    if missing_sheets or missing_columns or ambiguous:
        parts: list[str] = []
        if missing_sheets:
            parts.append(f"missing worksheets: {', '.join(missing_sheets)}")
        if missing_columns:
            parts.append(f"missing required headers; columns: {'; '.join(missing_columns)}")
        if ambiguous:
            parts.append(f"ambiguous headers: {'; '.join(ambiguous)}")
        raise EdctLoadError("Workbook rejected — " + " | ".join(parts))
    return EdctInspectionResult(
        resolved_headers=resolved,
        alias_matches=tuple(alias_matches),
    )


def inspect_edct_workbook(
    input_path: Path | str,
    settings: EdctHeaderSettings,
    progress: Callable[[str], None] | None = None,
) -> EdctInspectionResult:
    report = progress or (lambda _message: None)
    path = Path(input_path)
    if not path.exists():
        raise EdctLoadError(f"Workbook rejected — input file not found: {path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise EdctLoadError(
            "Workbook rejected — workbook is unreadable or password-protected"
        ) from exc
    report("Checking worksheet structure…")
    try:
        report("Checking required headers…")
        return _resolve_workbook_structure(workbook, settings)
    finally:
        workbook.close()


@dataclass(slots=True)
class LoadedEdctWorkbook:
    path: Path
    workbook: Workbook
    settings: EdctHeaderSettings
    supplier_headers: dict[str, int]
    pn_headers: dict[str, int]
    supplier_rows: tuple[int, ...]
    supplier_index_header: str
    cofor_template_column: int

    def supplier_header(self, canonical: str) -> str:
        return self.settings.get_header(SUPPLIER_LEVEL_SHEET, canonical)

    def supplier_value(self, row: int, canonical: str) -> object:
        configured = self.supplier_header(canonical)
        return (
            self.workbook[SUPPLIER_LEVEL_SHEET].cell(row, self.supplier_headers[configured]).value
        )

    def open_task_punches(self) -> set[str]:
        worksheet = self.workbook[OPEN_TASK_SHEET]
        configured = self.settings.get_header(OPEN_TASK_SHEET, OPEN_TASK_PUNCH_HEADER)
        column = header_map(worksheet, EDCT_HEADER_ROW)[configured]
        return {
            punch
            for row in range(EDCT_HEADER_ROW + 1, worksheet.max_row + 1)
            if (punch := normalized_punch(worksheet.cell(row, column).value))
        }

    def cofor_template_punches(self) -> set[str]:
        worksheet = self.workbook[EDCT_COFOR_TEMPLATE_SHEET]
        return {
            punch
            for row in range(EDCT_COFOR_TEMPLATE_HEADER_ROW + 1, worksheet.max_row + 1)
            if (punch := normalized_choice(worksheet.cell(row, self.cofor_template_column).value))
        }

    def close(self) -> None:
        self.workbook.close()


def load_edct_workbook(
    input_path: Path | str,
    settings: EdctHeaderSettings,
) -> LoadedEdctWorkbook:
    path = Path(input_path)
    if not path.exists():
        raise EdctLoadError(f"Input file not found: {path}")

    workbook = load_workbook(path, data_only=False)
    try:
        inspection = _resolve_workbook_structure(workbook, settings)
        resolved_settings = inspection.resolved_settings()
        supplier_headers = header_map(workbook[SUPPLIER_LEVEL_SHEET], EDCT_HEADER_ROW)
        pn_headers = header_map(workbook[EDCT_PN_SHEET], EDCT_PN_HEADER_ROW)
        supplier_index_header = resolved_settings.get_header(
            SUPPLIER_LEVEL_SHEET, EDCT_INDEX_COLUMN
        )
        template_header = resolved_settings.get_header(
            EDCT_COFOR_TEMPLATE_SHEET, EDCT_COFOR_TEMPLATE_PUNCH_HEADER
        )
        cofor_template_column = header_map(
            workbook[EDCT_COFOR_TEMPLATE_SHEET], EDCT_COFOR_TEMPLATE_HEADER_ROW
        )[template_header]
        supplier = workbook[SUPPLIER_LEVEL_SHEET]
        supplier_rows = tuple(
            row
            for row in range(EDCT_HEADER_ROW + 1, supplier.max_row + 1)
            if normalized_text(supplier.cell(row, supplier_headers[supplier_index_header]).value)
        )
        return LoadedEdctWorkbook(
            path=path,
            workbook=workbook,
            settings=resolved_settings,
            supplier_headers=supplier_headers,
            pn_headers=pn_headers,
            supplier_rows=supplier_rows,
            supplier_index_header=supplier_index_header,
            cofor_template_column=cofor_template_column,
        )
    except Exception:
        workbook.close()
        raise
