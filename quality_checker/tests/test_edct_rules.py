from __future__ import annotations

import re
import sqlite3
import tempfile
from contextlib import closing
from datetime import date
from pathlib import Path

from edct_support import (
    EdctTestCase,
    build_edct_workbook,
)
from openpyxl import load_workbook
from quality_checker.checkers.edct.config import (
    EDCT_EDI_MODE_VALUES,
    EDCT_FORMULA_COLUMNS,
    EDCT_PN_REQUIRED_COLUMNS,
    EDCT_PORTAL_COLUMNS,
    EDCT_PORTAL_VALUES,
    EDCT_REQUIRED_COLUMNS,
    EDCT_RULE_CATALOGUE_ROWS,
    EDCT_TRIPLE_STATUS_VALUES,
    EDCT_UNCHECKED_COLUMNS,
)
from quality_checker.checkers.edct.export import export_edct_result
from quality_checker.checkers.edct.runner import run_edct_analysis


class EdctRuleTests(EdctTestCase):
    def test_rule_documentation_lists_every_required_edct_column(self) -> None:
        documentation = (
            Path(__file__).resolve().parents[2] / "docs" / "EDCT_QUALITY_CHECKER.md"
        ).read_text(encoding="utf-8")
        inventory = documentation.split("## Configured column inventory", 1)[1].split(
            "## Formula rules", 1
        )[0]
        documented = set(re.findall(r"^- `([^`]+)`$", inventory, re.MULTILINE))
        self.assertEqual(documented, set(EDCT_REQUIRED_COLUMNS) | set(EDCT_PN_REQUIRED_COLUMNS))

    def test_rule_documentation_matches_formula_and_unchecked_configuration(self) -> None:
        documentation = (
            Path(__file__).resolve().parents[2] / "docs" / "EDCT_QUALITY_CHECKER.md"
        ).read_text(encoding="utf-8")
        formula_section = documentation.split("## Formula rules", 1)[1].split("## Field rules", 1)[
            0
        ]
        unchecked_section = documentation.split("## Explicitly unchecked fields", 1)[1].split(
            "## Verified sample smoke", 1
        )[0]
        self.assertEqual(
            set(re.findall(r"^- `([^`]+)`$", formula_section, re.MULTILINE)),
            set(EDCT_FORMULA_COLUMNS),
        )
        self.assertEqual(
            set(re.findall(r"^- `([^`]+)`$", unchecked_section, re.MULTILINE)),
            set(EDCT_UNCHECKED_COLUMNS),
        )
        for configured_value in (
            *EDCT_TRIPLE_STATUS_VALUES,
            *EDCT_PORTAL_VALUES,
            *EDCT_EDI_MODE_VALUES,
        ):
            self.assertIn(f"`{configured_value}`", documentation)
        field_rules = documentation.split("## Field rules", 1)[1].split(
            "## Explicitly unchecked fields", 1
        )[0]
        documented_rows = tuple(line for line in field_rules.splitlines() if line.startswith("| `"))

        def cells(line: str) -> tuple[str, ...]:
            return tuple(cell.strip() for cell in line.strip("|").split("|"))

        self.assertEqual(
            tuple(map(cells, documented_rows)),
            tuple(map(cells, EDCT_RULE_CATALOGUE_ROWS)),
        )

    def test_field_formats_are_reported_with_exact_column_names(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    3: {
                        "Effective kick-off date": "30/07/2026",
                        "eSupplierConnect": "NOT",
                        "B2B": "NOT",
                        "New supplier portal": "NOT",
                        "SPM": "NOT",
                        "iTMS": "NOT",
                        "OPEN TASK": "YES",
                        "Sales contact": "first@example.com, second@example.com",
                        "Seller COFOR": "ABC",
                        "First communication sent": "30/07/2026",
                        "Readiness Comments": "30/07/2026: contacted",
                        "Participants": "Name - person@example.com",
                    }
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(
                    path, connection=connection, analysis_date=date(2026, 7, 30)
                )

        self.assertEqual(result.row_results[("Supplier Level", 3)].check, 3)
        for column in (
            "Sales contact",
            "Seller COFOR",
            "Participants",
        ):
            self.assertIn(column, result.row_results[("Supplier Level", 3)].comment)
        self.assertIn(
            "first@example.com, second@example.com",
            result.row_results[("Supplier Level", 3)].comment,
        )
        self.assertIn(
            "Name - person@example.com", result.row_results[("Supplier Level", 3)].comment
        )

    def test_cofor_separator_accepts_excel_non_breaking_spaces(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    3: {
                        "Effective kick-off date": "30/07/2026",
                        "eSupplierConnect": "NOT",
                        "B2B": "NOT",
                        "New supplier portal": "NOT",
                        "SPM": "NOT",
                        "iTMS": "NOT",
                        "Seller COFOR": "A00E0I\u00a0 01",
                        "Manufacturer COFOR": "A00LYI \u00a001",
                        "Shipper COFOR": "A001EE \u00a006",
                    },
                    4: {
                        "Effective kick-off date": "30/07/2026",
                        "eSupplierConnect": "NOT",
                        "B2B": "NOT",
                        "New supplier portal": "NOT",
                        "SPM": "NOT",
                        "iTMS": "NOT",
                        "Seller COFOR": "A02FYY 01",
                    },
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(
                    path, connection=connection, analysis_date=date(2026, 7, 30)
                )

        self.assertEqual(result.row_results[("Supplier Level", 3)].check, 0)
        self.assertEqual(result.row_results[("Supplier Level", 4)].check, 1)
        self.assertIn(
            "Invalid COFOR format: Seller COFOR = A02FYY 01",
            result.row_results[("Supplier Level", 4)].comment,
        )

    def test_creation_of_cofors_request_date_accepts_only_supported_dates(self) -> None:
        scenarios = (
            ("empty", "", "unmatched", 0),
            ("native date", date(2026, 8, 10), "9999", 0),
            ("strict text", "10/08/2026", "9999", 0),
            ("future date", "10/08/2030", "9999", 0),
            ("dot date", "10.08.2026", "9999", 1),
            ("impossible date", "31/02/2026", "9999", 1),
            ("timestamp text", "10/08/2026 12:00", "9999", 1),
        )
        for label, request_date, supplier_punch, expected_check in scenarios:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(
                    Path(temp),
                    row_overrides={
                        3: {
                            "Supplier Punch code": supplier_punch,
                            "OPEN TASK": "",
                            "Creation of Cofors request date": request_date,
                        }
                    },
                )
                with closing(sqlite3.connect(":memory:")) as connection:
                    result = run_edct_analysis(path, connection=connection)

                self.assertEqual(result.row_results[("Supplier Level", 3)].check, expected_check)

    def test_request_date_requires_supplier_punch_in_cofor_template(self) -> None:
        scenarios = (
            ("matching punch", " 1003 ", 0),
            ("absent punch", "9999", 1),
        )
        for label, template_punch, expected_check in scenarios:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(
                    Path(temp),
                    row_overrides={
                        3: {
                            "Supplier Punch code": "1003",
                            "Creation of Cofors request date": "10/08/2026",
                        }
                    },
                )
                workbook = load_workbook(path)
                workbook["Template-Cofor-Creation"]["D3"] = template_punch
                workbook.save(path)
                workbook.close()

                with closing(sqlite3.connect(":memory:")) as connection:
                    result = run_edct_analysis(path, connection=connection)

                self.assertEqual(result.row_results[("Supplier Level", 3)].check, expected_check)
                if expected_check:
                    self.assertEqual(
                        result.rule_totals[("cofor_template", "Supplier Punch code")],
                        1,
                    )

    def test_matching_cofor_template_punch_requires_request_date(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(Path(temp))
            workbook = load_workbook(path)
            workbook["Template-Cofor-Creation"]["D3"] = " 1003 "
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)

            self.assertEqual(result.row_results[("Supplier Level", 3)].check, 1)
            self.assertEqual(
                result.rule_totals[("date_required", "Creation of Cofors request date")],
                1,
            )
            self.assertEqual(result.row_results[("Supplier Level", 4)].check, 0)

    def test_cofor_template_matching_is_normalized_but_not_fuzzy(self) -> None:
        scenarios = (
            ("case and outer whitespace", " Ab- 01 ", "aB- 01", 1),
            ("partial", "ABC123", "ABC", 0),
            ("leading zero", "00123", "123", 0),
            ("internal space", "AB 123", "AB123", 0),
            ("punctuation", "AB-123", "AB123", 0),
            ("numeric representation", "123.0", "123", 0),
        )
        for label, reference_punch, supplier_punch, expected_check in scenarios:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(
                    Path(temp),
                    row_overrides={
                        3: {
                            "Supplier Punch code": supplier_punch,
                            "OPEN TASK": "",
                        }
                    },
                )
                workbook = load_workbook(path)
                workbook["Template-Cofor-Creation"]["D3"] = reference_punch
                workbook.save(path)
                workbook.close()

                with closing(sqlite3.connect(":memory:")) as connection:
                    result = run_edct_analysis(path, connection=connection)

                self.assertEqual(result.row_results[("Supplier Level", 3)].check, expected_check)

    def test_email_and_dated_comment_boundaries(self) -> None:
        scenarios = (
            ("single email", {"Participants": "one@example.com"}, 0),
            ("email list", {"Participants": "one@example.com; two@example.org"}, 0),
            ("trailing separators", {"Participants": "one@example.com;;;"}, 0),
            ("spaced separators", {"Participants": "one@example.com ; two@example.org ;"}, 0),
            ("duplicate emails", {"Participants": "one@example.com;one@example.com;"}, 0),
            ("comma-separated email", {"Participants": "one@example.com,two@example.org"}, 1),
            ("display-name email", {"Participants": "Name <one@example.com>"}, 1),
            ("internal empty email", {"Participants": "one@example.com;;two@example.org"}, 1),
            ("separator-only email", {"Participants": ";;"}, 1),
            ("slash general comment", {"Comments": "30/07/2026: ready"}, 0),
            ("dot general comment", {"Comments": "30.07.2026: ready"}, 1),
            ("dot readiness comment", {"Readiness Comments": "30.07.2026: ready"}, 1),
            ("slash readiness comment", {"Readiness Comments": "30/07/2026: ready"}, 0),
        )
        for label, overrides, expected_check in scenarios:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(Path(temp), row_overrides={3: overrides})
                with closing(sqlite3.connect(":memory:")) as connection:
                    result = run_edct_analysis(
                        path,
                        connection=connection,
                        analysis_date=date(2026, 7, 30),
                    )
                self.assertEqual(result.row_results[("Supplier Level", 3)].check, expected_check)

    def test_catalogued_applicability_empty_and_accepted_value_policies(self) -> None:
        valid_portals = {
            "eSupplierConnect": "YES",
            "B2B": "NOT",
            "New supplier portal": "YES",
            "SPM": "NOT",
            "iTMS": "YES",
        }
        scenarios = (
            ("conditional email before trigger", {"Sales contact": "invalid"}, 0),
            (
                "conditional email after trigger",
                {
                    "Effective kick-off date": "30/07/2026",
                    "Sales contact": "invalid",
                    **valid_portals,
                },
                1,
            ),
            ("conditional cofor before trigger", {"Seller COFOR": "invalid"}, 0),
            (
                "conditional cofor after trigger",
                {
                    "Effective kick-off date": "30/07/2026",
                    "Seller COFOR": "invalid",
                    **valid_portals,
                },
                1,
            ),
            (
                "future effective date",
                {
                    "Effective kick-off date": "31/07/2026",
                    **valid_portals,
                },
                1,
            ),
            ("triple status before trigger", {"Triple Status": ""}, 0),
            (
                "triple status after trigger",
                {
                    "Cofor created date": "30/07/2026",
                    "Triple Status": "",
                    "EDI Mode": "WEB EDI",
                },
                1,
            ),
            ("overseas empty", {"Overseas": ""}, 0),
            ("overseas exact YES", {"Overseas": " YES "}, 0),
            ("overseas exact NOT", {"Overseas": "NOT"}, 0),
            ("overseas lowercase", {"Overseas": "yes"}, 1),
            ("overseas NO", {"Overseas": "NO"}, 1),
            ("supplier confirmation optional", {"Supplier Confimation": ""}, 0),
            ("supplier confirmation invalid", {"Supplier Confimation": "NO"}, 1),
            (
                "portals required after trigger",
                {"Effective kick-off date": "30/07/2026"},
                len(EDCT_PORTAL_COLUMNS),
            ),
            ("edi mode before trigger", {"EDI Mode": ""}, 0),
            (
                "edi mode after trigger",
                {
                    "Cofor created date": "30/07/2026",
                    "Triple Status": "Valid",
                    "EDI Mode": "",
                },
                1,
            ),
        )
        for label, overrides, expected_check in scenarios:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(Path(temp), row_overrides={3: overrides})
                with closing(sqlite3.connect(":memory:")) as connection:
                    result = run_edct_analysis(
                        path,
                        connection=connection,
                        analysis_date=date(2026, 7, 30),
                    )
                self.assertEqual(result.row_results[("Supplier Level", 3)].check, expected_check)

    def test_lifecycle_and_open_task_conditions_are_reported(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(
                directory,
                row_overrides={
                    3: {
                        "Cofor created date": "invalid date",
                        "Triple Status": "",
                        "EDI Mode": "",
                        "Effective kick-off date": "30/07/2026",
                        "Overseas": "YES",
                        "Supplier Confimation": "NO",
                        "OPEN TASK": "",
                    },
                    4: {"OPEN TASK": "NO"},
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                connection.row_factory = sqlite3.Row
                result = run_edct_analysis(
                    path, connection=connection, analysis_date=date(2026, 7, 30)
                )
                totals = connection.execute(
                    "SELECT column_name, fail_count FROM run_columns WHERE run_id = ?",
                    (result.run_id,),
                ).fetchall()
                output = export_edct_result(result, directory)

            exported = load_workbook(output)
            supplier = exported["Supplier Level"]
            headers = [cell.value for cell in supplier[2]]
            exported_checks = [
                supplier.cell(row, headers.index("Check") + 1).value for row in (3, 4)
            ]
            exported_comments = [
                supplier.cell(row, headers.index("Comment") + 1).value for row in (3, 4)
            ]
            exported.close()

        self.assertEqual(result.row_results[("Supplier Level", 3)].check, 10)
        self.assertEqual(result.row_results[("Supplier Level", 4)].check, 1)
        self.assertEqual(exported_checks, [10, 1])
        for column in (
            "Triple Status",
            "EDI Mode",
            "Supplier Confimation",
            "OPEN TASK",
            "eSupplierConnect",
            "B2B",
            "New supplier portal",
            "SPM",
            "iTMS",
        ):
            self.assertIn(column, result.row_results[("Supplier Level", 3)].comment)
            self.assertIn(column, exported_comments[0])
        self.assertIn("OPEN TASK", exported_comments[1])
        self.assertEqual(sum(row["fail_count"] for row in totals), 11)

    def test_formula_structure_uses_formula_reference_row(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={
                    4: {
                        "Starting date": '=IF(B4="","",B4)',
                        "Triplet COFOR": '= + if ( a4 = "" , "" , a4 )',
                    }
                },
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(
                    path, connection=connection, analysis_date=date(2026, 7, 30)
                )

        self.assertEqual(result.row_results[("Supplier Level", 3)].check, 0)
        self.assertEqual(result.row_results[("Supplier Level", 4)].check, 1)
        self.assertIn("Starting date", result.row_results[("Supplier Level", 4)].comment)
        self.assertNotIn("Triplet COFOR", result.row_results[("Supplier Level", 4)].comment)

    def test_formula_comparison_rejects_every_material_change(self) -> None:
        scenarios = (
            ("constant", "manual"),
            ("missing formula", None),
            ("changed function", '=AND(A4="")'),
            ("changed operator", '=IF(A4<>"","",A4)'),
            ("changed reference", '=IF(B4="","",B4)'),
            ("changed condition", '=IF(A4="x","",A4)'),
            ("changed quoted value", '=IF(A4="","changed",A4)'),
        )
        for label, formula in scenarios:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp:
                path = build_edct_workbook(
                    Path(temp),
                    row_overrides={4: {"Starting date": formula}},
                )
                with closing(sqlite3.connect(":memory:")) as connection:
                    result = run_edct_analysis(
                        path,
                        connection=connection,
                        analysis_date=date(2026, 7, 30),
                    )
                self.assertEqual(result.row_results[("Supplier Level", 4)].check, 1)
                self.assertIn("Starting date", result.row_results[("Supplier Level", 4)].comment)

    def test_missing_formula_reference_flags_every_assessed_row_and_continues(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = build_edct_workbook(
                Path(temp),
                row_overrides={3: {"Starting date": None}},
            )
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(
                    path, connection=connection, analysis_date=date(2026, 7, 30)
                )

        self.assertEqual(result.row_results[("Supplier Level", 3)].check, 1)
        self.assertEqual(result.row_results[("Supplier Level", 4)].check, 1)
