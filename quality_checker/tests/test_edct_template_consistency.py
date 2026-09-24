from __future__ import annotations

import sqlite3
import tempfile
from contextlib import closing
from pathlib import Path

from edct_support import EdctTestCase, build_edct_workbook
from openpyxl import load_workbook
from quality_checker.checkers.edct.export import export_edct_result
from quality_checker.checkers.edct.runner import run_edct_analysis


class TemplateConsistencyTests(EdctTestCase):
    def test_template_codes_are_assessed_and_exported(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(directory)
            workbook = load_workbook(path)
            template = workbook["Template-Cofor-Creation"]
            template["D3"] = " 1003.0 "
            template["E3"] = " Supplier 3 "
            template["D4"] = "absent"
            template["D5"] = None
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)
                self.addCleanup(result.workbook.close)
                self.assertEqual(result.row_results[("Template-Cofor-Creation", 3)].check, 0)
                missing = result.row_results[("Template-Cofor-Creation", 4)]
                self.assertEqual(missing.check, 1)
                self.assertIn("Punch Code = absent", missing.comment)
                self.assertNotIn(("Template-Cofor-Creation", 5), result.row_results)
                self.assertEqual(result.rule_totals[("template_supplier_code", "Punch Code")], 1)
                run = connection.execute(
                    "SELECT rows_total, rows_failed FROM runs WHERE id = ?", (result.run_id,)
                ).fetchone()
                self.assertEqual(run[0], len(result.assessed_rows))
                self.assertEqual(run[1], result.rows_failed)
                output = export_edct_result(result, directory)
            exported = load_workbook(output)
            self.addCleanup(exported.close)
            sheet = exported["Template-Cofor-Creation"]
            headers = {cell.value: cell.column for cell in sheet[2] if cell.value}
            self.assertEqual(sheet.cell(3, headers["Check"]).value, 0)
            self.assertEqual(sheet.cell(4, headers["Check"]).value, 1)
            self.assertIn("Punch Code = absent", sheet.cell(4, headers["Comment"]).value)

    def test_company_matches_any_supplier_for_code_and_missing_name_fails(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            path = build_edct_workbook(
                directory,
                row_overrides={
                    4: {"Index": "", "Supplier Punch code": 1003, "Supplier name": "Second Co"},
                },
            )
            workbook = load_workbook(path)
            supplier = workbook["Supplier Level"]
            supplier_headers = {cell.value: cell.column for cell in supplier[2] if cell.value}
            supplier.cell(5, supplier_headers["Supplier Punch code"], "unique")
            supplier.cell(5, supplier_headers["Supplier name"], "Only Unindexed")
            supplier.cell(6, supplier_headers["Supplier Punch code"], "unnamed")
            template = workbook["Template-Cofor-Creation"]
            template["D3"], template["E3"] = 1003, " second co "
            template["D4"], template["E4"] = 1003, "unknown"
            template["D5"], template["E5"] = 1003, None
            template["D6"], template["E6"] = "missing", "anything"
            template["D7"], template["E7"] = "unique", "only unindexed"
            template["D8"], template["E8"] = "unnamed", None
            workbook.save(path)
            workbook.close()
            with closing(sqlite3.connect(":memory:")) as connection:
                result = run_edct_analysis(path, connection=connection)
                self.addCleanup(result.workbook.close)
                self.assertEqual(result.row_results[("Template-Cofor-Creation", 3)].check, 0)
                for row in (4, 5):
                    failure = result.row_results[("Template-Cofor-Creation", row)]
                    self.assertEqual(failure.check, 1)
                    self.assertIn("Company name =", failure.comment)
                self.assertEqual(result.row_results[("Template-Cofor-Creation", 6)].check, 1)
                self.assertEqual(result.row_results[("Template-Cofor-Creation", 7)].check, 0)
                self.assertEqual(result.row_results[("Template-Cofor-Creation", 8)].check, 1)
                self.assertEqual(result.rule_totals[("template_supplier_name", "Company name")], 3)
                output = export_edct_result(result, directory)
            exported = load_workbook(output)
            self.addCleanup(exported.close)
            sheet = exported["Template-Cofor-Creation"]
            headers = {cell.value: cell.column for cell in sheet[2] if cell.value}
            self.assertIn("Company name = unknown", sheet.cell(4, headers["Comment"]).value)
